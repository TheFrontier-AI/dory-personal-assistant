"""Data studio: read a spreadsheet / CSV / JSON into tables, profile it, query it with SQL and chart it.

Pure logic (no Discord, no model): `data_bot.py` binds it to a request as agent tools.

Loading (`load_tables`): xlsx / xlsm (each sheet is a table), legacy xls (xlrd), csv / tsv (delimiter
sniffed; utf-8, utf-8-sig, cp874, latin-1), json (an array of objects). Table names are `t_` + an
ascii slug of the sheet name, column names an ascii slug of the header (a header with no ascii
letters, e.g. Thai, becomes `col_N`); the original text is kept as `label`.

Type rules (documented because the answers depend on them):

* null: empty, "-", dashes, "n/a", "null", "none", "nan", Excel error text;
* number: "1,234.50", "฿1,200", "1 200 บาท", "(1,200)" = -1200, "12%" = 0.12 (a ratio; the column
  gets unit "%"). A column is numeric when at least 90% of its non-null cells are; the rest become
  null and count as "mixed types". Digit strings with a leading zero or 11+ digits stay text (ids);
* date: dd/mm/yyyy (day first, the Thai and European order), yyyy-mm-dd, dd-mm-yyyy, dd.mm.yyyy,
  "15 ต.ค. 2568", "15 Oct 2025", with optional time. A year from 2400 to 2700 is Buddhist
  (Thai) and gets 543 subtracted: 15/10/2568 is 2025-10-15. Same for real Excel date cells;
* everything else is text.

Caps: file 50 MB, 200,000 cells per table, 20 tables, 200 columns; beyond a cap the first rows are
loaded and the profile says so.
"""

from __future__ import annotations

import base64
import contextlib
import csv
import gzip
import io
import json
import logging
import math
import re
import secrets
import threading
import time
import unicodedata
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pandas as pd

from . import bizfin
from .coo_store import DocStore
from .i18n import CATALOG, t

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_CELLS = 200_000
MAX_TABLES = 20
MAX_COLS = 200
MAX_SQL_COLS = 40
SQL_LIMIT = 200
EXPORT_LIMIT = 50_000
SQL_TIMEOUT_S = 10.0
SQL_MEMORY = "512MB"
SHOW_ROWS = 25
PROFILE_LINES = 40
MAX_CATEGORIES = 30
CHART_W, CHART_H = 1200, 700
CHART_MAX_BYTES = 1_000_000
DATA_EXTS = ("xlsx", "xlsm", "xls", "csv", "tsv", "json")
DATASETS = "dory_datasets"
DATASET_DATA = "dory_dataset_data"
MAX_DATASETS = 30
MAX_STORED_BYTES = 3 * 1024 * 1024  # gzip, all tables of one dataset
CHUNK_CHARS = 600_000  # base64 per Firestore document (1 MiB limit per document)


class DataError(ValueError):
    """A reason the user can be told (the message is English, for the model)."""


# ----------------------------------------------------------------------------- playbook
_PLAYBOOK: str | None = None


def playbook() -> str:
    """The data-analysis playbook for the system prompt (`skills/data_analysis.md`, read once)."""
    global _PLAYBOOK
    if _PLAYBOOK is None:
        path = Path(__file__).with_name("skills") / "data_analysis.md"
        _PLAYBOOK = "\n" + path.read_text(encoding="utf-8").strip() + "\n"
    return _PLAYBOOK


# ----------------------------------------------------------------------------- strings
STRINGS: dict[str, dict[str, str]] = {
    "data_other": {"th": "อื่นๆ", "en": "Other"},
    "data_note_rows": {
        "th": "ตาราง {table}: โหลดเฉพาะ {n} แถวแรก (ไฟล์ใหญ่เกินกำหนด)",
        "en": "Table {table}: only the first {n} rows were loaded (size cap).",
    },
    "data_note_tables": {
        "th": "โหลดเฉพาะ {n} ตารางแรก (ชีตที่เหลือถูกข้าม)",
        "en": "Only the first {n} tables were loaded; the other sheets were skipped.",
    },
    "data_note_cols": {
        "th": "ตาราง {table}: โหลดเฉพาะ {n} คอลัมน์แรก",
        "en": "Table {table}: only the first {n} columns were loaded.",
    },
    "data_note_noheader": {
        "th": "ตาราง {table}: หาแถวหัวตารางไม่เจอ ตั้งชื่อคอลัมน์ให้เองเป็น col_N",
        "en": "Table {table}: no header row found; columns were named col_N.",
    },
    "data_prof_head": {
        "th": "{name}{label}: {rows} แถว × {cols} คอลัมน์",
        "en": "{name}{label}: {rows} rows × {cols} columns",
    },
    "data_prof_nulls": {"th": "ว่าง {pct}%", "en": "{pct}% empty"},
    "data_prof_distinct": {"th": "{n} ค่าไม่ซ้ำ", "en": "{n} distinct"},
    "data_prof_mean": {"th": "เฉลี่ย {v}", "en": "mean {v}"},
    "data_prof_more_cols": {"th": "… อีก {n} คอลัมน์", "en": "… {n} more columns"},
    "data_prof_issues": {"th": "ข้อสังเกต", "en": "Issues"},
    "data_issue_duplicate_header": {
        "th": "หัวคอลัมน์ซ้ำ: {column}",
        "en": "duplicate header: {column}",
    },
    "data_issue_mixed_types": {
        "th": "{column}: มี {n} ค่าที่ไม่ตรงชนิดและถูกตัดเป็นค่าว่าง",
        "en": "{column}: {n} values did not fit the type and became empty",
    },
    "data_issue_empty_column": {"th": "คอลัมน์ว่าง: {column}", "en": "empty column: {column}"},
    "data_issue_total_row": {
        "th": "แถวที่น่าจะเป็นผลรวม: แถว {rows}",
        "en": "likely total row(s): {rows}",
    },
    "data_sql_more_rows": {"th": "… และอีก {n} แถว", "en": "… and {n} more rows"},
    "data_sql_capped": {
        "th": "(แสดงสูงสุด {n} แถวจากผลลัพธ์)",
        "en": "(result capped at {n} rows)",
    },
    "data_sql_empty": {"th": "(ไม่มีแถวในผลลัพธ์)", "en": "(no rows)"},
    "data_sql_cols_cut": {
        "th": "(แสดงเฉพาะ {n} คอลัมน์แรก)",
        "en": "(only the first {n} columns are shown)",
    },
    "data_err_not_select": {
        "th": "ใช้ได้เฉพาะคำสั่ง SELECT (หรือ WITH ... SELECT) คำสั่งเดียว อ่านอย่างเดียว",
        "en": "Only a single read-only SELECT (or WITH ... SELECT) statement is allowed.",
    },
    "data_err_empty_query": {"th": "ยังไม่มีคำสั่ง SQL", "en": "The SQL query is empty."},
    "data_err_unknown_column": {
        "th": "ไม่พบคอลัมน์นี้ ({detail}) คอลัมน์ที่มี: {columns}",
        "en": "Unknown column ({detail}). Available columns: {columns}",
    },
    "data_err_unknown_table": {
        "th": "ไม่พบตารางนี้ ({detail}) ตารางที่มี: {tables}",
        "en": "Unknown table ({detail}). Available tables: {tables}",
    },
    "data_err_timeout": {
        "th": "คำสั่งใช้เวลานานเกิน {n} วินาที จึงหยุดไว้ ลองกรองข้อมูลให้แคบลง",
        "en": "The query ran longer than {n} seconds and was stopped. Try filtering more.",
    },
    "data_err_memory": {
        "th": "คำสั่งใช้หน่วยความจำมากเกินไป ลองกรองหรือสรุปข้อมูลก่อน",
        "en": "The query needs too much memory. Filter or aggregate first.",
    },
    "data_err_sql": {"th": "SQL ผิดพลาด: {detail}", "en": "SQL error: {detail}"},
    "data_err_chart_col": {
        "th": "ไม่พบคอลัมน์ {column} ในผลลัพธ์ คอลัมน์ที่มี: {columns}",
        "en": "Column {column} is not in the result. Available columns: {columns}",
    },
    "data_err_chart_kind": {
        "th": "ชนิดกราฟที่ใช้ได้: bar, line, pie, scatter, hist",
        "en": "Chart kinds: bar, line, pie, scatter, hist.",
    },
    "data_err_chart_numeric": {
        "th": "กราฟนี้ต้องใช้คอลัมน์ตัวเลข: {column}",
        "en": "This chart needs a numeric column: {column}",
    },
    "data_err_chart_empty": {"th": "ไม่มีข้อมูลให้วาดกราฟ", "en": "There is no data to chart."},
    "data_err_chart_pie": {
        "th": "กราฟวงกลมใช้กับค่าที่ไม่ติดลบเท่านั้น",
        "en": "A pie chart needs non-negative values.",
    },
}
CATALOG.update(STRINGS)


# ----------------------------------------------------------------------------- model
@dataclass
class Column:
    name: str
    label: str
    type: str = "text"  # integer number bool date datetime text
    unit: str = ""  # "%" when the values are ratios read from "12%"
    bad: int = 0  # non-null cells that did not fit the type (now null)

    def dump(self) -> dict:
        return {
            "name": self.name,
            "label": self.label,
            "type": self.type,
            "unit": self.unit,
            "bad": self.bad,
        }


@dataclass
class Table:
    name: str
    label: str
    frame: pd.DataFrame
    columns: list[Column]
    truncated: int = 0  # rows loaded when a cap cut the table, else 0
    no_header: bool = False
    dup_headers: list[str] = field(default_factory=list)
    total_rows: list[int] = field(default_factory=list)  # 1-based rows that look like totals

    @property
    def rows(self) -> int:
        return len(self.frame)


@dataclass
class Dataset:
    filename: str
    tables: list[Table]
    notes: list[dict] = field(default_factory=list)  # {"code", ...params}

    def table(self, name: str) -> Table | None:
        want = (name or "").strip().casefold()
        return next(
            (x for x in self.tables if x.name.casefold() == want or x.label.casefold() == want),
            None,
        )


# ----------------------------------------------------------------------------- slugs
_RESERVED: set[str] | None = None
_FALLBACK_RESERVED = {
    "select", "from", "where", "group", "order", "by", "limit", "table", "all", "and", "or",
    "not", "in", "is", "as", "on", "join", "case", "when", "then", "else", "end", "union",
    "having", "distinct", "default", "check", "column", "constraint", "create", "primary",
    "references", "unique", "using", "with", "for", "offset", "desc", "asc", "between", "like",
    "null", "true", "false", "cast", "left", "right", "inner", "outer", "full", "cross",
}  # fmt: skip


def _reserved() -> set[str]:
    global _RESERVED
    if _RESERVED is None:
        try:
            import duckdb

            con = duckdb.connect(":memory:")
            rows = con.execute(
                "SELECT keyword_name FROM duckdb_keywords() WHERE keyword_category='reserved'"
            ).fetchall()
            con.close()
            _RESERVED = {r[0].lower() for r in rows} | _FALLBACK_RESERVED
        except Exception:
            _RESERVED = set(_FALLBACK_RESERVED)
    return _RESERVED


def slug(text: str, fallback: str) -> str:
    """Lower-case ascii identifier from `text`; `fallback` when no ascii letter or digit is left."""
    norm = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9]+", "_", norm.lower()).strip("_")[:40].strip("_")
    return s or fallback


def _unique(base: str, used: set[str]) -> str:
    name, n = base, 2
    while name in used:
        name = f"{base}_{n}"
        n += 1
    used.add(name)
    return name


def column_name(label: str, index: int, used: set[str]) -> str:
    s = slug(label, f"col_{index}")
    if s[0].isdigit():
        s = "c_" + s
    if s in _reserved():
        s += "_"
    return _unique(s, used)


# ----------------------------------------------------------------------------- value parsing
_NULLS = {"", "-", "–", "—", "n/a", "null", "none", "nan", "#n/a", "#div/0!", "#value!", "#ref!"}
_CUR = re.compile(r"฿|\$|€|£|¥|บาท|\bthb\b|\busd\b|\beur\b|\bbaht\b", re.I)
_PLAIN = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
_THOUS = re.compile(r"^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$")
_MONTHS = {
    "ม.ค.": 1, "ก.พ.": 2, "มี.ค.": 3, "เม.ย.": 4, "พ.ค.": 5, "มิ.ย.": 6, "ก.ค.": 7, "ส.ค.": 8,
    "ก.ย.": 9, "ต.ค.": 10, "พ.ย.": 11, "ธ.ค.": 12,
    "มกราคม": 1, "กุมภาพันธ์": 2, "มีนาคม": 3, "เมษายน": 4, "พฤษภาคม": 5, "มิถุนายน": 6,
    "กรกฎาคม": 7, "สิงหาคม": 8, "กันยายน": 9, "ตุลาคม": 10, "พฤศจิกายน": 11, "ธันวาคม": 12,
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9,
    "sept": 9, "oct": 10, "nov": 11, "dec": 12, "january": 1, "february": 2, "march": 3,
    "april": 4, "june": 6, "july": 7, "august": 8, "september": 9, "october": 10,
    "november": 11, "december": 12,
}  # fmt: skip
_DATE_NUM = re.compile(
    r"^(\d{1,4})[/\-.](\d{1,2})[/\-.](\d{1,4})(?:[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$"
)
_DATE_TXT = re.compile(
    r"^(\d{1,2})\s+([^\d\s]+?)\.?\s+(\d{4})(?:[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$"
)
_TOTAL_WORD = re.compile(
    r"^\s*(grand\s+|sub\s*-?\s*)?total\b|^\s*sum\s*$|^\s*รวม|^\s*ยอดรวม|^\s*ผลรวม", re.I
)


def is_null(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str):
        return v.strip().casefold() in _NULLS
    return False


def buddhist_fix(year: int) -> int:
    """Years 2400-2700 are Buddhist-era (BE = CE + 543): 2568 -> 2025."""
    return year - 543 if 2400 <= year <= 2700 else year


def parse_number(v: Any) -> tuple[float, bool] | None:
    """(value, was_percent) for a number or a number-like text, else None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int | float | Decimal):
        f = float(v)
        return None if math.isnan(f) or math.isinf(f) else (f, False)
    if not isinstance(v, str):
        return None
    s = v.strip().replace(" ", " ")
    neg = s.startswith("(") and s.endswith(")")
    if neg:
        s = s[1:-1]
    s = _CUR.sub("", s).strip()
    pct = s.endswith("%")
    if pct:
        s = s[:-1].strip()
    s = re.sub(r"(?<=\d) (?=\d{3}\b)", "", s)  # "1 200" -> "1200"
    if _THOUS.match(s):
        s = s.replace(",", "")
    if not _PLAIN.match(s):
        return None
    f = float(s)
    if neg:
        f = -f
    return (f / 100 if pct else f), pct


def parse_date(v: Any) -> datetime | None:
    """A datetime for a date cell or date-like text (day first; Buddhist years fixed), else None."""
    if isinstance(v, datetime):
        return _fix_year(v)
    if isinstance(v, date):
        return _fix_year(datetime(v.year, v.month, v.day))
    if not isinstance(v, str):
        return None
    s = v.strip()
    try:
        if m := _DATE_NUM.match(s):
            a, b, c = int(m[1]), int(m[2]), int(m[3])
            if len(m[1]) == 4:
                y, mo, d = a, b, c
            elif len(m[3]) == 4:
                d, mo, y = a, b, c
            elif len(m[3]) == 2:
                d, mo, y = a, b, 2000 + c
            else:
                return None
        elif m := _DATE_TXT.match(s):
            mo = _MONTHS.get(m[2].casefold().rstrip("."), _MONTHS.get(m[2].casefold() + "."))
            if mo is None:
                return None
            d, y = int(m[1]), int(m[3])
        else:
            return None
        hh, mm, ss = int(m[4] or 0), int(m[5] or 0), int(m[6] or 0)
        return datetime(buddhist_fix(y), mo, d, hh, mm, ss)
    except ValueError:
        return None


def _fix_year(dt: datetime) -> datetime:
    if 2400 <= dt.year <= 2700:
        try:
            return dt.replace(year=dt.year - 543)
        except ValueError:
            return dt
    return dt


def _text(v: Any) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    return str(v).strip()


def infer_column(vals: list, skip: frozenset[int] = frozenset()) -> tuple[str, list, str, int]:
    """(type, converted values, unit, bad count) for one column of raw cell values. Rows in `skip`
    (likely total rows) do not vote on the type and never count as bad; they are converted too."""
    clean = [None if is_null(v) else v for v in vals]
    nn = [v for i, v in enumerate(clean) if v is not None and i not in skip]
    if not nn:
        return "text", [None if v is None else _text(v) for v in clean], "", 0
    kind = _decide(nn)
    out, unit = _convert(kind, clean, nn)
    bad = 0
    if kind != "text":
        bad = sum(1 for i, (v, o) in enumerate(zip(clean, out, strict=True)) if v is not None and o is None and i not in skip)  # fmt: skip
    return kind, out, unit, bad


def _decide(nn: list) -> str:
    n = len(nn)
    if all(isinstance(v, bool) for v in nn):
        return "bool"
    idish = [v.strip() for v in nn if isinstance(v, str) and v.strip().isdigit()]
    if any((len(s) > 1 and s[0] == "0") or len(s) >= 11 for s in idish):
        return "text"  # ids, phone numbers, barcodes
    if sum(isinstance(v, datetime | date) for v in nn) >= 0.9 * n:
        return "datetime" if _has_time([parse_date(v) for v in nn]) else "date"
    if sum(parse_number(v) is not None for v in nn) >= 0.9 * n:
        nums = [parse_number(v) for v in nn]
        integral = all(
            x is None or (x[0].is_integer() and not x[1] and abs(x[0]) < 2**53) for x in nums
        )
        return "integer" if integral else "number"
    dates = [parse_date(v) for v in nn]
    if sum(d is not None for d in dates) >= 0.9 * n:
        return "datetime" if _has_time(dates) else "date"
    return "text"


def _has_time(vals: list) -> bool:
    return any(v is not None and (v.hour or v.minute or v.second) for v in vals)


def _convert(kind: str, clean: list, nn: list) -> tuple[list, str]:
    if kind == "bool":
        return [None if v is None or not isinstance(v, bool) else v for v in clean], ""
    if kind in ("integer", "number"):
        nums = [None if v is None else parse_number(v) for v in clean]
        pct = sum(1 for x in nums if x is not None and x[1])
        ok = sum(x is not None for x in nums)
        unit = "%" if pct and pct >= 0.5 * ok else ""
        return [None if x is None else x[0] for x in nums], unit
    if kind in ("date", "datetime"):
        return [None if v is None else parse_date(v) for v in clean], ""
    return [None if v is None else _text(v) for v in clean], ""


def _series(kind: str, vals: list) -> pd.Series:
    if kind == "integer":
        return pd.Series([None if v is None else int(v) for v in vals], dtype="Int64")
    if kind == "number":
        return pd.Series(vals, dtype="float64")
    if kind == "bool":
        return pd.Series(vals, dtype="boolean")
    if kind in ("date", "datetime"):
        return pd.Series(pd.to_datetime(pd.Series(vals, dtype="object")), dtype="datetime64[us]")
    return pd.Series(vals, dtype="object")


# ----------------------------------------------------------------------------- tables from rows
def _label_like(v: Any) -> bool:
    """A cell that reads as a header label: text that is not a number or a date."""
    return (
        isinstance(v, str) and bool(v.strip()) and parse_number(v) is None and parse_date(v) is None
    )


def _find_header(rows: list[list]) -> int | None:
    for i, row in enumerate(rows):
        cells = [c for c in row if not is_null(c) or isinstance(c, str) and c.strip()]
        if len(cells) < 2:
            continue
        if sum(_label_like(c) for c in cells) < 0.8 * len(cells):
            continue
        later = [sum(1 for c in r if not is_null(c)) for r in rows[i + 1 : i + 6]]
        if later and len(cells) < 0.5 * max(later):
            continue  # a title strip over a wider table
        return i
    return None


def _trim(rows: list[list]) -> list[list]:
    out = []
    for r in rows:
        r = list(r)
        while r and (r[-1] is None or isinstance(r[-1], str) and not r[-1].strip()):
            r.pop()
        out.append(r)
    return out


def _is_total_row(row: list) -> bool:
    first = next((c for c in row if not is_null(c)), None)
    return isinstance(first, str) and bool(_TOTAL_WORD.search(first))


def build_table(
    label: str, raw_rows: list[list], tname: str, *, header: int | None = -1, seen_all: bool = True
) -> Table | None:
    """One table from raw rows. `header` -1 = detect, an index = that row, None = no header row.
    None when the sheet has nothing in it."""
    rows = [r for r in _trim(raw_rows) if any(not is_null(c) for c in r)]
    if not rows:
        return None
    no_header = False
    if header == -1:
        idx = _find_header(rows)
        if idx is None:
            no_header = True
            head_cells: list = []
            data = rows
        else:
            head_cells, data = rows[idx], rows[idx + 1 :]
    elif header is None:
        no_header, head_cells, data = True, [], rows
    else:
        head_cells, data = rows[header], rows[header + 1 :]
    width = max(len(head_cells), *[len(r) for r in data])
    keep = [
        j
        for j in range(width)
        if (j < len(head_cells) and not is_null(head_cells[j]))
        or any(j < len(r) and not is_null(r[j]) for r in data)
    ]
    cut = 0
    if len(keep) > MAX_COLS:
        cut, keep = len(keep), keep[:MAX_COLS]
    cap_rows = max(1, MAX_CELLS // max(1, len(keep)))
    truncated = 0
    if len(data) > cap_rows or not seen_all:
        data, truncated = data[:cap_rows], min(len(data), cap_rows)
    total_idx = [i for i, r in enumerate(data) if _is_total_row(r)]
    used: set[str] = set()
    cols: list[Column] = []
    frames: dict[str, pd.Series] = {}
    labels: list[str] = []
    dups: list[str] = []
    for pos, j in enumerate(keep, 1):
        head = head_cells[j] if j < len(head_cells) else None
        lab = "" if is_null(head) else " ".join(str(head).split())
        if lab and lab.casefold() in (x.casefold() for x in labels):
            dups.append(lab)
        labels.append(lab)
        name = column_name(lab, pos, used)
        kind, vals, unit, bad = infer_column(
            [r[j] if j < len(r) else None for r in data], frozenset(total_idx)
        )
        cols.append(Column(name, lab or name, kind, unit, bad))
        frames[name] = _series(kind, vals)
    frame = pd.DataFrame(frames, index=pd.RangeIndex(len(data)))
    tb = Table(
        tname, label, frame, cols, truncated=truncated, no_header=no_header, dup_headers=dups,
        total_rows=[i + 1 for i in total_idx[:3]],
    )  # fmt: skip
    tb._cut_cols = cut  # type: ignore[attr-defined]
    return tb


# ----------------------------------------------------------------------------- readers
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp874"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _sniff(sample: str, ext: str) -> str:
    if ext == "tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in ",;\t|"}
        return max(counts, key=counts.get) if any(counts.values()) else ","


def _read_csv(data: bytes, ext: str) -> list[list]:
    text = _decode(data)
    delim = _sniff(text[:8192], ext)
    csv.field_size_limit(10_000_000)
    return [
        [c if c.strip() else None for c in row]
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    ]


def _read_json(data: bytes) -> list[list]:
    try:
        obj = json.loads(_decode(data))
    except ValueError as e:
        raise DataError("the file is not valid JSON") from e
    if isinstance(obj, dict):
        lists = [v for v in obj.values() if isinstance(v, list)]
        obj = lists[0] if len(lists) == 1 else [obj]
    if not isinstance(obj, list) or not obj or not all(isinstance(x, dict) for x in obj):
        raise DataError("JSON must be an array of objects")
    keys: list[str] = []
    for x in obj:
        for k in x:
            if k not in keys:
                keys.append(k)
    keys = keys[:MAX_COLS]

    def cell(v):
        return json.dumps(v, ensure_ascii=False) if isinstance(v, dict | list) else v

    return [list(keys)] + [[cell(x.get(k)) for k in keys] for x in obj]


def _xlsx_sheets(data: bytes) -> list[tuple[str, list[list], bool]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise DataError("the workbook could not be opened (corrupt or password-protected)") from e
    out = []
    try:
        for ws in wb.worksheets[:MAX_TABLES]:
            rows: list[list] = []
            limit = MAX_CELLS + 60
            complete = True
            for n, row in enumerate(ws.iter_rows(values_only=True), 1):
                rows.append(list(row[: MAX_COLS * 2]))
                if n == 60:
                    wide = max((len(_trim([r])[0]) for r in rows), default=1)
                    limit = MAX_CELLS // max(1, min(wide, MAX_COLS)) + 60
                if n > limit:
                    complete = False
                    break
            out.append((ws.title, rows, complete))
        extra = len(wb.worksheets) - MAX_TABLES
    finally:
        wb.close()
    return [*out, ("", [], extra > 0)] if extra > 0 else out


def _xls_sheets(data: bytes) -> list[tuple[str, list[list], bool]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as e:
        raise DataError("the .xls file could not be read; save it as .xlsx or CSV") from e
    out = []
    for sh in book.sheets()[:MAX_TABLES]:
        rows = []
        for i in range(min(sh.nrows, MAX_CELLS + 60)):
            row = []
            for c in sh.row(i)[: MAX_COLS * 2]:
                if c.ctype in (0, 5, 6):
                    row.append(None)
                elif c.ctype == 3:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode))
                    except Exception:
                        row.append(None)
                elif c.ctype == 4:
                    row.append(bool(c.value))
                else:
                    row.append(c.value)
            rows.append(row)
        out.append((sh.name, rows, sh.nrows <= MAX_CELLS + 60))
    return out


def file_ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""


def is_data_file(filename: str) -> bool:
    return file_ext(filename) in DATA_EXTS


def load_tables(data: bytes, filename: str) -> Dataset:
    """Tables of an xlsx / xls / csv / tsv / json file. Raises DataError with a reason."""
    if len(data) > MAX_FILE_BYTES:
        raise DataError("the file is over 50 MB")
    ext = file_ext(filename)
    if ext in ("xlsx", "xlsm"):
        sheets = _xlsx_sheets(data)
    elif ext == "xls":
        sheets = _xls_sheets(data)
    elif ext in ("csv", "tsv", "txt"):
        rows = _read_csv(data, ext)
        sheets = [(filename.rsplit(".", 1)[0] or "data", rows[: MAX_CELLS + 60], True)]
        sheets[0] = (sheets[0][0], sheets[0][1], len(rows) <= MAX_CELLS + 60)
    elif ext == "json":
        rows = _read_json(data)
        sheets = [(filename.rsplit(".", 1)[0] or "data", rows[: MAX_CELLS + 60], len(rows) <= MAX_CELLS + 60)]  # fmt: skip
    else:
        raise DataError("unsupported file type (xlsx, xls, csv, tsv and json can be analysed)")
    ds = Dataset(filename, [])
    used: set[str] = set()
    for pos, (label, rows, complete) in enumerate(sheets, 1):
        if not label and not rows:  # the marker of skipped sheets
            ds.notes.append({"code": "tables", "n": MAX_TABLES})
            continue
        tname = _unique("t_" + slug(label, str(pos)).removeprefix("t_"), used)
        hdr: int | None = 0 if ext == "json" else -1
        tb = build_table(label or tname, rows, tname, header=hdr, seen_all=complete)
        if tb is None:
            continue
        if tb.truncated:
            ds.notes.append({"code": "rows", "table": tb.name, "n": tb.truncated})
        if getattr(tb, "_cut_cols", 0):
            ds.notes.append({"code": "cols", "table": tb.name, "n": MAX_COLS})
        if tb.no_header:
            ds.notes.append({"code": "noheader", "table": tb.name})
        ds.tables.append(tb)
    if not ds.tables:
        raise DataError("no data found in the file")
    return ds


def quick_counts(data: bytes, filename: str) -> tuple[int, int]:
    """(sheets, rows) of a data file without building tables: for the one line the agent gets next
    to the file text. Raises DataError when the file cannot be opened."""
    ext = file_ext(filename)
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise DataError("the workbook could not be opened") from e
        try:
            return len(wb.worksheets), sum(int(ws.max_row or 0) for ws in wb.worksheets)
        finally:
            wb.close()
    if ext == "xls":
        import xlrd

        try:
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
        except Exception as e:
            raise DataError("the .xls file could not be read") from e
        return book.nsheets, sum(sh.nrows for sh in book.sheets())
    if ext in ("csv", "tsv", "txt"):
        return 1, max(0, data.count(b"\n"))
    if ext == "json":
        return 1, max(0, data.count(b"},"))
    raise DataError("unsupported file type")


# ----------------------------------------------------------------------------- profile
def _num(v: Any):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    f = float(v)
    return int(f) if f.is_integer() and abs(f) < 2**53 else round(f, 4)


def _py(v: Any) -> Any:
    """A plain Python value (json-friendly except dates, which stay datetime / date)."""
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, Decimal):
        return float(v)
    if hasattr(v, "item") and not isinstance(v, str | bytes):
        try:
            v = v.item()
        except Exception:
            return str(v)
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    return v


def jsonable(v: Any) -> Any:
    v = _py(v)
    if isinstance(v, datetime):
        return v.isoformat(sep=" ") if (v.hour or v.minute or v.second) else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bytes):
        return v.hex()
    if isinstance(v, list | tuple | dict | set):
        return str(v)
    return v


def _issues(tb: Table) -> list[dict]:
    out: list[dict] = []
    for d in tb.dup_headers:
        out.append({"code": "duplicate_header", "column": d})
    for c in tb.columns:
        if c.bad:
            out.append({"code": "mixed_types", "column": c.name, "n": c.bad})
        if tb.rows and tb.frame[c.name].isna().all():
            out.append({"code": "empty_column", "column": c.name})
    if tb.total_rows:
        out.append({"code": "total_row", "rows": tb.total_rows})
    return out


def _col_profile(tb: Table, c: Column) -> dict:
    s = tb.frame[c.name]
    n = len(s)
    nulls = int(s.isna().sum())
    d: dict = {
        "name": c.name,
        "label": c.label,
        "type": c.type,
        "unit": c.unit,
        "null_pct": round(100 * nulls / n, 1) if n else 0.0,
        "distinct": int(s.nunique(dropna=True)),
    }
    nn = s.dropna()
    if len(nn):
        if c.type in ("integer", "number"):
            f = nn.astype("float64")
            d.update(min=_num(f.min()), max=_num(f.max()), mean=_num(f.mean()), median=_num(f.median()))  # fmt: skip
        elif c.type in ("date", "datetime"):
            d.update(min=jsonable(nn.min()), max=jsonable(nn.max()))
        elif c.type == "text" and d["distinct"] <= 30:
            top = nn.value_counts().head(5)
            d["top"] = [[str(k), int(v)] for k, v in top.items()]
    return d


def profile(ds: Dataset, lang: str = "en") -> dict:
    """Per table: rows, columns (name, label, type, null %, distinct, min / max / mean / median or
    date range or top 5 values), 5 sample rows and detected issues."""
    tables = []
    for tb in ds.tables:
        sample = [
            {c.name: jsonable(tb.frame.at[i, c.name]) for c in tb.columns}
            for i in range(min(5, tb.rows))
        ]
        tables.append(
            {
                "name": tb.name,
                "label": tb.label,
                "rows": tb.rows,
                "truncated": bool(tb.truncated),
                "columns": [_col_profile(tb, c) for c in tb.columns],
                "sample": sample,
                "issues": _issues(tb),
            }
        )
    return {
        "filename": ds.filename,
        "notes": [note_text(n, lang) for n in ds.notes],
        "tables": tables,
    }


def note_text(note: dict, lang: str) -> str:
    return t("data_note_" + note["code"], lang, **{k: v for k, v in note.items() if k != "code"})


def issue_text(issue: dict, lang: str) -> str:
    p = {k: v for k, v in issue.items() if k != "code"}
    if "rows" in p:
        p["rows"] = ", ".join(str(x) for x in p["rows"])
    return t("data_issue_" + issue["code"], lang, **p)


def fmt_value(v: Any) -> str:
    v = jsonable(v)
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).lower()
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        s = f"{v:,.2f}"
        return s.rstrip("0").rstrip(".") if "." in s else s
    return str(v)


def profile_md(ds: Dataset, lang: str = "th", prof: dict | None = None) -> str:
    """The profile as at most PROFILE_LINES lines of Markdown for Discord."""
    prof = prof or profile(ds, lang)
    lines: list[str] = list(prof["notes"][:3])
    tabs = prof["tables"]
    room = max(PROFILE_LINES - len(lines), 4)
    per = max(1, room // max(1, len(tabs)))
    for tp in tabs:
        label = f" ({tp['label']})" if tp["label"] and tp["label"] != tp["name"] else ""
        block = [
            "**"
            + t(
                "data_prof_head",
                lang,
                name=tp["name"],
                label=label,
                rows=f"{tp['rows']:,}",
                cols=len(tp["columns"]),
            )  # fmt: skip
            + "**"
        ]
        issues = [issue_text(i, lang) for i in tp["issues"]]
        if issues:
            block.append(f"{t('data_prof_issues', lang)}: " + "; ".join(issues[:3]))
        cols = tp["columns"]
        slots = max(1, per - len(block) - 1)
        shown = cols if len(cols) <= slots else cols[: max(1, slots - 1)]
        for c in shown:
            block.append(_col_line(c, lang))
        if len(shown) < len(cols):
            block.append(t("data_prof_more_cols", lang, n=len(cols) - len(shown)))
        lines += block
    return "\n".join(lines[:PROFILE_LINES])


def _col_line(c: dict, lang: str) -> str:
    head = f"- `{c['name']}`" + (
        f" ({c['label']})" if c["label"] and c["label"] != c["name"] else ""
    )
    bits = [c["type"] + (" %" if c.get("unit") else "")]
    if "mean" in c:
        bits.append(f"{fmt_value(c['min'])} → {fmt_value(c['max'])}")
        bits.append(t("data_prof_mean", lang, v=fmt_value(c["mean"])))
    elif c["type"] in ("date", "datetime") and "min" in c:
        bits.append(f"{c['min']} → {c['max']}")
    elif c.get("top"):
        bits.append(", ".join(f"{k} ({n})" for k, n in c["top"][:3]))
    if c["null_pct"]:
        bits.append(t("data_prof_nulls", lang, pct=fmt_value(c["null_pct"])))
    if not c.get("top"):
        bits.append(t("data_prof_distinct", lang, n=f"{c['distinct']:,}"))
    return f"{head} · " + " · ".join(bits)


# ----------------------------------------------------------------------------- SQL
@dataclass
class Result:
    columns: list[str]
    rows: list[list]  # native Python values (datetime / date stay dates)
    truncated: bool  # more rows existed than `limit`
    cols_cut: bool = False
    frame: pd.DataFrame | None = None
    query: str = ""
    limit: int = SQL_LIMIT

    def json_rows(self) -> list[list]:
        return [[jsonable(v) for v in r] for r in self.rows]


class SqlError(DataError):
    def __init__(self, message: str, kind: str = "sql"):
        super().__init__(message)
        self.kind = kind


_FORBIDDEN = re.compile(
    r"\b(attach|detach|copy|install|load|pragma|create|insert|update|delete|drop|alter|export|"
    r"import|vacuum|checkpoint|truncate|call|read_csv\w*|read_parquet|read_json\w*|read_text|"
    r"read_blob|parquet_scan|csv_scan|sniff_csv|glob|getenv|current_setting)\b",
    re.I,
)
_COMMENTS = re.compile(r"--[^\n]*|/\*.*?\*/", re.S)
_QUOTED = re.compile(r"'(?:[^']|'')*'|\"(?:[^\"]|\"\")*\"")


def check_select(query: str, lang: str = "en") -> str:
    """The query cleaned for execution; SqlError unless it is one read-only SELECT / WITH."""
    q = _COMMENTS.sub(" ", query or "").strip()
    while q.endswith(";"):
        q = q[:-1].rstrip()
    if not q:
        raise SqlError(t("data_err_empty_query", lang), "empty")
    bare = _QUOTED.sub("''", q)
    if ";" in bare or not re.match(r"(?is)^\(*\s*(select|with|from|values)\b", bare):
        raise SqlError(t("data_err_not_select", lang), "not_select")
    if _FORBIDDEN.search(bare):
        raise SqlError(t("data_err_not_select", lang), "not_select")
    try:
        import duckdb

        stmts = duckdb.connect(":memory:").extract_statements(q)
        if len(stmts) != 1 or stmts[0].type != duckdb.StatementType.SELECT:
            raise SqlError(t("data_err_not_select", lang), "not_select")
    except SqlError:
        raise
    except Exception:  # the parser error surfaces (with a hint) when the query runs
        pass
    return q


def _connect(ds: Dataset):
    import duckdb

    con = duckdb.connect(":memory:")
    for tb in ds.tables:
        con.register(tb.name, tb.frame)
    con.execute(f"SET memory_limit='{SQL_MEMORY}'")
    con.execute("SET threads=2")
    con.execute("SET enable_external_access=false")
    con.execute("SET lock_configuration=true")
    return con


def _schema_hint(ds: Dataset, only: str = "") -> str:
    parts = [
        f"{tb.name}: " + ", ".join(c.name for c in tb.columns[:30])
        for tb in ds.tables
        if not only or tb.name == only
    ]
    return "; ".join(parts)[:900]


def _explain(e: Exception, ds: Dataset, lang: str) -> SqlError:
    msg = str(e)
    low = msg.lower()
    detail = re.sub(r"\s+", " ", msg.split("\n")[0])[:200]
    if "out of memory" in low or "memory limit" in low:
        return SqlError(t("data_err_memory", lang), "memory")
    if "referenced column" in low or "column" in low and "not found" in low:
        return SqlError(
            t("data_err_unknown_column", lang, detail=detail, columns=_schema_hint(ds)),
            "unknown_column",
        )
    if "table with name" in low or "catalog error" in low and "table" in low:
        return SqlError(
            t(
                "data_err_unknown_table",
                lang,
                detail=detail,
                tables=", ".join(x.name for x in ds.tables),
            ),  # fmt: skip
            "unknown_table",
        )
    return SqlError(t("data_err_sql", lang, detail=detail), "sql")


def run_sql(
    ds: Dataset,
    query: str,
    limit: int = SQL_LIMIT,
    *,
    lang: str = "en",
    timeout: float = SQL_TIMEOUT_S,
) -> Result:
    """Run one read-only SELECT / WITH over the dataset's tables with DuckDB (in memory, 512 MB,
    no file or network access, `timeout` seconds). At most `limit` rows and 40 columns come back."""
    q = check_select(query, lang)
    limit = max(1, int(limit))
    con = _connect(ds)
    box: dict = {}

    def work():
        try:
            box["df"] = con.execute(f"SELECT * FROM ({q}) AS _q LIMIT {limit + 1}").fetchdf()
        except Exception as e:  # reported below
            box["err"] = e

    th = threading.Thread(target=work, daemon=True)
    started = time.monotonic()
    th.start()
    th.join(timeout)
    if th.is_alive():
        with contextlib.suppress(Exception):
            con.interrupt()
        th.join(3)
        raise SqlError(t("data_err_timeout", lang, n=f"{timeout:g}"), "timeout")
    try:
        if "err" in box:
            raise _explain(box["err"], ds, lang)
    finally:
        with contextlib.suppress(Exception):
            con.close()
    df: pd.DataFrame = box["df"]
    more = len(df) > limit
    df = df.iloc[:limit]
    cut = df.shape[1] > MAX_SQL_COLS
    if cut:
        df = df.iloc[:, :MAX_SQL_COLS]
    df.columns = _dedupe_names([str(c) for c in df.columns])
    rows = [[_py(v) for v in r] for r in df.itertuples(index=False, name=None)]
    logger.info("data sql ok rows=%d in %.2fs", len(rows), time.monotonic() - started)
    return Result(list(df.columns), rows, more, cut, df, q, limit)


def _dedupe_names(names: list[str]) -> list[str]:
    used: set[str] = set()
    return [_unique(n or "col", used) for n in names]


def _width(s: str) -> int:
    w = 0
    for ch in s:
        if unicodedata.combining(ch) or unicodedata.category(ch) in ("Mn", "Me", "Cf"):
            continue
        w += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return w


def _pad(s: str, w: int, right: bool) -> str:
    gap = " " * max(0, w - _width(s))
    return gap + s if right else s + gap


def _clip(s: str, n: int = 28) -> str:
    s = s.replace("\n", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


def result_md(res: Result, lang: str = "th") -> str:
    """The result as a code-block table (25 rows at most) with a line for what is left out."""
    if not res.rows:
        return t("data_sql_empty", lang)
    shown = res.rows[:SHOW_ROWS]
    cells = [[_clip(fmt_value(v)) for v in r] for r in shown]
    heads = [_clip(c) for c in res.columns]
    right = [all(isinstance(r[i], int | float) and not isinstance(r[i], bool) or r[i] is None for r in shown) for i in range(len(heads))]  # fmt: skip
    widths = [max([_width(h), *[_width(r[i]) for r in cells]]) for i, h in enumerate(heads)]
    out = [
        " | ".join(_pad(h, widths[i], right[i]) for i, h in enumerate(heads)),
        "-+-".join("-" * w for w in widths),
        *[" | ".join(_pad(c, widths[i], right[i]) for i, c in enumerate(r)) for r in cells],
    ]
    text = "```\n" + "\n".join(out) + "\n```"
    left = len(res.rows) - len(shown)
    notes = []
    if left > 0:
        notes.append(t("data_sql_more_rows", lang, n=f"{left:,}"))
    if res.truncated:
        notes.append(t("data_sql_capped", lang, n=f"{res.limit:,}"))
    if res.cols_cut:
        notes.append(t("data_sql_cols_cut", lang, n=MAX_SQL_COLS))
    return text + ("\n" + " ".join(notes) if notes else "")


def result_xlsx(res: Result, sheet: str = "Result") -> bytes:
    rows = [dict(zip(res.columns, r, strict=True)) for r in res.rows]
    return bizfin.to_xlsx({sheet: rows or [dict.fromkeys(res.columns)]})


def result_csv(res: Result) -> bytes:
    return bizfin.to_csv([dict(zip(res.columns, r, strict=True)) for r in res.rows])


# ----------------------------------------------------------------------------- charts
KINDS = ("bar", "line", "pie", "scatter", "hist")
PALETTE = ["#0072B2", "#E69F00", "#009E73", "#D55E00", "#CC79A7", "#56B4E9", "#F0E442", "#7A7A7A"]
_SCRIPT_RANGES = {
    "zh": [(0x4E00, 0x9FFF), (0x3400, 0x4DBF)],
    "ja": [(0x3040, 0x30FF)],
    "ko": [(0xAC00, 0xD7AF), (0x1100, 0x11FF)],
    "hi": [(0x0900, 0x097F)],
}


def scripts_in(texts: Iterable[str]) -> list[str]:
    """The font scripts (zh ja ko hi) that the texts need besides Thai / Latin."""
    found: list[str] = []
    for s in texts:
        for ch in str(s):
            cp = ord(ch)
            for lang, rngs in _SCRIPT_RANGES.items():
                if lang not in found and any(a <= cp <= b for a, b in rngs):
                    found.append(lang)
    return found


def chart_font_paths(texts: Iterable[str], ui_lang: str = "th", configured: str = "") -> list[str]:
    """Font files for the chart in fallback order: Thai / Latin face, then the scripts used."""
    from .pdf_report import find_fonts

    paths: list[str] = []
    need = scripts_in(texts)
    if ui_lang in _SCRIPT_RANGES and ui_lang not in need:
        need.append(ui_lang)
    for lang in [None, *need]:
        try:
            fs = find_fonts(configured, lang)
        except Exception:
            continue
        if fs.regular and fs.regular not in paths:
            paths.append(fs.regular)
    return paths


def _lookup(frame: pd.DataFrame, name: str | None) -> str | None:
    if not name:
        return None
    for c in frame.columns:
        if str(c) == name:
            return c
    for c in frame.columns:
        if str(c).casefold() == str(name).strip().casefold():
            return c
    return None


def _is_numeric(s: pd.Series) -> bool:
    return pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s)


def _tick_fmt(v, _pos=None) -> str:
    if abs(v) >= 1000 or float(v).is_integer():
        return f"{v:,.0f}"
    return f"{v:,.2f}"


def cap_categories(df: pd.DataFrame, x, ys: list, other: str) -> pd.DataFrame:
    """At most MAX_CATEGORIES rows: the top 29 by the first value column plus one summed `other`
    row. Shorter frames come back unchanged."""
    if len(df) <= MAX_CATEGORIES:
        return df
    top = df.sort_values(ys[0], ascending=False)
    rest = top.iloc[MAX_CATEGORIES - 1 :]
    row = {c: rest[c].sum() for c in ys}
    row[x] = other
    return pd.concat([top.iloc[: MAX_CATEGORIES - 1], pd.DataFrame([row])], ignore_index=True)


def chart_png(
    data: Result | pd.DataFrame,
    kind: str = "bar",
    x: str | None = None,
    y: list[str] | str | None = None,
    title: str = "",
    lang: str = "th",
    *,
    font_paths: list[str] | None = None,
    configured_font: str = "",
) -> bytes:
    """A 1200x700 PNG chart of a query result (or frame): bar, line, pie, scatter or hist. Thai /
    CJK / Devanagari text uses the fonts of `pdf_report.find_fonts`; at most 30 categories (the
    rest summed into "other"); numbers with thousands separators. Raises DataError."""
    from matplotlib import font_manager
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    from matplotlib.ticker import FuncFormatter

    kind = (kind or "bar").strip().lower()
    if kind not in KINDS:
        raise DataError(t("data_err_chart_kind", lang))
    df = data.frame if isinstance(data, Result) else data
    if df is None or df.empty or df.shape[1] == 0:
        raise DataError(t("data_err_chart_empty", lang))
    df = df.copy()
    cols = ", ".join(str(c) for c in df.columns)
    ys_in = [y] if isinstance(y, str) else list(y or [])
    ys = []
    for name in ys_in:
        c = _lookup(df, str(name).strip())
        if c is None:
            raise DataError(t("data_err_chart_col", lang, column=name, columns=cols))
        ys.append(c)
    xc = _lookup(df, x) if x else None
    if x and xc is None:
        raise DataError(t("data_err_chart_col", lang, column=x, columns=cols))
    if xc is None and kind != "hist":
        xc = df.columns[0]
    if not ys:
        ys = [c for c in df.columns if c != xc and _is_numeric(df[c])][:6]
        if not ys and kind == "hist" and xc is not None:
            ys = [xc]
        if not ys:
            raise DataError(t("data_err_chart_numeric", lang, column=cols))
    if kind == "hist":
        ys = ys[:1]
        xc = None
    if kind == "pie":
        ys = ys[:1]
    for c in ys:
        if not _is_numeric(df[c]):
            coerced = pd.to_numeric(df[c], errors="coerce")
            if coerced.notna().sum() == 0:
                raise DataError(t("data_err_chart_numeric", lang, column=str(c)))
            df[c] = coerced
    if kind == "scatter" and (xc is None or not _is_numeric(df[xc])):
        coerced = pd.to_numeric(df[xc], errors="coerce") if xc is not None else None
        if coerced is None or coerced.notna().sum() == 0:
            raise DataError(t("data_err_chart_numeric", lang, column=str(xc)))
        df[xc] = coerced

    categorical = kind in ("bar", "pie") or (
        kind == "line"
        and xc is not None
        and not (_is_numeric(df[xc]) or pd.api.types.is_datetime64_any_dtype(df[xc]))  # fmt: skip
    )
    if kind in ("bar", "pie") and xc is not None:
        df = cap_categories(df, xc, ys, t("data_other", lang))
    if kind == "line" and len(df) > 2000:
        df = df.iloc[:: math.ceil(len(df) / 2000)]

    texts = [title, *[str(c) for c in ys]]
    if xc is not None:
        texts += [str(xc)] + [str(v) for v in df[xc].head(60)]
    paths = font_paths if font_paths is not None else chart_font_paths(texts, lang, configured_font)
    families = []
    for p in paths:
        try:
            font_manager.fontManager.addfont(p)
            families.append(font_manager.FontProperties(fname=p).get_name())
        except Exception:
            logger.warning("chart font %s not usable", p)
    families.append("DejaVu Sans")

    fig = Figure(figsize=(CHART_W / 100, CHART_H / 100), dpi=100, facecolor="white")
    FigureCanvasAgg(fig)
    fmt = FuncFormatter(_tick_fmt)
    with _rc(families):
        ax = fig.add_subplot(111)
        if title:
            ax.set_title(title, fontsize=17, pad=14)
        if kind == "bar":
            labels = [_clip(str(v), 24) for v in df[xc]]
            n = len(ys)
            width = 0.8 / n
            pos = list(range(len(labels)))
            for k, c in enumerate(ys):
                xs = [p + (k - (n - 1) / 2) * width for p in pos]
                bars = ax.bar(xs, df[c].fillna(0), width * 0.95, label=str(c), color=PALETTE[k % 8])
                if len(labels) <= 12 and n == 1:
                    ax.bar_label(bars, labels=[_tick_fmt(v) for v in df[c].fillna(0)], fontsize=10, padding=3)  # fmt: skip
            ax.set_xticks(pos, labels, rotation=0 if len(labels) <= 8 else 45, ha="center" if len(labels) <= 8 else "right")  # fmt: skip
            ax.yaxis.set_major_formatter(fmt)
            ax.set_xlabel(str(xc))
        elif kind == "line":
            if categorical:
                xs_ = list(range(len(df)))
                ax.set_xticks(xs_[:: max(1, len(xs_) // 12)], [_clip(str(v), 18) for v in df[xc]][:: max(1, len(xs_) // 12)], rotation=45, ha="right")  # fmt: skip
            else:
                xs_ = df[xc]
            for k, c in enumerate(ys):
                ax.plot(xs_, df[c], marker="o" if len(df) <= 40 else None, lw=2, label=str(c), color=PALETTE[k % 8])  # fmt: skip
            ax.yaxis.set_major_formatter(fmt)
            ax.set_xlabel(str(xc))
            if pd.api.types.is_datetime64_any_dtype(df[xc]):
                fig.autofmt_xdate()
        elif kind == "pie":
            vals = df[ys[0]].fillna(0)
            if (vals < 0).any():
                raise DataError(t("data_err_chart_pie", lang))
            if vals.sum() <= 0:
                raise DataError(t("data_err_chart_empty", lang))
            labels = [_clip(str(v), 24) for v in df[xc]]
            ax.pie(
                vals, labels=None, colors=[PALETTE[i % 8] for i in range(len(vals))], startangle=90,
                counterclock=False, autopct=lambda p: f"{p:.1f}%" if p >= 3 else "",
                wedgeprops={"linewidth": 1, "edgecolor": "white"}, pctdistance=0.8,
                textprops={"fontsize": 11, "color": "white"},
            )  # fmt: skip
            ax.legend(labels, loc="center left", bbox_to_anchor=(1.0, 0.5), fontsize=11, frameon=False, ncol=1 if len(labels) <= 15 else 2)  # fmt: skip
            ax.set_aspect("equal")
        elif kind == "scatter":
            for k, c in enumerate(ys):
                ax.scatter(df[xc], df[c], s=36, alpha=0.8, label=str(c), color=PALETTE[k % 8])
            ax.xaxis.set_major_formatter(fmt)
            ax.yaxis.set_major_formatter(fmt)
            ax.set_xlabel(str(xc))
        else:  # hist
            vals = df[ys[0]].dropna()
            if vals.empty:
                raise DataError(t("data_err_chart_empty", lang))
            ax.hist(vals, bins=min(20, max(5, int(math.sqrt(len(vals))))), color=PALETTE[0], edgecolor="white")  # fmt: skip
            ax.xaxis.set_major_formatter(fmt)
            ax.set_xlabel(str(ys[0]))
        if kind != "pie":
            ax.grid(axis="y", alpha=0.3)
            ax.spines[["top", "right"]].set_visible(False)
            if len(ys) > 1 and kind != "hist":
                ax.legend(frameon=False)
            elif len(ys) == 1 and kind in ("bar", "line", "scatter"):
                ax.set_ylabel(str(ys[0]))
        fig.tight_layout()
        return _png(fig)


def _rc(families: list[str]):
    import matplotlib

    return matplotlib.rc_context({"font.family": families, "axes.unicode_minus": False})


def _png(fig) -> bytes:
    for dpi in (100, 80, 60):
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=dpi)
        if buf.tell() <= CHART_MAX_BYTES:
            return buf.getvalue()
    raise DataError("the chart image is too large")


# ----------------------------------------------------------------------------- stored datasets
def _dump_table(tb: Table) -> bytes:
    buf = io.StringIO()
    tb.frame.to_csv(buf, index=False, na_rep="\\N")
    return gzip.compress(buf.getvalue().encode("utf-8"), 6)


def _load_table(meta: dict, raw: bytes) -> Table:
    text = gzip.decompress(raw).decode("utf-8")
    df = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False, na_values=["\\N"])
    cols = [Column(**c) for c in meta["columns"]]
    out = {}
    for c in cols:
        s = df[c.name] if c.name in df else pd.Series([None] * len(df), dtype="object")
        if c.type == "integer":
            out[c.name] = pd.to_numeric(s, errors="coerce").astype("Int64")
        elif c.type == "number":
            out[c.name] = pd.to_numeric(s, errors="coerce").astype("float64")
        elif c.type == "bool":
            out[c.name] = s.map({"True": True, "False": False}).astype("boolean")
        elif c.type in ("date", "datetime"):
            out[c.name] = pd.to_datetime(s, errors="coerce").astype("datetime64[us]")
        else:
            out[c.name] = s.astype("object").where(s.notna(), None)
    frame = pd.DataFrame(out, index=pd.RangeIndex(len(df)))
    return Table(meta["name"], meta["label"], frame, cols, truncated=meta.get("truncated", 0))


class DatasetRepo:
    """Saved datasets per guild on the DocStore (SQLite or Firestore): one record per dataset
    (`dory_datasets`, key `ds-xxxxxx`) and the gzip CSV of each table in base64 chunks of at most
    600 kB each (`dory_dataset_data`, key `ds-xxxxxx_<table>_<n>`; a Firestore document holds
    1 MiB)."""

    def __init__(self, docs: DocStore, clock: Callable[[], datetime] | None = None):
        self.docs = docs
        self._clock = clock or (lambda: datetime.now(UTC))

    def save(
        self,
        guild_id: int,
        ds: Dataset,
        *,
        name: str,
        owner_id: int,
        owner_name: str = "",
        channel_id: int = 0,
        prof: dict | None = None,
        did: str | None = None,
    ) -> dict:
        """Store `ds`; returns the record. DataError when it is over 3 MB compressed."""
        blobs = [(tb, _dump_table(tb)) for tb in ds.tables]
        if sum(len(b) for _tb, b in blobs) > MAX_STORED_BYTES:
            raise DataError("the dataset is over 3 MB compressed, so it cannot be kept")
        did = did or "ds-" + secrets.token_hex(3)
        while self.docs.get(DATASETS, guild_id, did) is not None:
            did = "ds-" + secrets.token_hex(3)
        metas = []
        for tb, blob in blobs:
            b64 = base64.b64encode(blob).decode("ascii")
            parts = [b64[i : i + CHUNK_CHARS] for i in range(0, len(b64), CHUNK_CHARS)] or [""]
            for n, part in enumerate(parts):
                self.docs.put(DATASET_DATA, guild_id, f"{did}_{tb.name}_{n}", {"b64": part})
            metas.append(
                {
                    "name": tb.name,
                    "label": tb.label,
                    "rows": tb.rows,
                    "truncated": tb.truncated,
                    "chunks": len(parts),
                    "columns": [c.dump() for c in tb.columns],
                }
            )
        rec = {
            "id": did,
            "name": (name or ds.filename or did).strip()[:80],
            "filename": ds.filename,
            "owner_id": int(owner_id),
            "owner_name": owner_name,
            "channel_id": int(channel_id),
            "created_at": self._clock().isoformat(timespec="seconds"),
            "pinned": False,
            "tables": metas,
            "notes": ds.notes,
            "profile": _clean_json(prof or profile(ds, "en")),
        }
        self.docs.put(DATASETS, guild_id, did, rec)
        self._prune(guild_id, keep=did)
        return rec

    def get(self, guild_id: int, did: str) -> dict | None:
        return self.docs.get(DATASETS, guild_id, did)

    def all(self, guild_id: int) -> list[dict]:
        return sorted(self.docs.find(DATASETS, guild_id), key=lambda r: r.get("created_at", ""), reverse=True)  # fmt: skip

    def load(self, guild_id: int, rec: dict) -> Dataset:
        tables = []
        for meta in rec["tables"]:
            b64 = "".join(
                (
                    self.docs.get(DATASET_DATA, guild_id, f"{rec['id']}_{meta['name']}_{n}") or {}
                ).get("b64", "")  # fmt: skip
                for n in range(meta["chunks"])
            )
            if not b64:
                raise DataError("the saved data is missing")
            tables.append(_load_table(meta, base64.b64decode(b64)))
        return Dataset(rec.get("filename", ""), tables, list(rec.get("notes") or []))

    def delete(self, guild_id: int, rec: dict) -> None:
        for meta in rec["tables"]:
            for n in range(meta["chunks"]):
                self.docs.delete(DATASET_DATA, guild_id, f"{rec['id']}_{meta['name']}_{n}")
        self.docs.delete(DATASETS, guild_id, rec["id"])

    def _prune(self, guild_id: int, keep: str) -> None:
        recs = self.all(guild_id)
        extra = len(recs) - MAX_DATASETS
        if extra <= 0:
            return
        old = [r for r in reversed(recs) if not r.get("pinned") and r["id"] != keep]
        for r in old[:extra]:
            self.delete(guild_id, r)

    def find(self, guild_id: int, text: str) -> list[dict]:
        """Datasets whose id, name or file name contains `text` (case-insensitive), newest first."""
        low = (text or "").strip().casefold()
        return [
            r
            for r in self.all(guild_id)
            if low
            and (
                low == r["id"]
                or low in r["name"].casefold()
                or low in r.get("filename", "").casefold()
            )  # fmt: skip
        ]


def _clean_json(obj: Any) -> Any:
    return json.loads(json.dumps(obj, ensure_ascii=False, default=jsonable))


def can_read(
    rec: dict, user_id: int, channel_readable: Callable[[int], bool], private: bool = False
) -> bool:
    """May `user_id` open the dataset in this conversation? A dataset that came from a channel:
    only when the asker AND everybody who will read the answer can read that channel (the owner
    too, so a public room never shows a hidden channel's data). A DM / private-chat dataset
    (channel 0): its owner, and only in a private conversation (`private`)."""
    cid = int(rec.get("channel_id") or 0)
    if cid == 0:
        return rec.get("owner_id") == user_id and private
    return bool(channel_readable(cid))
