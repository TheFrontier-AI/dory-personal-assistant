"""Documents (PDF, docx, xlsx, pptx, text) posted in remembered channels: read the text for free,
fall back to the cheap multimodal model for scanned PDFs, and keep a short searchable note in the
memory log. Blocking functions (run with asyncio.to_thread); bytes are never stored.
"""

from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass

from google.genai import types
from pydantic import BaseModel

from .llm import generate

logger = logging.getLogger(__name__)

MAX_DOC_BYTES = 20 * 1024 * 1024
OCR_MAX_PAGES = 30
MIN_TEXT_CHARS = 20  # a text layer shorter than this (in total) means a scanned PDF
KEY_TEXT_CHARS = 1500
NOTE_CHARS = 2600
PROMPT_TEXT_CHARS = 12_000
OPEN_FILE_CHARS = 30_000
_MAX_CELLS = 20_000
_EXT_KIND = {
    "pdf": "pdf", "docx": "docx", "xlsx": "xlsx", "pptx": "pptx", "doc": "doc", "xls": "xls",
    "txt": "text", "csv": "text", "md": "text", "tsv": "text", "json": "text", "log": "text",
}  # fmt: skip
_LABEL = {
    "pdf": "PDF", "docx": "DOCX", "xlsx": "XLSX", "pptx": "PPTX", "doc": "DOC", "xls": "XLS",
    "text": "ข้อความ",
}  # fmt: skip
_UNIT = {"pdf": "หน้า", "docx": "หน้า", "doc": "หน้า", "xlsx": "ชีต", "xls": "ชีต", "pptx": "สไลด์"}
MAX_ROWS = 500

PROMPT = (
    "You read one document posted in a team's Discord chat so it can be searched later. Reply as "
    "JSON with: doc_type (e.g. invoice, receipt, quote, contract, report, letter, spreadsheet, "
    "slides, notes, other), summary (1-3 lines, in the document's language), parties (people / "
    "companies named, empty if none), amounts (every key amount WITH currency, e.g. 'total "
    "1,250.50 THB'), dates (key dates, ISO when clear), references (invoice / PO / reference "
    "numbers), amount_value (the main total as a plain number, no currency; null if none), "
    "addressed_to (people - not companies - the document is addressed to or names as the "
    "recipient / responsible person, comma separated, empty if none) and text ({text_rule}). For "
    "doc_type also use legal (court / lawyer / legal notice), official (government / "
    "authority letter) or company (the company's own records: registration certificate, "
    "articles of association / memorandum, company regulations, shareholder register or "
    "agreement, licences) when that fits. Use only what is in the document; never guess unreadable "
    "parts."
)
TEXT_RULE_LAYER = "leave empty"
TEXT_RULE_OCR = "the document's important text transcribed verbatim, at most about 1500 characters"


class DocNote(BaseModel):
    doc_type: str = "other"
    summary: str = ""
    parties: str = ""
    amounts: str = ""
    dates: str = ""
    references: str = ""
    text: str = ""
    amount_value: float | None = None  # the main total as a number (importance threshold)
    addressed_to: str = ""  # people it is addressed to / names (who it concerns)


class DocUnreadable(ValueError):
    """The document cannot be turned into text (old .doc without antiword, .xls, corrupt)."""


@dataclass
class DocText:
    text: str
    pages: int  # pages / sheets / slides (see _UNIT)


def doc_kind(att_or_name, content_type: str | None = None) -> str | None:
    """pdf / docx / xlsx / pptx / text for an attachment (or file name), else None."""
    name = att_or_name if isinstance(att_or_name, str) else getattr(att_or_name, "filename", "")
    ct = content_type
    if not isinstance(att_or_name, str):
        ct = str(getattr(att_or_name, "content_type", "") or "")
    name = str(name or "")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in _EXT_KIND:
        return _EXT_KIND[ext]
    ct = (ct or "").split(";")[0].strip().lower()
    if ct == "application/pdf":
        return "pdf"
    if ct in ("text/plain", "text/csv"):
        return "text"
    return None


def kind_label(kind: str) -> str:
    return _LABEL.get(kind, kind)


# ------------------------------------------------------------------ free text extraction
def extract_pdf(data: bytes) -> DocText:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    out, size = [], 0
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:  # one bad page must not lose the rest
            t = ""
        out.append(t)
        size += len(t)
        if size > OPEN_FILE_CHARS * 2:
            break
    return DocText(_clean("\n".join(out)), pages)


def extract_docx(data: bytes) -> DocText:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = []
    for sec in doc.sections:
        parts += [p.text for p in sec.header.paragraphs if p.text.strip()]
    parts += [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    for sec in doc.sections:
        parts += [p.text for p in sec.footer.paragraphs if p.text.strip()]
    text = _clean("\n".join(parts))
    return DocText(text, max(1, len(text) // 3000))


def extract_xlsx(data: bytes) -> DocText:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, sheets = [], len(wb.worksheets)
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= MAX_ROWS:
                parts.append(f"[... แสดง {MAX_ROWS} แถวแรกของชีตนี้ / first {MAX_ROWS} rows only]")
                break
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                parts.append(",".join(vals).rstrip(","))
    wb.close()
    return DocText(_clean("\n".join(parts)), max(1, sheets))


def extract_pptx(data: bytes) -> DocText:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for n, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""
        parts.append(f"## สไลด์ {n}: {title}".rstrip(": "))
        for shape in slide.shapes:
            if shape is slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[โน้ต] {notes}")
    return DocText(_clean("\n".join(parts)), max(1, len(prs.slides)))


def extract_doc(data: bytes) -> DocText:
    """Old Word (.doc) through `antiword` when installed."""
    import shutil
    import subprocess

    if shutil.which("antiword") is None:
        raise DocUnreadable("old .doc needs antiword; save the file as .docx or PDF")
    proc = subprocess.run(["antiword", "-"], input=data, capture_output=True, timeout=60)
    if proc.returncode != 0:
        raise DocUnreadable("antiword could not read the file; save it as .docx or PDF")
    text = _clean(proc.stdout.decode("utf-8", "replace"))
    return DocText(text, max(1, len(text) // 3000))


def extract_xls(data: bytes) -> DocText:
    raise DocUnreadable("old .xls cannot be read; save the file as .xlsx or CSV")


def extract_text_file(data: bytes) -> DocText:
    text = _clean(data.decode("utf-8", errors="replace"))
    return DocText(text, max(1, len(text) // 3000))


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text or "")
    return re.sub(r"\n\s*\n+", "\n", text).strip()


_EXTRACT = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "doc": extract_doc,
    "xls": extract_xls,
    "xlsx": extract_xlsx,
    "pptx": extract_pptx,
    "text": extract_text_file,
}


def extract_text(data: bytes, kind: str) -> DocText:
    """Text layer of a document for free. Raises on a corrupt / encrypted file."""
    return _EXTRACT[kind](data)


def extract_document(
    data: bytes, filename: str = "", mime: str | None = None
) -> tuple[str, str, int]:
    """(kind, text, pages) for any supported document: the one reader used by the addressed
    path, the memory notes, the backlog and open_file. Raises DocUnreadable / other errors."""
    kind = doc_kind(filename, mime)
    if kind is None:
        raise DocUnreadable("unsupported document type")
    dt = extract_text(data, kind)
    return kind, dt.text, dt.pages


def clip_for_agent(text: str, max_chars: int = OPEN_FILE_CHARS) -> str:
    """Text for the agent, with a visible note when it is cut."""
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + f"\n[ตัดทอน: แสดง {max_chars} ตัวอักษรแรกจาก {len(text)} / truncated]"


def unit_label(kind: str) -> str:
    return _UNIT.get(kind, "หน้า")


def is_scanned(kind: str, dt: DocText) -> bool:
    return kind == "pdf" and len(dt.text) < MIN_TEXT_CHARS


def first_pages(data: bytes, n: int = OCR_MAX_PAGES) -> bytes:
    """The PDF limited to its first `n` pages (the bytes unchanged when it is not longer)."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) <= n:
        return data
    w = PdfWriter()
    for page in reader.pages[:n]:
        w.add_page(page)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


# ------------------------------------------------------------------ model
def analyze(client, model: str, data: bytes, kind: str, dt: DocText) -> DocNote:
    """One cheap call: summary and key fields. A scanned PDF is sent as a PDF part (first 30
    pages, <= 20 MB) and also OCR'd; everything else sends only its text layer."""
    scanned = is_scanned(kind, dt)
    if scanned:
        if len(data) > MAX_DOC_BYTES:
            raise ValueError("scanned PDF too large")
        contents = [
            types.Part.from_bytes(data=first_pages(data), mime_type="application/pdf"),
            PROMPT.format(text_rule=TEXT_RULE_OCR),
        ]
    else:
        contents = [
            f"{PROMPT.format(text_rule=TEXT_RULE_LAYER)}\n\nDocument text:\n{dt.text[:PROMPT_TEXT_CHARS]}"
        ]
    resp = generate(
        client,
        model,
        contents,
        types.GenerateContentConfig(
            temperature=0,
            max_output_tokens=900,
            response_mime_type="application/json",
            response_schema=DocNote,
        ),
    )
    return DocNote.model_validate(json.loads(resp.text))


def ocr_text(client, model: str, data: bytes) -> str:
    """Plain text of a scanned PDF (first 30 pages) for the open_file tool."""
    resp = generate(
        client,
        model,
        [
            types.Part.from_bytes(data=first_pages(data), mime_type="application/pdf"),
            "Transcribe all legible text of this document verbatim, in reading order. "
            "Do not summarise and do not guess unreadable parts.",
        ],
        types.GenerateContentConfig(temperature=0, max_output_tokens=16000),
    )
    return (getattr(resp, "text", None) or "").strip()


# ------------------------------------------------------------------ memory note
def doc_note(
    name: str, kind: str, dt: DocText, note: DocNote | None, max_chars: int = NOTE_CHARS
) -> str:
    """ "[PDF: name, N หน้า, ประเภท] summary | ข้อความสำคัญ: first ~1500 chars of the text"."""
    pages = f", {dt.pages} {unit_label(kind)}" if kind != "text" else ""
    dtype = f", {note.doc_type.strip()[:30]}" if note and note.doc_type.strip() else ""
    head = f"[{kind_label(kind)}: {name}{pages}{dtype}]"
    bits = []
    if note:
        bits.append(" ".join(note.summary.split()))
        for label, v in (
            ("คู่ค้า/ผู้เกี่ยวข้อง", note.parties),
            ("ยอด", note.amounts),
            ("วันที่", note.dates),
            ("เลขที่", note.references),
        ):
            if v.strip():
                bits.append(f"{label}: {' '.join(v.split())}")
    body = " | ".join(b for b in bits if b)
    text = " ".join((dt.text or (note.text if note else "")).split())[:KEY_TEXT_CHARS]
    out = head + (f" {body}" if body else "")
    if text:
        out += f" | ข้อความสำคัญ: {text}"
    return out[:max_chars]


def receipt_doc_note(name: str, kind: str, dt: DocText, receipt_line: str) -> str:
    """The receipt pipeline's note (receipt_note) plus the text read for free; no model call."""
    pages = f", {dt.pages} {unit_label(kind)}" if kind != "text" else ""
    text = " ".join(dt.text.split())[:KEY_TEXT_CHARS]
    out = f"[{kind_label(kind)}: {name}{pages}] {receipt_line}"
    if text:
        out += f" | ข้อความสำคัญ: {text}"
    return out[:NOTE_CHARS]
