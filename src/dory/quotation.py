"""Quotations (ใบเสนอราคา): made from natural language, as a PDF and an editable .docx.

A quotation is a spec (customer, lines, discount, VAT, terms) kept per guild in the DocStore kind
`dory_quotes` (SQLite or Firestore) with its whole revision history. The files are NOT stored:
they are rendered from the spec of a revision on demand. Numbers are `QT-YYYY-NNNN` per guild and
never reused; a revision keeps the base number and adds `-r2`, `-r3`...

All money is `Decimal` (`bizfin` helpers); everything the model passes is validated here.

Who may do what (decided in code, never by the model):
* issue / revise: executives, or members with a role listed in the guild's `quote_roles`;
* read / list / re-send: the author, executives, and anyone who can read the channel where the
  quotation was made (same idea as the idea documents).

The company header (name, address, tax id, bank account, signer) comes from the company playbook
profile (`company_brain`); missing fields print a placeholder and are reported to the caller.
"""

from __future__ import annotations

import io
import logging
import re
import threading
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

from fpdf import FPDF
from pydantic import BaseModel, Field

from .bizfin import VAT_RATE, D, FinError, q2
from .company_brain import CompanyRepo
from .coo_store import QUOTES, DocStore
from .i18n import t
from .ideas_pdf import _DocPdf
from .pdf_report import FontSet

logger = logging.getLogger(__name__)

MAX_LINES = 50
MAX_CHARGES = 10
MAX_REVISIONS = 30
MAX_LIST = 30
MAX_TEXT = 300
MAX_NOTES = 1500
DEFAULT_VALID_DAYS = 30
VAT_MODES = ("excluded", "included", "none")
NUMBER_RE = re.compile(r"^\s*QT[-\s]?(\d{4})[-\s]?(\d{1,6})(?:[-\s]?r(\d{1,3}))?\s*$", re.I)
REQUIRED_PROFILE = ("legal_name", "address", "tax_id", "signer_name")
HUNDRED = Decimal(100)
MAX_AMOUNT = Decimal("1000000000000")  # a trillion: far beyond any quotation, keeps text sane


# ------------------------------------------------------------------ text helpers
def _txt(v: Any, limit: int = MAX_TEXT) -> str:
    s = " ".join(str(v or "").replace("\x00", " ").split())
    return s if len(s) <= limit else s[: limit - 1].rstrip() + "…"


_DIGITS = ["ศูนย์", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]
_PLACES = ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน"]


def _thai_int(n: int) -> str:
    """A whole number in Thai words (below a million per group; millions recurse)."""
    if n == 0:
        return ""
    if n >= 1_000_000:
        high, low = divmod(n, 1_000_000)
        return _thai_int(high) + "ล้าน" + ("เอ็ด" if low == 1 else _thai_int(low))
    s = str(n)
    out = []
    for i, ch in enumerate(s):
        d, pos = int(ch), len(s) - i - 1
        if d == 0:
            continue
        if pos == 1 and d == 2:
            out.append("ยี่สิบ")
        elif pos == 1 and d == 1:
            out.append("สิบ")
        elif pos == 0 and d == 1 and len(s) > 1:
            out.append("เอ็ด")  # 11 = สิบเอ็ด, 101 = หนึ่งร้อยเอ็ด
        else:
            out.append(_DIGITS[d] + _PLACES[pos])
    return "".join(out)


def thai_baht_text(amount: Decimal) -> str:
    """ "สามหมื่นห้าพันห้าร้อยเจ็ดสิบเจ็ดบาทห้าสิบสตางค์" for 35,577.50; "...บาทถ้วน" without satang."""
    amount = q2(abs(D(amount)))
    baht, satang = divmod(int(amount * 100), 100)
    words = (_thai_int(baht) or _DIGITS[0]) + "บาท"
    return words + (_thai_int(satang) + "สตางค์" if satang else "ถ้วน")


def fmt_money(d: Decimal) -> str:
    return f"{q2(d):,.2f}"


def fmt_qty(d: Decimal) -> str:
    s = f"{d.normalize():,f}"
    return s.rstrip("0").rstrip(".") if "." in s else s


def fmt_pct(d: Decimal) -> str:
    s = f"{d.normalize():f}"
    return s.rstrip("0").rstrip(".") if "." in s else s


def fmt_date(d: date) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year}"


# ------------------------------------------------------------------ model
class Line(BaseModel):
    description: str
    qty: Decimal
    unit: str = ""
    unit_price: Decimal
    discount_percent: Decimal = Decimal(0)
    discount_amount: Decimal = Decimal(0)


class Charge(BaseModel):
    description: str
    amount: Decimal


class Customer(BaseModel):
    name: str
    address: str = ""
    contact: str = ""
    tax_id: str = ""


class QuoteSpec(BaseModel):
    customer: Customer
    lines: list[Line]
    charges: list[Charge] = Field(default_factory=list)
    discount_percent: Decimal = Decimal(0)  # on the total, after the line discounts
    discount_amount: Decimal = Decimal(0)
    currency: str = "THB"
    vat: str = "excluded"  # excluded: add VAT on top | included: prices include VAT | none
    vat_rate: Decimal = VAT_RATE * 100
    valid_days: int = DEFAULT_VALID_DAYS
    payment_terms: str = ""
    notes: str = ""


class Revision(BaseModel):
    rev: int
    spec: QuoteSpec
    total: str
    issued: str  # ISO date (guild time zone)
    by_id: int
    by_name: str = ""
    at: str = ""
    summary: str = ""


class QuoteDoc(BaseModel):
    number: str  # base number QT-2026-0001
    year: int
    seq: int
    channel_id: int = 0
    author_id: int
    author_name: str = ""
    created_at: str = ""
    updated_at: str = ""
    revisions: list[Revision] = Field(default_factory=list)

    @property
    def latest(self) -> Revision:
        return self.revisions[-1]

    def label(self, rev: int | None = None) -> str:
        n = rev or self.latest.rev
        return self.number if n <= 1 else f"{self.number}-r{n}"


def parse_number(text: str) -> tuple[str, int | None] | None:
    """("QT-2026-0001", revision or None) from "QT-2026-0001", "qt 2026 1 r2"..., else None."""
    m = NUMBER_RE.match(text or "")
    if not m:
        return None
    return f"QT-{int(m.group(1))}-{int(m.group(2)):04d}", (int(m.group(3)) if m.group(3) else None)


# ------------------------------------------------------------------ validation
def _num(v: Any, name: str, *, lo: Decimal = Decimal(0), hi: Decimal = MAX_AMOUNT) -> Decimal:
    d = D(v, name)
    if d < lo or d > hi:
        raise FinError(f"{name} is out of range")
    return d


def _pct(v: Any, name: str) -> Decimal:
    return _num(0 if v in (None, "") else v, name, hi=HUNDRED)


def build_line(raw: dict, n: int) -> Line:
    if not isinstance(raw, dict):
        raise FinError(f"line {n}: must be an object")
    desc = _txt(raw.get("description") or raw.get("name") or raw.get("item"))
    if not desc:
        raise FinError(f"line {n}: description is required")
    qty = D(raw.get("qty", raw.get("quantity", 1)), f"line {n} qty")
    if qty <= 0 or qty > MAX_AMOUNT:
        raise FinError(f"line {n}: qty must be greater than 0")
    price = _num(raw.get("unit_price", raw.get("price")), f"line {n} unit_price")
    dp = _pct(raw.get("discount_percent"), f"line {n} discount_percent")
    da = _num(raw.get("discount_amount") or 0, f"line {n} discount_amount")
    if dp and da:
        raise FinError(f"line {n}: give a discount percent or an amount, not both")
    if da > q2(qty * price):
        raise FinError(f"line {n}: the discount is more than the line amount")
    return Line(
        description=desc,
        qty=qty,
        unit=_txt(raw.get("unit"), 30),
        unit_price=price,
        discount_percent=dp,
        discount_amount=da,
    )


def build_charge(raw: dict, n: int) -> Charge:
    if not isinstance(raw, dict):
        raise FinError(f"charge {n}: must be an object")
    desc = _txt(raw.get("description") or raw.get("name"))
    if not desc:
        raise FinError(f"charge {n}: description is required")
    return Charge(description=desc, amount=_num(raw.get("amount"), f"charge {n} amount"))


def build_lines(raw: Any) -> list[Line]:
    if not isinstance(raw, list) or not raw:
        raise FinError("at least one item is required")
    if len(raw) > MAX_LINES:
        raise FinError(f"at most {MAX_LINES} lines")
    return [build_line(r, i + 1) for i, r in enumerate(raw)]


def build_charges(raw: Any) -> list[Charge]:
    if not raw:
        return []
    if not isinstance(raw, list) or len(raw) > MAX_CHARGES:
        raise FinError(f"at most {MAX_CHARGES} extra charges")
    return [build_charge(r, i + 1) for i, r in enumerate(raw)]


def clean_currency(v: Any) -> str:
    cur = _txt(v or "THB", 3).upper()
    if not re.fullmatch(r"[A-Z]{3}", cur):
        raise FinError("currency must be a 3-letter code such as THB")
    return cur


def clean_vat(v: Any) -> str:
    mode = str(v or "excluded").strip().casefold()
    mode = {
        "exclusive": "excluded",
        "exclude": "excluded",
        "add": "excluded",
        "inclusive": "included",
        "include": "included",
        "no": "none",
        "without": "none",
        "false": "none",
    }.get(mode, mode)
    if mode not in VAT_MODES:
        raise FinError("vat must be excluded, included or none")
    return mode


def check_spec(spec: QuoteSpec) -> QuoteSpec:
    """Final limits of a whole spec (also after a revision)."""
    if not 1 <= len(spec.lines) <= MAX_LINES:
        raise FinError(f"a quotation needs 1 to {MAX_LINES} lines")
    if len(spec.charges) > MAX_CHARGES:
        raise FinError(f"at most {MAX_CHARGES} extra charges")
    if not spec.customer.name:
        raise FinError("the customer name is required")
    if not 1 <= spec.valid_days <= 365:
        raise FinError("valid_days must be between 1 and 365")
    if not 0 <= spec.vat_rate <= 30:
        raise FinError("vat_rate_percent must be between 0 and 30")
    if spec.discount_percent and spec.discount_amount:
        raise FinError("give the total discount as a percent or an amount, not both")
    totals = compute(spec)  # also rejects a discount bigger than the amount
    if totals.total > MAX_AMOUNT:
        raise FinError("the total is too large")
    return spec


# ------------------------------------------------------------------ money
@dataclass
class Totals:
    lines: list[tuple[Decimal, Decimal, Decimal]]  # (gross, discount, net) per line
    subtotal: Decimal  # sum of gross line amounts
    line_discounts: Decimal
    total_discount: Decimal
    discount: Decimal  # all discounts
    charges: Decimal
    base: Decimal  # after discounts, with charges (VAT base when VAT is on top)
    vat: Decimal
    total: Decimal
    vat_mode: str
    vat_rate: Decimal


def compute(spec: QuoteSpec) -> Totals:
    """Totals in Decimal, every amount rounded half-up to satang as it is printed."""
    rows: list[tuple[Decimal, Decimal, Decimal]] = []
    for ln in spec.lines:
        gross = q2(ln.qty * ln.unit_price)
        disc = (
            q2(gross * ln.discount_percent / HUNDRED) if ln.discount_percent else ln.discount_amount
        )
        disc = q2(disc)
        if disc > gross:
            raise FinError("a line discount is more than the line amount")
        rows.append((gross, disc, gross - disc))
    subtotal = sum((r[0] for r in rows), Decimal(0))
    line_disc = sum((r[1] for r in rows), Decimal(0))
    after_lines = subtotal - line_disc
    if spec.discount_percent:
        total_disc = q2(after_lines * spec.discount_percent / HUNDRED)
    else:
        total_disc = q2(spec.discount_amount)
    if total_disc > after_lines:
        raise FinError("the discount is more than the total")
    charges = sum((q2(c.amount) for c in spec.charges), Decimal(0))
    base = after_lines - total_disc + charges
    rate = spec.vat_rate / HUNDRED
    if spec.vat == "excluded":
        vat = q2(base * rate)
        total = base + vat
    elif spec.vat == "included":
        total = base
        vat = q2(total * rate / (1 + rate))
    else:
        vat, total = Decimal(0), base
    return Totals(
        rows,
        subtotal,
        line_disc,
        total_disc,
        line_disc + total_disc,
        charges,
        base,
        vat,
        total,
        spec.vat,
        spec.vat_rate,
    )


# ------------------------------------------------------------------ rendering
@dataclass
class View:
    """What is printed besides the spec."""

    label: str  # QT-2026-0001-r2
    rev: int
    issued: date
    prepared_by: str
    company: dict[str, str]
    org_name: str = ""
    lang: str = "th"

    @property
    def company_name(self) -> str:
        return self.company.get("legal_name") or self.org_name or "-"


def missing_profile(company: dict[str, str]) -> list[str]:
    return [f for f in REQUIRED_PROFILE if not (company.get(f) or "").strip()]


def valid_until(issued: date, spec: QuoteSpec) -> date:
    return issued + timedelta(days=spec.valid_days)


def _company_lines(v: View) -> list[str]:
    c, lang = v.company, v.lang
    lines = [c.get("address") or t("quote_missing", lang, field=t("quote_f_address", lang))]
    tax = c.get("tax_id") or t("quote_missing", lang, field=t("quote_f_tax_id", lang))
    lines.append(f"{t('quote_tax_id', lang)}: {tax}")
    contact = [
        f"{t(k, lang)}: {c[f]}"
        for k, f in (("quote_phone", "phone"), ("quote_email", "email"), ("quote_web", "website"))
        if c.get(f)
    ]
    if contact:
        lines.append("  ".join(contact))
    return lines


def _customer_lines(spec: QuoteSpec, lang: str) -> list[str]:
    c = spec.customer
    out = [c.name]
    if c.address:
        out.append(c.address)
    if c.contact:
        out.append(f"{t('quote_attn', lang)}: {c.contact}")
    if c.tax_id:
        out.append(f"{t('quote_tax_id', lang)}: {c.tax_id}")
    return out


def _bank_line(company: dict[str, str], lang: str) -> str:
    parts = [
        company.get("bank_name", ""),
        company.get("bank_account_name", ""),
        company.get("bank_account_number", ""),
    ]
    return " / ".join(p for p in parts if p)


def _terms(spec: QuoteSpec, v: View) -> list[tuple[str, str]]:
    """(label, text) rows of the terms block."""
    lang = v.lang
    until = fmt_date(valid_until(v.issued, spec))
    rows = [
        (
            t("quote_valid_until", lang),
            t("quote_validity", lang, days=spec.valid_days, until=until),
        ),
        (t("quote_payment", lang), spec.payment_terms or t("quote_default_payment", lang)),
    ]
    bank = _bank_line(v.company, lang)
    if bank:
        rows.append((t("quote_bank", lang), bank))
    if spec.notes:
        rows.append((t("quote_notes", lang), spec.notes))
    return rows


def _total_rows(spec: QuoteSpec, tot: Totals, lang: str) -> list[tuple[str, str, bool]]:
    """(label, amount text, bold) rows of the totals block."""
    cur = spec.currency
    rows = [(t("quote_subtotal", lang), fmt_money(tot.subtotal), False)]
    if tot.discount:
        label = t("quote_discount", lang)
        if spec.discount_percent and not tot.line_discounts:
            label = t("quote_discount_pct", lang, pct=fmt_pct(spec.discount_percent))
        rows.append((label, "-" + fmt_money(tot.discount), False))
    for c in spec.charges:
        rows.append((c.description, fmt_money(q2(c.amount)), False))
    if tot.vat_mode == "excluded":
        rows.append((t("quote_after_discount", lang), fmt_money(tot.base), False))
        rows.append((t("quote_vat", lang, rate=fmt_pct(tot.vat_rate)), fmt_money(tot.vat), False))
    elif tot.vat_mode == "included":
        rows.append(
            (t("quote_vat_incl", lang, rate=fmt_pct(tot.vat_rate)), fmt_money(tot.vat), False)
        )
    else:
        rows.append((t("quote_no_vat", lang), "", False))
    rows.append((f"{t('quote_total', lang)} ({cur})", fmt_money(tot.total), True))
    return rows


def _words(spec: QuoteSpec, tot: Totals, lang: str) -> str:
    if spec.currency == "THB" and lang == "th":
        return thai_baht_text(tot.total)
    return ""


def _header_title(v: View) -> str:
    title = t("quote_title", v.lang)
    return title + (f" ({t('quote_rev', v.lang, n=v.rev)})" if v.rev > 1 else "")


def _item_row(i: int, ln: Line, amounts: tuple[Decimal, Decimal, Decimal], lang: str) -> list[str]:
    gross, disc, _net = amounts
    if disc:
        shown = f"{fmt_pct(ln.discount_percent)}%" if ln.discount_percent else fmt_money(disc)
    else:
        shown = "-"
    return [
        str(i),
        ln.description,
        fmt_qty(ln.qty),
        ln.unit or "-",
        fmt_money(ln.unit_price),
        shown,
        fmt_money(gross - disc),
    ]


class _QuotePdf(_DocPdf):
    number = ""

    def header(self) -> None:  # page 1 prints its own letterhead
        if self.page_no() > 1:
            self.set_font(self.family, "", 9)
            self.set_text_color(110)
            self.cell(0, 6, self._txt(f"{self.org}  ·  {self.number}"))
            self.ln(7)
            self.set_text_color(0)

    def kv(self, label: str, value: str, label_w: float = 38) -> None:
        self.set_font(self.family, "B", 10)
        self.set_x(self.l_margin)
        self.cell(label_w, 5.5, self._txt(label))
        self.ml(value, 10, 5.5, indent=label_w)

    def need(self, mm: float) -> None:
        if self.get_y() + mm > self.h - 18:
            self.add_page()


def quote_pdf(spec: QuoteSpec, v: View, fonts: FontSet, generated_at: datetime) -> bytes:
    lang, tot = v.lang, compute(spec)
    pdf = _QuotePdf(fonts, v.company_name, lang, generated_at)
    pdf.number = v.label
    pdf.add_page(orientation="P")
    logo = v.company.get("logo_url", "")
    # TODO(logo): the logo is not fetched (no safe image fetch helper yet); the name is printed.
    _ = logo
    pdf.ml(v.company_name, 16, 8, bold=True)
    for line in _company_lines(v):
        pdf.ml(line, 9.5, 5)
    pdf.ln(3)
    pdf.set_draw_color(150)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
    pdf.ln(3)
    pdf.ml(_header_title(v), 18, 9, bold=True)
    pdf.kv(t("quote_no", lang), v.label)
    pdf.kv(t("quote_date", lang), fmt_date(v.issued))
    pdf.kv(t("quote_valid_until", lang), fmt_date(valid_until(v.issued, spec)))
    if v.prepared_by:
        pdf.kv(t("quote_prepared_by", lang), v.prepared_by)
    pdf.ln(2)
    cust = _customer_lines(spec, lang)
    pdf.kv(t("quote_customer", lang), cust[0])
    for extra in cust[1:]:
        pdf.set_x(pdf.l_margin)
        pdf.ml(extra, 10, 5.5, indent=38)
    pdf.ln(3)

    # ---- items
    widths = (10, 70, 18, 16, 24, 18, 26)
    head = [
        t(k, lang)
        for k in (
            "quote_col_no",
            "quote_col_desc",
            "quote_col_qty",
            "quote_col_unit",
            "quote_col_price",
            "quote_col_disc",
            "quote_col_amount",
        )
    ]
    pdf.set_font(pdf.family, "", 9)
    rows = []
    for i, (ln, amounts) in enumerate(zip(spec.lines, tot.lines, strict=True), 1):
        row = _item_row(i, ln, amounts, lang)
        row[1] = "\n".join(pdf.wrap(pdf._txt(row[1]), widths[1] - 3))
        rows.append(row)
    pdf.set_x(pdf.l_margin)
    aligns = ("CENTER", "LEFT", "RIGHT", "CENTER", "RIGHT", "RIGHT", "RIGHT")
    pdf.set_font(pdf.family, "B", 9)
    with FPDF.table(pdf, col_widths=widths, text_align=aligns, line_height=5) as table:
        hrow = table.row()
        for h in head:
            hrow.cell(pdf._txt(h))
        pdf.set_font(pdf.family, "", 9)
        for r in rows:
            row = table.row()
            for c in r:
                row.cell(pdf._txt(c))
    pdf.ln(3)

    # ---- totals
    pdf.need(12 + 6.5 * (len(spec.charges) + 6))
    label_w, val_w = 70.0, 34.0
    x0 = pdf.w - pdf.r_margin - label_w - val_w
    for label, amount, bold in _total_rows(spec, tot, lang):
        pdf.set_font(pdf.family, "B" if bold else "", 10)
        pdf.set_x(x0)
        lab = pdf.wrap(pdf._txt(label), label_w - 2)[0]
        pdf.cell(label_w, 6.5, lab, border="T" if bold else 0)
        pdf.cell(val_w, 6.5, amount, align="R", border="T" if bold else 0)
        pdf.ln(6.5)
    words = _words(spec, tot, lang)
    if words:
        pdf.ln(1)
        pdf.ml(f"{t('quote_words', lang)}: ({words})", 10, 5.5, bold=True)
    pdf.ln(3)

    # ---- terms
    pdf.need(40)
    pdf.ml(t("quote_terms", lang), 11, 6.5, bold=True)
    for label, text in _terms(spec, v):
        pdf.kv(label, text, label_w=42)

    # ---- signatures
    pdf.need(50)
    pdf.ln(10)
    col = (pdf.w - pdf.l_margin - pdf.r_margin) / 2
    y = pdf.get_y()
    c = v.company
    signer = c.get("signer_name") or t("quote_missing", lang, field=t("quote_f_signer_name", lang))
    blocks = [
        (t("quote_accepted_by", lang), "", ""),
        (t("quote_signed_by", lang), signer, c.get("signer_title", "")),
    ]
    for i, (title, name, role) in enumerate(blocks):
        x = pdf.l_margin + i * col
        pdf.set_xy(x + 4, y)
        pdf.set_font(pdf.family, "", 10)
        pdf.cell(col - 8, 6, pdf._txt(title), align="C")
        pdf.line(x + 8, y + 18, x + col - 8, y + 18)
        pdf.set_xy(x + 4, y + 19)
        pdf.cell(
            col - 8,
            5.5,
            pdf._txt(f"({name})" if name else "(                         )"),
            align="C",
        )
        pdf.set_xy(x + 4, y + 25)
        pdf.cell(col - 8, 5.5, pdf._txt(role), align="C")
        pdf.set_xy(x + 4, y + 31)
        pdf.cell(col - 8, 5.5, pdf._txt(t("quote_date_line", lang)), align="C")
    return bytes(pdf.output())


def quote_docx(spec: QuoteSpec, v: View) -> bytes:
    """The same content as the PDF in a simple editable Word file."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    lang, tot = v.lang, compute(spec)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "TH Sarabun New"
    normal.font.size = Pt(14)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        from docx.oxml import OxmlElement

        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), "TH Sarabun New")  # Tahoma is the fallback Word picks without it

    def para(text: str = "", bold: bool = False, size: int | None = None, align=None) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if align is not None:
            p.alignment = align

    para(v.company_name, bold=True, size=20)
    for line in _company_lines(v):
        para(line, size=12)
    para(_header_title(v), bold=True, size=22)
    for label, value in (
        (t("quote_no", lang), v.label),
        (t("quote_date", lang), fmt_date(v.issued)),
        (t("quote_valid_until", lang), fmt_date(valid_until(v.issued, spec))),
        (t("quote_prepared_by", lang), v.prepared_by),
    ):
        if value:
            para(f"{label}: {value}")
    cust = _customer_lines(spec, lang)
    para(f"{t('quote_customer', lang)}: {cust[0]}", bold=True)
    for extra in cust[1:]:
        para(extra)

    head = [
        t(k, lang)
        for k in (
            "quote_col_no",
            "quote_col_desc",
            "quote_col_qty",
            "quote_col_unit",
            "quote_col_price",
            "quote_col_disc",
            "quote_col_amount",
        )
    ]
    table = doc.add_table(rows=1 + len(spec.lines), cols=len(head))
    table.style = "Table Grid"
    for c, h in enumerate(head):
        cell = table.cell(0, c)
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for i, (ln, amounts) in enumerate(zip(spec.lines, tot.lines, strict=True), 1):
        for c, val in enumerate(_item_row(i, ln, amounts, lang)):
            cell = table.cell(i, c)
            cell.text = val
            if c >= 2 and c != 3:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    rows = _total_rows(spec, tot, lang)
    totals = doc.add_table(rows=len(rows), cols=2)
    totals.style = "Table Grid"
    for i, (label, amount, bold) in enumerate(rows):
        totals.cell(i, 0).text = label
        totals.cell(i, 1).text = amount
        totals.cell(i, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if bold:
            for c in (0, 1):
                for r in totals.cell(i, c).paragraphs[0].runs:
                    r.bold = True
    words = _words(spec, tot, lang)
    if words:
        para(f"{t('quote_words', lang)}: ({words})", bold=True)
    para(t("quote_terms", lang), bold=True, size=16)
    for label, text in _terms(spec, v):
        para(f"{label}: {text}")
    c = v.company
    signer = c.get("signer_name") or t("quote_missing", lang, field=t("quote_f_signer_name", lang))
    sig = doc.add_table(rows=1, cols=2)
    sig.cell(
        0, 0
    ).text = f"\n\n{t('quote_accepted_by', lang)}\n(                         )\n{t('quote_date_line', lang)}"
    sig.cell(0, 1).text = (
        f"\n\n{t('quote_signed_by', lang)}\n({signer})\n{c.get('signer_title', '')}\n"
        f"{t('quote_date_line', lang)}"
    )
    for cell in sig.rows[0].cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ store
class QuoteRepo:
    """Quotations of every guild in the DocStore kind `dory_quotes` (key = base number). One bot
    process: the counter is guarded by a lock and by never reusing a number that exists."""

    def __init__(self, docs: DocStore, clock: Callable[[], datetime] | None = None):
        self.docs = docs
        self._clock = clock or datetime.now
        self._lock = threading.Lock()

    def now(self) -> str:
        return self._clock().isoformat(timespec="seconds")

    def all(self, guild_id: int) -> list[QuoteDoc]:
        out: list[QuoteDoc] = []
        for d in self.docs.find(QUOTES, guild_id, lambda d: "number" in d and "revisions" in d):
            try:
                out.append(QuoteDoc(**{k: v for k, v in d.items() if k in QuoteDoc.model_fields}))
            except ValueError:
                logger.warning("unreadable quotation skipped")
        return out

    def get(self, guild_id: int, number: str) -> QuoteDoc | None:
        d = self.docs.get(QUOTES, guild_id, number)
        if not d or "revisions" not in d:
            return None
        try:
            return QuoteDoc(**{k: v for k, v in d.items() if k in QuoteDoc.model_fields})
        except ValueError:
            return None

    def save(self, guild_id: int, q: QuoteDoc) -> None:
        q.revisions = q.revisions[-MAX_REVISIONS:]
        self.docs.put(QUOTES, guild_id, q.number, q.model_dump(mode="json"))

    def create(
        self, guild_id: int, spec: QuoteSpec, *, year: int, issued: date, by_id: int,
        by_name: str, channel_id: int,
    ) -> QuoteDoc:  # fmt: skip
        """The next number of the guild for `year` (never reused) with revision 1."""
        with self._lock:
            ckey = f"_seq_{year}"
            seen = (self.docs.get(QUOTES, guild_id, ckey) or {}).get("n", 0)
            top = max((q.seq for q in self.all(guild_id) if q.year == year), default=0)
            seq = max(int(seen), top) + 1
            number = f"QT-{year}-{seq:04d}"
            while self.docs.get(QUOTES, guild_id, number) is not None:
                seq += 1
                number = f"QT-{year}-{seq:04d}"
            self.docs.put(QUOTES, guild_id, ckey, {"n": seq})
            now = self.now()
            q = QuoteDoc(
                number=number, year=year, seq=seq, channel_id=channel_id, author_id=by_id,
                author_name=by_name, created_at=now, updated_at=now,
                revisions=[self._rev(1, spec, issued, by_id, by_name, now, "")],
            )  # fmt: skip
            self.save(guild_id, q)
            return q

    def add_revision(
        self, guild_id: int, q: QuoteDoc, spec: QuoteSpec, *, issued: date, by_id: int,
        by_name: str, summary: str,
    ) -> QuoteDoc:  # fmt: skip
        with self._lock:
            fresh = self.get(guild_id, q.number) or q
            now = self.now()
            fresh.revisions.append(
                self._rev(fresh.latest.rev + 1, spec, issued, by_id, by_name, now, summary)
            )
            fresh.updated_at = now
            self.save(guild_id, fresh)
            return fresh

    @staticmethod
    def _rev(rev, spec, issued, by_id, by_name, now, summary) -> Revision:
        return Revision(
            rev=rev, spec=spec, total=str(compute(spec).total), issued=issued.isoformat(),
            by_id=by_id, by_name=by_name, at=now, summary=summary[:200],
        )  # fmt: skip


def can_read(
    q: QuoteDoc, user_id: int, is_exec: bool, channel_readable: Callable[[int], bool]
) -> bool:
    """The author, executives, and anyone who can read the channel the quotation was made in."""
    if is_exec or q.author_id == user_id:
        return True
    try:
        return bool(q.channel_id) and bool(channel_readable(q.channel_id))
    except Exception:
        return False


# ------------------------------------------------------------------ tools
@dataclass
class QuoteCtx:
    """Everything the tools need, bound in code to one guild, one channel and one speaker."""

    guild_id: int
    guild_name: str
    asker_id: int
    asker_name: str
    is_exec: bool
    can_issue: bool  # executive or a member with one of the quote roles
    quote_roles: list[str]
    channel_id: int
    repo: QuoteRepo
    company: CompanyRepo
    lang: str
    now: Callable[[], datetime]
    channel_readable: Callable[[int], bool] = lambda cid: False
    fonts: Callable[[], FontSet] = lambda: FontSet(None, None, False)
    outbox: list[tuple[str, bytes]] = field(default_factory=list)

    # ---- access
    def _refusal(self) -> dict:
        roles = (
            t("quote_refusal_roles", self.lang, roles=", ".join(self.quote_roles))
            if self.quote_roles
            else ""
        )
        return {
            "status": "not_allowed",
            "message": t("quote_refusal", self.lang, roles=roles),
            "instruction": "Tell the user this politely in their language; do not issue anything.",
        }

    def readable(self) -> list[QuoteDoc]:
        docs = [
            q
            for q in self.repo.all(self.guild_id)
            if can_read(q, self.asker_id, self.is_exec, self.channel_readable)
        ]
        docs.sort(key=lambda q: q.updated_at, reverse=True)
        return docs

    def _find(self, number: str) -> tuple[QuoteDoc | None, int | None, dict | None]:
        """(quotation, requested revision, error) for a number, or the latest one made in this
        channel when `number` is empty."""
        if (number or "").strip():
            parsed = parse_number(number)
            if parsed is None:
                return (
                    None,
                    None,
                    {"status": "error", "error": "not a quotation number (QT-YYYY-NNNN)"},
                )
            base, rev = parsed
            q = self.repo.get(self.guild_id, base)
            if q is None or not can_read(q, self.asker_id, self.is_exec, self.channel_readable):
                return None, None, {"status": "not_found", "number": base}
            if rev is not None and not 1 <= rev <= q.latest.rev:
                return None, None, {"status": "not_found", "number": f"{base}-r{rev}"}
            return q, rev, None
        here = [q for q in self.readable() if q.channel_id == self.channel_id]
        if not here:
            return None, None, {"status": "not_found", "error": "no quotation in this channel yet"}
        return here[0], None, None

    # ---- rendering
    def _company(self) -> dict[str, str]:
        return dict(self.company.get(self.guild_id).profile)

    def _view(self, q: QuoteDoc, rev: Revision, company: dict[str, str]) -> View:
        return View(
            label=q.label(rev.rev), rev=rev.rev, issued=date.fromisoformat(rev.issued),
            prepared_by=rev.by_name, company=company, org_name=self.guild_name, lang=self.lang,
        )  # fmt: skip

    def render(self, q: QuoteDoc, rev: Revision) -> list[str]:
        """Attach the PDF and the .docx of one revision; returns the file names."""
        company = self._company()
        view = self._view(q, rev, company)
        names = []
        files = (
            ("pdf", lambda: quote_pdf(rev.spec, view, self.fonts(), self.now())),
            ("docx", lambda: quote_docx(rev.spec, view)),
        )
        for ext, build in files:
            name = f"{view.label}.{ext}"
            self.outbox.append((name, build()))
            names.append(name)
        return names

    def _result(self, q: QuoteDoc, rev: Revision, status: str = "ok") -> dict:
        spec, tot = rev.spec, compute(rev.spec)
        company = self._company()
        out: dict[str, Any] = {
            "status": status,
            "number": q.label(rev.rev),
            "base_number": q.number,
            "revision": rev.rev,
            "customer": spec.customer.name,
            "currency": spec.currency,
            "subtotal": fmt_money(tot.subtotal),
            "discount": fmt_money(tot.discount),
            "charges": fmt_money(tot.charges),
            "vat": fmt_money(tot.vat),
            "vat_mode": spec.vat,
            "total": fmt_money(tot.total),
            "valid_until": fmt_date(valid_until(date.fromisoformat(rev.issued), spec)),
            "files": [],
        }
        gaps = missing_profile(company)
        if gaps:
            out["missing_company_fields"] = gaps
            out["teach_hint"] = t("quote_teach_hint", self.lang)
        return out

    def _issue_guard(self) -> dict | None:
        return None if self.can_issue else self._refusal()

    # ---- make
    def make(self, args: dict) -> dict:
        if (g := self._issue_guard()) is not None:
            return g
        spec = self._spec_from(args)
        issued = self.now().date()
        q = self.repo.create(
            self.guild_id, spec, year=issued.year, issued=issued, by_id=self.asker_id,
            by_name=self.asker_name, channel_id=self.channel_id,
        )  # fmt: skip
        res = self._result(q, q.latest)
        res["files"] = self.render(q, q.latest)
        return res

    def _spec_from(self, a: dict) -> QuoteSpec:
        name = _txt(a.get("customer_name"), 120)
        if not name:
            raise FinError("the customer name is required")
        dp, da = (
            _pct(a.get("discount_percent"), "discount_percent"),
            _num(a.get("discount_amount") or 0, "discount_amount"),
        )
        spec = QuoteSpec(
            customer=Customer(
                name=name,
                address=_txt(a.get("customer_address")),
                contact=_txt(a.get("customer_contact"), 120),
                tax_id=_txt(a.get("customer_tax_id"), 30),
            ),
            lines=build_lines(a.get("items")),
            charges=build_charges(a.get("extra_charges")),
            discount_percent=dp,
            discount_amount=da,
            currency=clean_currency(a.get("currency")),
            vat=clean_vat(a.get("vat")),
            vat_rate=_num(
                7 if a.get("vat_rate_percent") in (None, "") else a.get("vat_rate_percent"),
                "vat_rate_percent",
                hi=Decimal(30),
            ),
            valid_days=_days(a.get("valid_days")),
            payment_terms=_txt(a.get("payment_terms")),
            notes=_txt(a.get("notes"), MAX_NOTES),
        )
        return check_spec(spec)

    # ---- revise
    def revise(self, number: str, a: dict) -> dict:
        if (g := self._issue_guard()) is not None:
            return g
        q, _rev, err = self._find(number)
        if err is not None:
            return err
        spec = self._apply_changes(q.latest.spec.model_copy(deep=True), a)
        summary = _summary_of(a)
        issued = self.now().date()
        q = self.repo.add_revision(
            self.guild_id, q, spec, issued=issued, by_id=self.asker_id, by_name=self.asker_name,
            summary=summary,
        )  # fmt: skip
        res = self._result(q, q.latest)
        res["files"] = self.render(q, q.latest)
        return res

    def _apply_changes(self, spec: QuoteSpec, a: dict) -> QuoteSpec:
        changed = False
        for ch in a.get("item_changes") or []:
            if not isinstance(ch, dict):
                raise FinError("item_changes must be a list of objects")
            idx = self._line_index(spec, ch)
            ln = spec.lines[idx].model_dump()
            for key in ("description", "unit", "qty", "unit_price"):
                if ch.get(key) not in (None, ""):
                    ln[key] = ch[key]
            if ch.get("discount_percent") is not None:
                ln["discount_percent"], ln["discount_amount"] = ch["discount_percent"], 0
            if ch.get("discount_amount") is not None:
                ln["discount_amount"], ln["discount_percent"] = ch["discount_amount"], 0
            spec.lines[idx] = build_line({k: str(v) for k, v in ln.items()}, idx + 1)
            changed = True
        gone = sorted({int(i) for i in (a.get("remove_lines") or [])}, reverse=True)
        for i in gone:
            if not 1 <= i <= len(spec.lines):
                raise FinError(f"there is no line {i}")
            del spec.lines[i - 1]
            changed = True
        if a.get("add_items"):
            spec.lines += build_lines(a["add_items"])
            changed = True
        if a.get("add_charges"):
            spec.charges += build_charges(a["add_charges"])
            changed = True
        for i in sorted({int(i) for i in (a.get("remove_charges") or [])}, reverse=True):
            if not 1 <= i <= len(spec.charges):
                raise FinError(f"there is no extra charge {i}")
            del spec.charges[i - 1]
            changed = True
        if a.get("discount_percent") is not None:
            spec.discount_percent, spec.discount_amount = (
                _pct(a["discount_percent"], "discount_percent"),
                Decimal(0),
            )
            changed = True
        if a.get("discount_amount") is not None:
            spec.discount_amount, spec.discount_percent = (
                _num(a["discount_amount"], "discount_amount"),
                Decimal(0),
            )
            changed = True
        for key, field_name, limit in (
            ("customer_name", "name", 120),
            ("customer_address", "address", MAX_TEXT),
            ("customer_contact", "contact", 120),
            ("customer_tax_id", "tax_id", 30),
        ):
            if a.get(key) is not None:
                setattr(spec.customer, field_name, _txt(a[key], limit))
                changed = True
        if a.get("payment_terms") is not None:
            spec.payment_terms, changed = _txt(a["payment_terms"]), True
        if a.get("notes") is not None:
            spec.notes, changed = _txt(a["notes"], MAX_NOTES), True
        if a.get("valid_days") is not None:
            spec.valid_days, changed = _days(a["valid_days"]), True
        if a.get("vat") is not None:
            spec.vat, changed = clean_vat(a["vat"]), True
        if a.get("vat_rate_percent") is not None:
            spec.vat_rate = _num(a["vat_rate_percent"], "vat_rate_percent", hi=Decimal(30))
            changed = True
        if a.get("currency") is not None:
            spec.currency, changed = clean_currency(a["currency"]), True
        if not changed:
            raise FinError("nothing to change")
        return check_spec(spec)

    @staticmethod
    def _line_index(spec: QuoteSpec, ch: dict) -> int:
        if ch.get("line") not in (None, ""):
            n = int(D(ch["line"], "line"))
            if not 1 <= n <= len(spec.lines):
                raise FinError(f"there is no line {n}")
            return n - 1
        if len(spec.lines) == 1:
            return 0
        key = _txt(ch.get("match"), 80).casefold()
        hits = [i for i, ln in enumerate(spec.lines) if key and key in ln.description.casefold()]
        if len(hits) != 1:
            raise FinError("say which line (its number) to change")
        return hits[0]

    # ---- read
    def listing(self, customer: str = "", since: str = "") -> dict:
        low = _txt(customer, 80).casefold()
        start: date | None = None
        if (since or "").strip():
            try:
                start = date.fromisoformat(since.strip()[:10])
            except ValueError:
                return {"status": "error", "error": "since must be a date like 2026-10-01"}
        rows = []
        for q in self.readable():
            r = q.latest
            if low and low not in r.spec.customer.name.casefold():
                continue
            if start and date.fromisoformat(q.created_at[:10] or "1970-01-01") < start:
                continue
            rows.append(
                {
                    "number": q.label(r.rev),
                    "customer": r.spec.customer.name,
                    "total": fmt_money(Decimal(r.total)),
                    "currency": r.spec.currency,
                    "issued": r.issued,
                    "by": q.author_name,
                    "revisions": r.rev,
                }
            )
        return {"status": "ok", "count": len(rows), "quotations": rows[:MAX_LIST]}

    def get(self, number: str) -> dict:
        q, rev_no, err = self._find(number)
        if err is not None:
            return err
        rev = next((r for r in q.revisions if r.rev == (rev_no or q.latest.rev)), q.latest)
        res = self._result(q, rev, "ok")
        res["files"] = self.render(q, rev)
        return res


def _days(v: Any) -> int:
    if v in (None, ""):
        return DEFAULT_VALID_DAYS
    d = D(v, "valid_days")
    if d != d.to_integral_value() or not 1 <= d <= 365:
        raise FinError("valid_days must be a whole number between 1 and 365")
    return int(d)


def _summary_of(a: dict) -> str:
    keys = [k for k, v in a.items() if v not in (None, "", [], {})]
    return ", ".join(sorted(keys))


def quotation_tools(ctx: QuoteCtx) -> list[Callable]:
    """The functions `_agent_kwargs` appends to `extra_tools` (docstrings = tool descriptions).
    The PDF and .docx are collected in `ctx.outbox` (`outbox` attribute: see `qa.ask_async`)."""

    def safe(fn: Callable[[], dict]) -> dict:
        try:
            return fn()
        except FinError as e:
            return {"status": "error", "error": str(e)}
        except (ValueError, TypeError, ArithmeticError) as e:
            return {"status": "error", "error": f"invalid input: {e}"[:200]}
        except Exception:
            logger.exception("quotation tool failed")
            return {"status": "error", "error": "the quotation tool failed; try again"}

    def make_quotation(
        customer_name: str,
        items: list[dict[str, Any]],
        customer_address: str = "",
        customer_contact: str = "",
        customer_tax_id: str = "",
        discount_percent: float | None = None,
        discount_amount: float | None = None,
        extra_charges: list[dict[str, Any]] | None = None,
        vat: str = "excluded",
        vat_rate_percent: float | None = None,
        currency: str = "THB",
        valid_days: int = 30,
        payment_terms: str = "",
        notes: str = "",
    ) -> dict:
        """Make a quotation (ใบเสนอราคา) as a PDF and an editable Word file, attached to the reply,
        with a running number QT-YYYY-NNNN. Only executives and members with an issuing role may;
        otherwise the result says so (tell the user kindly). Use it whenever someone asks for a
        quotation / price offer / ใบเสนอราคา with a customer and items. Never invent prices or
        quantities: ask for what is missing. Reply short: the number, the customer, the total and
        that the files are attached; if the result lists missing_company_fields, say which company
        details to teach Dory (and the teach_hint). Takes a few seconds.

        Args:
            customer_name: the customer (company or person).
            items: lines, each {description, qty, unit, unit_price, discount_percent?,
                discount_amount?} (a line discount is a percent OR an amount); 1-50 lines.
            customer_address: optional.
            customer_contact: optional contact person.
            customer_tax_id: optional.
            discount_percent: discount on the whole quotation in percent (0-100), after line
                discounts. "ส่วนลด 5%" -> 5.
            discount_amount: discount on the whole quotation as an amount (not with a percent).
            extra_charges: optional extra lines such as shipping: [{description, amount}].
            vat: "excluded" (default: 7% VAT added on top), "included" (the prices already include
                VAT) or "none" (no VAT).
            vat_rate_percent: VAT rate when not 7.
            currency: 3-letter code, default THB.
            valid_days: how many days the price stands (default 30; "ยืนราคา N วัน").
            payment_terms: free text; empty uses the default (pay within 30 days of the invoice).
            notes: optional notes printed on the document.
        """
        args = dict(
            customer_name=customer_name, items=items, customer_address=customer_address,
            customer_contact=customer_contact, customer_tax_id=customer_tax_id,
            discount_percent=discount_percent, discount_amount=discount_amount,
            extra_charges=extra_charges, vat=vat, vat_rate_percent=vat_rate_percent,
            currency=currency, valid_days=valid_days, payment_terms=payment_terms, notes=notes,
        )  # fmt: skip
        return safe(lambda: ctx.make(args))

    def revise_quotation(
        number: str = "",
        item_changes: list[dict[str, Any]] | None = None,
        add_items: list[dict[str, Any]] | None = None,
        remove_lines: list[int] | None = None,
        add_charges: list[dict[str, Any]] | None = None,
        remove_charges: list[int] | None = None,
        discount_percent: float | None = None,
        discount_amount: float | None = None,
        customer_name: str | None = None,
        customer_address: str | None = None,
        customer_contact: str | None = None,
        customer_tax_id: str | None = None,
        valid_days: int | None = None,
        payment_terms: str | None = None,
        notes: str | None = None,
        vat: str | None = None,
        vat_rate_percent: float | None = None,
        currency: str | None = None,
    ) -> dict:
        """Change a quotation: makes a NEW REVISION (same base number, -r2, -r3...) and attaches
        the new PDF and Word file. "เปลี่ยนจำนวนเป็น 120", "เพิ่มค่าส่ง 500", "ลดเหลือ 3%". Only
        executives and issuing roles. Leave everything you do not change out. Reply short: the
        new number and the new total.

        Args:
            number: QT-2026-0001 (empty = the latest quotation made in this channel).
            item_changes: [{line (1-based; may be left out when there is one line), qty?, unit?,
                unit_price?, description?, discount_percent?, discount_amount?}].
            add_items: new lines, same shape as in make_quotation.
            remove_lines: line numbers (1-based) to delete.
            add_charges: extra charges [{description, amount}], e.g. shipping.
            remove_charges: extra charge numbers (1-based) to delete.
            discount_percent: new discount on the whole quotation in percent (0 removes it).
            discount_amount: new discount on the whole quotation as an amount.
            customer_name: new customer name.
            customer_address: new customer address.
            customer_contact: new contact person.
            customer_tax_id: new customer tax id.
            valid_days: new number of days the price stands.
            payment_terms: new payment terms.
            notes: new notes (empty text clears them).
            vat: excluded | included | none.
            vat_rate_percent: VAT rate.
            currency: 3-letter code.
        """
        args = dict(
            item_changes=item_changes, add_items=add_items, remove_lines=remove_lines,
            add_charges=add_charges, remove_charges=remove_charges,
            discount_percent=discount_percent, discount_amount=discount_amount,
            customer_name=customer_name, customer_address=customer_address,
            customer_contact=customer_contact, customer_tax_id=customer_tax_id,
            valid_days=valid_days, payment_terms=payment_terms, notes=notes, vat=vat,
            vat_rate_percent=vat_rate_percent, currency=currency,
        )  # fmt: skip
        return safe(lambda: ctx.revise(number, args))

    def list_quotations(customer: str = "", since: str = "") -> dict:
        """List the quotations the user may see (newest first): number, customer, total, date,
        who made it, revisions.

        Args:
            customer: optional part of the customer name.
            since: optional date YYYY-MM-DD; only quotations made on or after it.
        """
        return safe(lambda: ctx.listing(customer, since))

    def get_quotation(number: str = "") -> dict:
        """Send a quotation's PDF and Word files again (re-rendered from the saved data). "QT-2026-0001"
        is the latest revision, "QT-2026-0001-r2" a specific one; empty = the latest quotation made
        in this channel.

        Args:
            number: the quotation number.
        """
        return safe(lambda: ctx.get(number))

    tools = [make_quotation, revise_quotation, list_quotations, get_quotation]
    for fn in tools:
        fn.outbox = ctx.outbox  # type: ignore[attr-defined]  # qa.ask_async attaches these files
    return tools
