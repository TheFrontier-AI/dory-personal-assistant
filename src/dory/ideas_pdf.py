"""Idea documents as PDF / DOCX / Markdown: a small Markdown renderer on the Thai-capable PDF base
of `pdf_report`, plus the 9-box Business Model Canvas page."""

from __future__ import annotations

import io
import re
from datetime import datetime

from fpdf import FPDF
from fpdf.enums import XPos, YPos

from .frameworks import KINDS, split_h2
from .pdf_report import FontSet, _Pdf

_EMOJI = re.compile("[\U0001f000-\U0001faff☀-➿️‍]")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_BULLET = re.compile(r"^(\s*)(?:[-*•]|\d+[.)])\s+(.*)$")


def plain(s: str) -> str:
    """Inline Markdown and emoji removed (the PDF fonts have no emoji)."""
    s = _EMOJI.sub("", s)
    s = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r"\1 (\2)", s)
    s = re.sub(r"(\*\*|__)(.+?)\1", r"\2", s)
    s = re.sub(r"(?<![\w*])[*_](.+?)[*_](?![\w*])", r"\1", s)
    return re.sub(r"`([^`]*)`", r"\1", s).strip()


def _cells(row: str) -> list[str]:
    return [plain(c) for c in row.strip().strip("|").split("|")]


_NO_BREAK_BEFORE = set("ะัาำิีึืฺุู็่้๊๋์ํ๎ๆฯ")  # Thai vowel signs / tone marks / repeat mark
_NO_BREAK_AFTER = set("เแโใไ")  # leading vowels belong to the consonant after them


def _clusters(word: str) -> list[str]:
    """The word cut into pieces that may start a line (Thai: never before a vowel sign or tone
    mark, never right after a leading vowel)."""
    out: list[str] = []
    for i, ch in enumerate(word):
        if out and ch not in _NO_BREAK_BEFORE and word[i - 1] not in _NO_BREAK_AFTER:
            out.append(ch)
        elif out:
            out[-1] += ch
        else:
            out.append(ch)
    return out


class _DocPdf(_Pdf):
    def wrap(self, text: str, width: float) -> list[str]:
        """Lines that fit `width` (mm) in the current font: at spaces, and inside a word only when
        it cannot go on a line of its own (Thai has no spaces between phrases)."""
        sw = self.get_string_width
        lines: list[str] = []
        cur = ""
        for word in text.split():
            trial = f"{cur} {word}" if cur else word
            if sw(trial) <= width:
                cur = trial
                continue
            if cur and sw(word) <= width:
                lines.append(cur)
                cur = word
                continue
            buf = f"{cur} " if cur else ""  # a long word: fill the line, then continue below
            for cluster in _clusters(word):
                if sw(buf + cluster) > width and buf.strip():
                    lines.append(buf.rstrip())
                    buf = ""
                buf += cluster
            cur = buf
        if cur:
            lines.append(cur)
        return lines or [""]

    def ml(
        self, text: str, size: float, line_h: float, *, bold=False, indent=0.0, prefix=""
    ) -> None:
        """A wrapped, left-aligned paragraph (optionally with a hanging bullet `prefix`)."""
        self.set_font(self.family, "B" if bold else "", size)
        text = self._txt(text)
        width = self.w - self.l_margin - self.r_margin - indent
        pw = self.get_string_width(prefix) if prefix else 0
        lines = self.wrap(text, width - pw)
        for i, line in enumerate(lines):
            self.set_x(self.l_margin + indent)
            if prefix:
                self.cell(pw, line_h, prefix if i == 0 else "")
            self.cell(0, line_h, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    def h1(self, text: str) -> None:
        self.ml(text, 18, 9, bold=True)

    def h2(self, text: str) -> None:
        self.ln(3)
        self.ml(text, 12, 7, bold=True)

    def p(self, text: str, size: int = 10, bold: bool = False) -> None:
        self.ml(text, size, 5.5, bold=bold)

    def write_markdown(self, md: str, *, demote: int = 0) -> None:
        lines = md.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip():
                i += 1
                continue
            if (
                line.lstrip().startswith("|")
                and i + 1 < len(lines)
                and _TABLE_SEP.match(lines[i + 1])
            ):
                rows = [_cells(line)]
                i += 2
                while i < len(lines) and lines[i].lstrip().startswith("|"):
                    rows.append(_cells(lines[i]))
                    i += 1
                self._md_table(rows)
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", line)
            if m:
                level = min(len(m.group(1)) + demote, 4)
                text = plain(m.group(2))
                if level == 1:
                    self.h1(text)
                elif level == 2:
                    self.h2(text)
                else:
                    self.ln(1.5)
                    self.p(text, size=11 if level == 3 else 10, bold=True)
                i += 1
                continue
            b = _BULLET.match(line)
            if b:
                indent = min(len(b.group(1)) // 2, 3) * 5
                self.ml(plain(b.group(2)), 10, 5.5, indent=indent, prefix="- ")
                i += 1
                continue
            para = [line]
            i += 1
            while (
                i < len(lines)
                and lines[i].strip()
                and not re.match(r"^(#{1,6}\s|\s*\|)", lines[i])
                and not _BULLET.match(lines[i])
            ):
                para.append(lines[i].strip())
                i += 1
            self.p(plain(" ".join(para)))
            self.ln(1)

    def _md_table(self, rows: list[list[str]]) -> None:
        width = self.w - self.l_margin - self.r_margin
        n = max(len(r) for r in rows)
        rows = [r + [""] * (n - len(r)) for r in rows]
        self.set_font(self.family, "", 9)
        self.set_x(self.l_margin)
        with FPDF.table(
            self, col_widths=tuple([width / n] * n), text_align="LEFT", line_height=5
        ) as table:
            for k, r in enumerate(rows):
                row = table.row()
                for c in r:
                    if k == 0:
                        self.set_font(self.family, "B", 9)
                    else:
                        self.set_font(self.family, "", 9)
                    row.cell(self._txt(c))
        self.ln(2)

    def image_page(self, data: bytes, caption: str, heading: str) -> None:
        """One page: the section heading, the image fitted to the page, the caption below."""
        from PIL import Image

        self.add_page(orientation="P")
        self.h2(heading)
        with Image.open(io.BytesIO(data)) as im:
            iw, ih = im.size
        width = self.w - self.l_margin - self.r_margin
        height = self.h - self.get_y() - self.b_margin - 24  # room for the caption
        scale = min(width / iw, height / ih)
        w, h = iw * scale, ih * scale
        y = self.get_y() + 2
        self.image(io.BytesIO(data), x=self.l_margin + (width - w) / 2, y=y, w=w, h=h)
        self.set_y(y + h + 3)
        self.ml(caption, 10, 5.5)

    # ---- the 9-box canvas
    def _box(self, x: float, y: float, w: float, h: float, title: str, bullets: list[str]) -> None:
        self.set_draw_color(120)
        self.rect(x, y, w, h)
        inner = w - 3
        yy = y + 1.2
        self.set_font(self.family, "B", 9)
        self.set_text_color(30, 60, 120)
        for line in self.wrap(self._txt(title), inner):
            self.set_xy(x + 1.5, yy)
            self.cell(inner, 4.5, line)
            yy += 4.5
        self.set_text_color(0)
        self.set_font(self.family, "", 7.5)
        line_h, limit = 3.7, y + h - 1.5
        pw = self.get_string_width("- ")
        for b in bullets:
            lines = self.wrap(self._txt(b), inner - pw)
            if yy + line_h * len(lines) > limit:
                if yy + line_h <= limit:
                    self.set_xy(x + 1.5, yy)
                    self.cell(inner, line_h, "...")
                break
            for i, line in enumerate(lines):
                self.set_xy(x + 1.5, yy)
                self.cell(pw, line_h, "- " if i == 0 else "")
                self.cell(inner - pw, line_h, line)
                yy += line_h

    def canvas_page(self, title: str, blocks: dict[str, list[str]], heads: dict[str, str]) -> None:
        """`blocks`: English block name -> bullets; `heads`: English name -> printed title."""
        self.add_page(orientation="L")
        self.set_font(self.family, "B", 14)
        self.cell(0, 8, self._txt(title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        x0, y0 = self.l_margin, self.get_y() + 1
        total_w = self.w - self.l_margin - self.r_margin
        col = total_w / 5
        top_h = 118.0
        bottom_h = max(self.h - y0 - 18 - top_h, 30.0)

        def b(name: str, x: float, y: float, w: float, h: float) -> None:
            self._box(x, y, w, h, heads.get(name, name), blocks.get(name, []))

        b("Key Partnerships", x0, y0, col, top_h)
        b("Key Activities", x0 + col, y0, col, top_h / 2)
        b("Key Resources", x0 + col, y0 + top_h / 2, col, top_h / 2)
        b("Value Propositions", x0 + 2 * col, y0, col, top_h)
        b("Customer Relationships", x0 + 3 * col, y0, col, top_h / 2)
        b("Channels", x0 + 3 * col, y0 + top_h / 2, col, top_h / 2)
        b("Customer Segments", x0 + 4 * col, y0, col, top_h)
        half = total_w / 2
        b("Cost Structure", x0, y0 + top_h, half, bottom_h)
        b("Revenue Streams", x0 + half, y0 + top_h, half, bottom_h)


def canvas_blocks(markdown: str, lang: str) -> tuple[dict[str, list[str]], dict[str, str]]:
    """Bullets per BMC block from a Markdown document with the nine `##` headings."""
    pairs = KINDS["bmc"][2]
    by_head = {re.sub(r"[^\w฀-๿]+", "", h.casefold()): body for h, body in split_h2(markdown)}
    blocks: dict[str, list[str]] = {}
    heads: dict[str, str] = {}
    for th, en in pairs:
        heads[en] = f"{th} ({en})" if lang == "th" else en
        body = ""
        for key, text in by_head.items():
            if (
                re.sub(r"[^\w฀-๿]+", "", en.casefold()) in key
                or re.sub(r"[^\w฀-๿]+", "", th) in key
            ):
                body = text
                break
        items = []
        for line in body.splitlines():
            m = _BULLET.match(line)
            s = plain(m.group(2) if m else line)
            if s and not s.startswith("|") and not s.startswith("#"):
                items.append(s)
        blocks[en] = items
    return blocks, heads


def markdown_pdf(
    *,
    org: str,
    title: str,
    markdown: str,
    lang: str,
    generated_at: datetime,
    fonts: FontSet,
    canvas: str | None = None,
    byline: str = "",
    demote: int = 0,
    images: list[tuple[bytes, str]] | None = None,
    images_title: str = "",
) -> bytes:
    """PDF of a Markdown document. `canvas`: Markdown holding a BMC; its nine-box page is
    inserted first. `demote`: shift the headings of the body down (a section inside a document).
    `images`: [(JPEG/PNG bytes, caption)] added after the body, one per page."""
    pdf = _DocPdf(fonts, org, lang, generated_at)
    if canvas:
        blocks, heads = canvas_blocks(canvas, lang)
        pdf.canvas_page(title, blocks, heads)
    pdf.add_page(orientation="P")
    body = markdown.strip()
    if not re.match(r"^#\s", body):
        pdf.h1(title)
    if byline:
        pdf.p(byline, size=9)
    pdf.write_markdown(body, demote=demote)
    for data, caption in images or []:  # concept images: one per page, caption below
        try:
            pdf.image_page(data, plain(caption), images_title)
        except Exception:  # an unreadable image never loses the document
            continue
    return bytes(pdf.output())


def markdown_docx(title: str, markdown: str) -> bytes:
    """A simple .docx of the Markdown (headings, bullets, tables, paragraphs)."""
    from docx import Document

    doc = Document()
    if not re.match(r"^\s*#\s", markdown):  # the Markdown carries its own title otherwise
        doc.add_heading(title, level=0)
    lines = markdown.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|") and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            rows = [_cells(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_cells(lines[i]))
                i += 1
            n = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=n)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            doc.add_heading(plain(m.group(2)), level=min(len(m.group(1)), 4))
        elif _BULLET.match(line):
            style = "List Number" if re.match(r"^\s*\d", line) else "List Bullet"
            doc.add_paragraph(plain(_BULLET.match(line).group(2)), style=style)
        else:
            doc.add_paragraph(plain(line))
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
