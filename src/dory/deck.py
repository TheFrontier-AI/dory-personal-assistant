"""Slide decks: the outline the chat agent sends (`DeckSpec`) and its rendering to `.pptx`.

The agent's `create_deck` tool (dory.qa) validates the outline with these models, fetches chart
numbers from the ledger itself (the model never types numbers into a chart) and calls the
renderer. Everything here is pure: no network, no model calls, nothing written to disk.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from typing import Annotated, Literal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field, StringConstraints, field_validator, model_validator

from .i18n import t

MAX_TITLE_CHARS = 80
MAX_SUBTITLE_CHARS = 120
MAX_SLIDES = 12
MAX_BULLETS = 6
MAX_BULLET_CHARS = 140
MAX_NOTES_CHARS = 1000
MAX_TOP = 10
MAX_DAY_RANGE = 62  # days one day-by-day chart may span
MAX_DECK_BYTES = 2_000_000
MAX_SLUG_CHARS = 40

# Look: fonts (Thai goes in the complex-script slot), palette, 16:9 geometry.
LATIN_FONT = "Arial"
THAI_FONT = "Leelawadee UI"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x2E, 0x86, 0xDE)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xC9, 0xD3, 0xE3)
CHART_COLORS = [
    ACCENT,
    RGBColor(0xF5, 0x9E, 0x0B),
    RGBColor(0x10, 0xB9, 0x81),
    RGBColor(0xEF, 0x44, 0x44),
    RGBColor(0x8B, 0x5C, 0xF6),
    RGBColor(0x64, 0x74, 0x8B),
]
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.8)
BODY_TOP, BODY_H = Inches(1.8), Inches(4.9)
THAI_MONTHS = [
    "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
]  # fmt: skip
EN_MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]  # fmt: skip

Bullet = Annotated[str, StringConstraints(max_length=MAX_BULLET_CHARS)]


@dataclass
class ChartData:
    """What one chart slide plots: category labels and one value list per currency."""

    categories: list[str] = field(default_factory=list)
    series: dict[str, list[float]] = field(default_factory=dict)  # currency -> values

    @property
    def empty(self) -> bool:
        return not self.categories


class ChartQuery(BaseModel):
    """Which ledger numbers a chart slide shows; the numbers themselves are fetched in code."""

    group_by: Literal["category", "vendor", "submitter", "day"]
    start_date: str
    end_date: str
    style: Literal["bar", "pie", "line"] = "bar"
    top: int = Field(8, ge=1, le=MAX_TOP)

    @field_validator("start_date", "end_date")
    @classmethod
    def _iso(cls, v: str) -> str:
        try:
            return date.fromisoformat(v).isoformat()
        except ValueError as e:
            raise ValueError("must be an ISO date YYYY-MM-DD") from e

    @model_validator(mode="after")
    def _range(self) -> ChartQuery:
        a, b = date.fromisoformat(self.start_date), date.fromisoformat(self.end_date)
        if b < a:
            raise ValueError("end_date is before start_date")
        if self.group_by == "day" and (b - a).days + 1 > MAX_DAY_RANGE:
            raise ValueError(f"a day-by-day chart may span at most {MAX_DAY_RANGE} days")
        return self


class Slide(BaseModel):
    kind: Literal["section", "bullets", "two_column", "chart"]
    title: str = Field(max_length=MAX_TITLE_CHARS)
    bullets: list[Bullet] = Field(default_factory=list, max_length=MAX_BULLETS)
    right: list[Bullet] = Field(default_factory=list, max_length=MAX_BULLETS)
    chart: ChartQuery | None = None
    notes: str = Field("", max_length=MAX_NOTES_CHARS)

    @model_validator(mode="after")
    def _kind_fields(self) -> Slide:
        if self.kind == "chart" and self.chart is None:
            raise ValueError("chart is required for a chart slide")
        if self.kind in ("bullets", "two_column") and not self.bullets:
            raise ValueError("bullets needs at least one item")
        return self


class DeckSpec(BaseModel):
    title: str = Field(max_length=MAX_TITLE_CHARS)
    subtitle: str = Field("", max_length=MAX_SUBTITLE_CHARS)
    slides: list[Slide] = Field(min_length=1, max_length=MAX_SLIDES)


def deck_filename(title: str, today: date) -> str:
    """An ASCII slug of the title plus the date: "spending-oct-2026-20261016.pptx", or
    "dory-deck-20261016.pptx" when the title has no ASCII words (e.g. all Thai)."""
    ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    slug = ""
    for word in re.findall(r"[a-z0-9]+", ascii_title.lower()):
        if len(slug) + len(word) + (1 if slug else 0) > MAX_SLUG_CHARS:
            break
        slug = f"{slug}-{word}" if slug else word
    return f"{slug or 'dory-deck'}-{today:%Y%m%d}.pptx"


def chart_data(summary: dict, query: ChartQuery, lang: str) -> ChartData:
    """What a chart slide plots, from a ledger summary (`store.build_summary` shape: groups
    biggest first, or by date for `day`). Non-day charts keep the `top` groups and sum the rest
    per currency into "Other"; day charts keep every day, labelled MM-DD. One series per
    currency, THB first."""
    groups = list(summary.get("groups") or [])
    if not groups:
        return ChartData()
    day = query.group_by == "day"
    kept, rest = (groups, []) if day else (groups[: query.top], groups[query.top :])

    def by(g: dict) -> dict:
        return g.get("totals_by_currency") or {}

    def label(key) -> str:
        if key in (None, ""):
            return "—"
        return str(key)[5:10] if day else str(key)

    currencies = sorted({c for g in groups for c in by(g)}, key=lambda c: (c != "THB", c))
    categories = [label(g.get("key")) for g in kept]
    series = {c: [float(by(g).get(c, 0)) for g in kept] for c in currencies}
    if rest:
        categories.append(t("deck_other", lang))
        for c in currencies:
            series[c].append(round(sum(by(g).get(c, 0) for g in rest), 2))
    return ChartData(categories, series)


# ---------------------------------------------------------------------------- rendering
def clean_text(s: str, keep_newlines: bool = False) -> str:
    """Text safe for the slide XML: control characters dropped (python-pptx rejects them), line
    breaks turned into spaces unless `keep_newlines` (speaker notes), runs of spaces collapsed."""
    s = (s or "").replace("\r", "\n").replace("\t", " ")
    if not keep_newlines:
        s = s.replace("\n", " ")
    s = "".join(ch for ch in s if ch == "\n" or unicodedata.category(ch) not in ("Cc", "Cs"))
    return re.sub(r"[ ]{2,}", " ", s).strip()


def format_deck_date(d: date, lang: str) -> str:
    """ "16 October 2026" (en) or "16 ตุลาคม 2569" (th, Buddhist Era)."""
    if lang == "th":
        return f"{d.day} {THAI_MONTHS[d.month - 1]} {d.year + 543}"
    return f"{d.day} {EN_MONTHS[d.month - 1]} {d.year}"


def _thai_font(rpr) -> None:
    """Put the Thai font in the complex-script slot (`a:cs`, right after `a:latin`) of a run's
    or a chart's character properties; python-pptx does not expose it."""
    for old in rpr.findall(qn("a:cs")):
        rpr.remove(old)
    cs = OxmlElement("a:cs")
    cs.set("typeface", THAI_FONT)
    rpr.find(qn("a:latin")).addnext(cs)


def _run(paragraph, text: str, size, color, bold: bool = False, raw: bool = False):
    """One styled run. `raw` keeps fixed internal text as is (e.g. the "• " marker); everything
    else goes through clean_text."""
    run = paragraph.add_run()
    run.text = text if raw else clean_text(text)
    font = run.font
    font.size, font.bold, font.name = size, bold, LATIN_FONT
    font.color.rgb = color
    _thai_font(run._r.get_or_add_rPr())
    return run


def _textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    frame = slide.shapes.add_textbox(x, y, w, h).text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    return frame


def _rect(slide, x, y, w, h, color) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _background(slide, color) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _title_size(text: str, big: int) -> int:
    return big if len(text) <= 40 else big - 8


def _bullet_size(items: list[str]) -> int:
    chars = sum(len(x) for x in items)
    return 24 if chars <= 200 else 20 if chars <= 400 else 16


def _title_slide(prs, spec: DeckSpec, lang: str, today: date) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _background(slide, NAVY)
    _rect(slide, MARGIN, Inches(2.0), Inches(1.2), Inches(0.08), ACCENT)
    width = SLIDE_W - 2 * MARGIN
    frame = _textbox(slide, MARGIN, Inches(2.2), width, Inches(3.0))
    _run(frame.paragraphs[0], spec.title, Pt(_title_size(spec.title, 44)), WHITE, bold=True)
    if spec.subtitle:  # same box, so it follows a one- or two-line title closely
        sub = frame.add_paragraph()
        sub.space_before = Pt(14)
        _run(sub, spec.subtitle, Pt(24), LIGHT)
    foot = _textbox(slide, MARGIN, Inches(6.4), width, Inches(0.5))
    _run(
        foot.paragraphs[0],
        t("deck_footer", lang, date=format_deck_date(today, lang)),
        Pt(14),
        LIGHT,
    )


def _page_number(slide, n: int, color=MUTED) -> None:
    frame = _textbox(slide, SLIDE_W - MARGIN - Inches(1), Inches(6.85), Inches(1), Inches(0.4))
    frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    _run(frame.paragraphs[0], str(n), Pt(12), color)


def _bullet_list(slide, items: list[str], x, w, size: int) -> None:
    frame = _textbox(slide, x, BODY_TOP, w, BODY_H)
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(10)
        _run(p, "• ", Pt(size), ACCENT, bold=True, raw=True)
        _run(p, item, Pt(size), INK)


CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def _chart(slide, query: ChartQuery, data: ChartData | None, lang: str) -> None:
    """A native (editable) chart in the body area, or the no-data text when there is nothing."""
    x, w = MARGIN, SLIDE_W - 2 * MARGIN
    if data is None or data.empty:
        frame = _textbox(slide, x, BODY_TOP, w, BODY_H, MSO_ANCHOR.MIDDLE)
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        _run(frame.paragraphs[0], t("deck_no_data", lang), Pt(24), MUTED)
        return
    currencies = list(data.series)[:1] if query.style == "pie" else list(data.series)
    cd = CategoryChartData(number_format="#,##0")
    cd.categories = data.categories
    for cur in currencies:
        cd.add_series(cur, data.series[cur])
    chart = slide.shapes.add_chart(CHART_TYPES[query.style], x, BODY_TOP, w, BODY_H, cd).chart
    chart.font.size, chart.font.name = Pt(14), LATIN_FONT
    _thai_font(chart.font._rPr)
    chart.has_legend = len(currencies) > 1 or query.style == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "#,##0"
    plot.data_labels.number_format_is_linked = False
    if query.style == "pie":
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = CHART_COLORS[i % len(CHART_COLORS)]
        return
    for i, series in enumerate(plot.series):
        color = CHART_COLORS[i % len(CHART_COLORS)]
        if query.style == "line":
            series.format.line.color.rgb = color
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color


def _content_slide(prs, s: Slide, n: int, chart: ChartData | None, lang: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if s.kind == "section":
        _background(slide, ACCENT)
        frame = _textbox(
            slide, MARGIN, Inches(2.5), SLIDE_W - 2 * MARGIN, Inches(2.5), MSO_ANCHOR.MIDDLE
        )
        frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        _run(frame.paragraphs[0], s.title, Pt(_title_size(s.title, 44)), WHITE, bold=True)
    else:
        _rect(slide, 0, 0, Inches(0.25), SLIDE_H, ACCENT)
        title = _textbox(slide, MARGIN, Inches(0.5), SLIDE_W - 2 * MARGIN, Inches(1.1))
        _run(title.paragraphs[0], s.title, Pt(_title_size(s.title, 36)), INK, bold=True)
        width = SLIDE_W - 2 * MARGIN
        size = _bullet_size([*s.bullets, *s.right])
        if s.kind == "bullets":
            _bullet_list(slide, s.bullets, MARGIN, width, size)
        elif s.kind == "two_column":
            col = (width - Inches(0.5)) / 2
            _bullet_list(slide, s.bullets, MARGIN, col, size)
            _bullet_list(slide, s.right, MARGIN + col + Inches(0.5), col, size)
        elif s.kind == "chart" and s.chart is not None:
            _chart(slide, s.chart, chart, lang)
    if s.notes:
        slide.notes_slide.notes_text_frame.text = clean_text(s.notes, keep_newlines=True)
    _page_number(slide, n, LIGHT if s.kind == "section" else MUTED)
    return slide


def render_deck(spec: DeckSpec, charts: dict[int, ChartData], *, lang: str, today: date) -> bytes:
    """The deck as `.pptx` bytes: a title slide, then one slide per outline slide. `charts` maps
    the index of a chart slide to its data (fetched from the ledger by the caller)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _title_slide(prs, spec, lang, today)
    for i, s in enumerate(spec.slides):
        _content_slide(prs, s, i + 2, charts.get(i), lang)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
