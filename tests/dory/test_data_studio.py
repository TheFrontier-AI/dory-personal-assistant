"""Data studio logic (`dory.data_studio`): loading, types, profile, SQL, charts, stored datasets."""

from __future__ import annotations

import io
import time
import warnings
from datetime import UTC, datetime

import pytest
from fake_firestore import FakeFirestore
from openpyxl import Workbook, load_workbook
from PIL import Image

from dory import data_studio as ds
from dory.coo_store import FirestoreDocStore, SQLiteDocStore
from dory.i18n import CATALOG, EXTRA, t


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


SALES = [
    ["รายงานยอดขายประจำเดือน"],
    [None, None, "ปี 2568"],
    ["วันที่", "สินค้า", "จำนวน", "ราคา", "ส่วนลด"],
    ["15/10/2568", "กาแฟ", "1,200", "฿1,200.50", "12%"],
    ["16/10/2568", "ชา", "30", "฿20", "-"],
    ["17/10/2568", "กาแฟ", "5", "(40)", "0%"],
    ["รวม", None, "1,235", None, None],
]


def sales_ds() -> ds.Dataset:
    return ds.load_tables(xlsx({"ยอดขาย Q3": SALES}), "sales.xlsx")


# ------------------------------------------------------------------ loading
def test_header_row_is_found_below_title_rows_and_names_are_slugged():
    d = sales_ds()
    (tb,) = d.tables
    assert (tb.name, tb.label) == ("t_q3", "ยอดขาย Q3")
    assert [c.name for c in tb.columns] == ["col_1", "col_2", "col_3", "col_4", "col_5"]
    assert [c.label for c in tb.columns] == ["วันที่", "สินค้า", "จำนวน", "ราคา", "ส่วนลด"]
    assert tb.rows == 4  # three sales + the total row (kept, flagged)


def test_ascii_headers_become_slugs_duplicates_and_reserved_words_are_safe():
    csv = b"Order Date,Total ($),select,Total ($),2025\n2025-01-02,5,a,6,1\n"
    d = ds.load_tables(csv, "x.csv")
    names = [c.name for c in d.tables[0].columns]
    assert names == ["order_date", "total", "select_", "total_2", "c_2025"]
    prof = ds.profile(d)
    assert {"code": "duplicate_header", "column": "Total ($)"} in prof["tables"][0]["issues"]


def test_thai_numbers_currency_percent_and_buddhist_dates():
    (tb,) = sales_ds().tables
    types = {c.label: c.type for c in tb.columns}
    assert types == {
        "วันที่": "date",
        "สินค้า": "text",
        "จำนวน": "integer",
        "ราคา": "number",
        "ส่วนลด": "number",
    }
    f = tb.frame
    assert list(f["col_3"]) == [1200, 30, 5, 1235]
    assert list(f["col_4"][:3]) == [1200.5, 20.0, -40.0]  # ฿, and (40) is negative
    assert f["col_5"][0] == pytest.approx(0.12) and f["col_5"][2] == 0 and f["col_5"].isna()[1]
    assert tb.columns[4].unit == "%"
    assert str(f["col_1"][0].date()) == "2025-10-15"  # 2568 - 543
    assert str(f["col_1"][2].date()) == "2025-10-17"


@pytest.mark.parametrize(
    "text,iso",
    [
        ("15/10/2568", "2025-10-15"),
        ("2568-10-15", "2025-10-15"),
        ("15-10-2025", "2025-10-15"),
        ("15.10.2025", "2025-10-15"),
        ("03/04/2025", "2025-04-03"),  # day first
        ("15 ต.ค. 2568", "2025-10-15"),
        ("15 Oct 2025", "2025-10-15"),
        ("15/10/2568 14:30", "2025-10-15"),
    ],
)
def test_date_rules(text, iso):
    assert ds.parse_date(text).date().isoformat() == iso


@pytest.mark.parametrize("text", ["32/13/2025", "hello", "2025", ""])
def test_not_a_date(text):
    assert ds.parse_date(text) is None


@pytest.mark.parametrize(
    "text,val",
    [
        ("1,234.50", (1234.5, False)),
        ("฿1,200", (1200.0, False)),
        ("1 200 บาท", (1200.0, False)),
        ("12%", (0.12, True)),
        ("(1,200)", (-1200.0, False)),
        ("-5", (-5.0, False)),
        ("1e3", (1000.0, False)),
    ],
)
def test_number_rules(text, val):
    assert ds.parse_number(text) == pytest.approx(val)


def test_not_a_number_and_nulls():
    assert ds.parse_number("abc") is None and ds.parse_number("1,23") is None
    assert all(ds.is_null(v) for v in ["", "-", " – ", "N/A", None, float("nan"), "null"])
    assert not ds.is_null(0) and not ds.is_null("0")


def test_ids_with_leading_zeros_and_long_digit_strings_stay_text():
    d = ds.load_tables(b"id,phone,n\n007,0812345678,1\n010,0899999999,2\n", "x.csv")
    assert [c.type for c in d.tables[0].columns] == ["text", "text", "integer"]
    assert d.tables[0].frame["id"][0] == "007"
    d2 = ds.load_tables(b"card,n\n1234567890123,1\n1234567890124,2\n", "x.csv")
    assert d2.tables[0].columns[0].type == "text"


def test_mixed_types_become_null_and_are_reported():
    rows = "\n".join(["v,k"] + [f"{i},a" for i in range(20)] + ["oops,b"])
    d = ds.load_tables(rows.encode(), "m.csv")
    c = d.tables[0].columns[0]
    assert (c.type, c.bad) == ("integer", 1)
    issues = ds.profile(d)["tables"][0]["issues"]
    assert {"code": "mixed_types", "column": "v", "n": 1} in issues


def test_empty_columns_and_total_rows_are_flagged():
    d = ds.load_tables(b"name,amt,note\nA,10,\nB,20,\nTotal,30,\n", "x.csv")
    issues = ds.profile(d)["tables"][0]["issues"]
    assert {"code": "empty_column", "column": "note"} in issues
    assert {"code": "total_row", "rows": [3]} in issues


def test_csv_encodings_and_delimiters():
    text = "ชื่อ;ยอด\nกาแฟ;10\nชา;20\n"
    for enc in ("utf-8", "utf-8-sig", "cp874"):
        d = ds.load_tables(text.encode(enc), "t.csv")
        assert d.tables[0].columns[0].label == "ชื่อ" and d.tables[0].rows == 2, enc
    tsv = ds.load_tables(b"a\tb\n1\t2\n3\t4\n", "t.tsv")
    assert [c.name for c in tsv.tables[0].columns] == ["a", "b"] and tsv.tables[0].rows == 2


def test_json_array_of_objects():
    data = b'[{"name":"A","qty":2,"tags":["x"]},{"name":"B","qty":3,"extra":1}]'
    (tb,) = ds.load_tables(data, "items.json").tables
    assert [c.name for c in tb.columns] == ["name", "qty", "tags", "extra"]
    assert tb.rows == 2 and tb.columns[1].type == "integer"
    for bad in (b"not json", b"[1, 2]", b"[]"):
        with pytest.raises(ds.DataError):
            ds.load_tables(bad, "x.json")


def test_unsupported_empty_corrupt_and_too_big():
    with pytest.raises(ds.DataError, match="unsupported"):
        ds.load_tables(b"x", "a.pdf")
    with pytest.raises(ds.DataError, match="no data"):
        ds.load_tables(xlsx({"Empty": [[None]]}), "e.xlsx")
    with pytest.raises(ds.DataError, match="could not be opened"):
        ds.load_tables(b"not a workbook", "e.xlsx")
    with pytest.raises(ds.DataError, match="50 MB"):
        ds.load_tables(b"x" * (ds.MAX_FILE_BYTES + 1), "big.csv")


def test_cell_cap_loads_the_first_rows_and_says_so(monkeypatch):
    monkeypatch.setattr(ds, "MAX_CELLS", 300)
    body = "a,b,c\n" + "\n".join(f"{i},{i},{i}" for i in range(500))
    d = ds.load_tables(body.encode(), "big.csv")
    assert d.tables[0].rows == 100 and d.tables[0].truncated == 100
    prof = ds.profile(d, "en")
    assert any("first 100 rows" in n for n in prof["notes"]) and prof["tables"][0]["truncated"]
    assert "100" in ds.profile_md(d, "th")


def test_table_cap_and_sheet_name_slugs(monkeypatch):
    sheets = {f"Sheet {i}": [["a", "b"], [1, 2]] for i in range(25)}
    d = ds.load_tables(xlsx(sheets), "many.xlsx")
    assert len(d.tables) == ds.MAX_TABLES == 20
    assert d.tables[0].name == "t_sheet_0" and {"code": "tables", "n": 20} in d.notes
    thai = ds.load_tables(xlsx({"ยอดขาย": [["a", "b"], [1, 2]], "ต้นทุน": [["a", "b"], [1, 2]]}), "t.xlsx")  # fmt: skip
    assert [x.name for x in thai.tables] == ["t_1", "t_2"]
    assert [x.label for x in thai.tables] == ["ยอดขาย", "ต้นทุน"]


def test_no_header_row_gets_generated_names():
    d = ds.load_tables(b"1,2\n3,4\n5,6\n", "n.csv")
    assert [c.name for c in d.tables[0].columns] == ["col_1", "col_2"]
    assert d.tables[0].rows == 3 and {"code": "noheader", "table": "t_n"} in d.notes


def test_quick_counts():
    assert ds.quick_counts(xlsx({"a": [[1, 2], [3, 4]], "b": [[1]]}), "x.xlsx") == (2, 3)
    assert ds.quick_counts(b"a,b\n1,2\n", "x.csv")[0] == 1
    with pytest.raises(ds.DataError):
        ds.quick_counts(b"junk", "x.xlsx")


def test_legacy_xls_is_read_through_xlrd():
    import xlrd  # noqa: F401  (tiny dependency; a missing one must fail loudly)

    with pytest.raises(ds.DataError, match=r"\.xls"):
        ds.load_tables(b"not an xls", "old.xls")


# ------------------------------------------------------------------ profile
def test_profile_shape():
    prof = ds.profile(sales_ds(), "en")
    (tp,) = prof["tables"]
    assert tp["rows"] == 4 and len(tp["sample"]) == 4 and len(tp["columns"]) == 5
    qty = next(c for c in tp["columns"] if c["name"] == "col_3")
    assert (qty["min"], qty["max"], qty["type"]) == (5, 1235, "integer")
    assert qty["mean"] == 617.5 and qty["null_pct"] == 0.0
    date_col = tp["columns"][0]
    assert (date_col["min"], date_col["max"]) == ("2025-10-15", "2025-10-17")
    prod = tp["columns"][1]
    assert prod["top"][0] == ["กาแฟ", 2] and prod["null_pct"] == 25.0
    assert tp["sample"][0]["col_1"] == "2025-10-15"
    assert {"code": "total_row", "rows": [4]} in tp["issues"]


def test_profile_md_is_short_and_bilingual():
    wide = [[f"c{i}" for i in range(60)], list(range(60))]
    d = ds.load_tables(xlsx({f"s{i}": wide for i in range(6)}), "wide.xlsx")
    for lang in ("th", "en"):
        md = ds.profile_md(d, lang)
        assert 0 < len(md.splitlines()) <= ds.PROFILE_LINES, lang
        assert "t_s0" in md
    th, en = ds.profile_md(sales_ds(), "th"), ds.profile_md(sales_ds(), "en")
    assert "แถว" in th and "rows" in en and "`col_3` (จำนวน)" in th
    assert "likely total row" in en and "แถวที่น่าจะเป็นผลรวม" in th


# ------------------------------------------------------------------ sql
def test_sql_happy_path_over_a_sheet_and_across_sheets():
    d = ds.load_tables(
        xlsx(
            {
                "Sales": [["item", "qty"], ["a", 1], ["b", 2], ["a", 5]],
                "Prices": [["item", "price"], ["a", 10], ["b", 20]],
            }
        ),
        "x.xlsx",
    )
    r = ds.run_sql(d, "SELECT item, SUM(qty) AS q FROM t_sales GROUP BY item ORDER BY item")
    assert r.columns == ["item", "q"] and r.rows == [["a", 6], ["b", 2]] and not r.truncated
    r = ds.run_sql(
        d,
        "WITH s AS (SELECT item, SUM(qty) q FROM t_sales GROUP BY 1) "
        "SELECT s.item, s.q * p.price AS value FROM s JOIN t_prices p USING (item) ORDER BY 1;",
    )
    assert r.rows == [["a", 60], ["b", 40]]
    assert ds.run_sql(d, "select 'drop table x; create' as note").rows == [["drop table x; create"]]
    assert ds.run_sql(d, "select 1 -- comment\n").rows == [[1]]


@pytest.mark.parametrize(
    "q",
    [
        "DROP TABLE t_sales",
        "INSERT INTO t_sales VALUES ('z', 1)",
        "UPDATE t_sales SET qty = 0",
        "DELETE FROM t_sales",
        "CREATE TABLE x AS SELECT 1",
        "SELECT 1; SELECT 2",
        "SELECT 1; DROP TABLE t_sales",
        "ATTACH 'x.db'",
        "COPY t_sales TO '/tmp/x.csv'",
        "INSTALL httpfs",
        "LOAD httpfs",
        "PRAGMA database_list",
        "EXPORT DATABASE '/tmp/x'",
        "SELECT * FROM read_csv('/etc/passwd')",
        "SELECT * FROM glob('/*')",
        "",
        "   ;  ",
        "SET memory_limit='10GB'",
        "-- just a comment",
    ],
)
def test_sql_rejects_everything_but_one_select(q):
    d = ds.load_tables(xlsx({"Sales": [["item", "qty"], ["a", 1]]}), "x.xlsx")
    with pytest.raises(ds.SqlError):
        ds.run_sql(d, q)
    assert ds.run_sql(d, "SELECT COUNT(*) FROM t_sales").rows == [[1]]  # nothing was changed


def test_sql_cannot_read_files_even_with_a_plain_select():
    d = ds.load_tables(b"a,b\n1,2\n", "x.csv")
    with pytest.raises(ds.SqlError):
        ds.run_sql(d, "SELECT * FROM '/etc/hostname'")


def test_unknown_column_and_table_hints_list_what_exists():
    d = ds.load_tables(b"price,qty\n1,2\n", "x.csv")
    with pytest.raises(ds.SqlError) as e:
        ds.run_sql(d, "SELECT nope FROM t_x")
    assert e.value.kind == "unknown_column" and "price, qty" in str(e.value)
    with pytest.raises(ds.SqlError) as e:
        ds.run_sql(d, "SELECT * FROM missing_table", lang="th")
    assert e.value.kind == "unknown_table" and "t_x" in str(e.value) and "ไม่พบตาราง" in str(e.value)
    with pytest.raises(ds.SqlError) as e:
        ds.run_sql(d, "SELECT FROM WHERE")
    assert e.value.kind == "sql"


def test_limit_truncation_and_column_cap():
    d = ds.load_tables(b"a\n1\n", "x.csv")
    r = ds.run_sql(d, "SELECT range AS n FROM range(1000)", limit=50)
    assert len(r.rows) == 50 and r.truncated
    r = ds.run_sql(d, "SELECT range AS n FROM range(50)", limit=50)
    assert len(r.rows) == 50 and not r.truncated
    wide = "SELECT " + ", ".join(f"{i} AS c{i}" for i in range(60))
    r = ds.run_sql(d, wide)
    assert len(r.columns) == ds.MAX_SQL_COLS and r.cols_cut
    dup = ds.run_sql(d, "SELECT 1 AS a, 2 AS a")
    assert dup.columns == ["a", "a_1"]


def test_a_slow_query_is_interrupted_at_the_timeout():
    d = ds.load_tables(b"a\n1\n", "x.csv")
    started = time.monotonic()
    with pytest.raises(ds.SqlError) as e:
        ds.run_sql(
            d,
            "SELECT COUNT(*) FROM range(100000000000) a, range(1000000) b",
            timeout=0.4,
        )
    assert e.value.kind == "timeout" and "0.4" in str(e.value)
    assert time.monotonic() - started < 6


def test_result_md_table_rows_notes_and_thai_alignment():
    d = ds.load_tables(b"a\n1\n", "x.csv")
    r = ds.run_sql(d, "SELECT range AS n, 'กาแฟ' AS ชื่อ FROM range(40)", limit=30)
    md = ds.result_md(r, "en")
    lines = md.splitlines()
    assert lines[0] == "```" and lines[-3] == "```" == False or "```" in md
    body = [ln for ln in lines if " | " in ln]
    assert len(body) == 1 + ds.SHOW_ROWS  # header + 25 rows
    assert "5 more rows" in md and "capped at 30" in md
    assert "และอีก" in ds.result_md(r, "th")
    assert ds.result_md(ds.run_sql(d, "SELECT 1 WHERE false"), "en") == t("data_sql_empty", "en")
    assert ds.fmt_value(1234567) == "1,234,567" and ds.fmt_value(1234.5) == "1,234.5"


def test_result_xlsx_and_csv_round_trip():
    d = sales_ds()
    r = ds.run_sql(d, "SELECT col_1, col_2, col_3 FROM t_q3 WHERE col_2 IS NOT NULL ORDER BY 1")
    wb = load_workbook(io.BytesIO(ds.result_xlsx(r)))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("col_1", "col_2", "col_3")
    assert rows[1][1:] == ("กาแฟ", 1200) and rows[1][0].year == 2025
    csv = ds.result_csv(r).decode("utf-8-sig").splitlines()
    assert csv[0] == "col_1,col_2,col_3" and "กาแฟ" in csv[1]
    empty = ds.run_sql(d, "SELECT col_2 FROM t_q3 WHERE false")
    assert ds.result_xlsx(empty)  # header only


# ------------------------------------------------------------------ charts
def png_info(data: bytes) -> tuple[str, tuple[int, int]]:
    im = Image.open(io.BytesIO(data))
    return im.format, im.size


@pytest.mark.parametrize("kind", ds.KINDS)
def test_every_chart_kind_makes_a_png_with_thai_labels_and_no_missing_glyphs(kind):
    d = ds.load_tables(
        "เดือน,ยอดขาย,ต้นทุน\nมกราคม,1200000,800000\nกุมภาพันธ์,1500000,900000\nมีนาคม,900000,700000\n".encode(),  # fmt: skip
        "t.csv",
    )
    r = ds.run_sql(d, "SELECT col_1 AS เดือน, col_2 AS ยอดขาย, col_3 AS ต้นทุน FROM t_t")
    ycols = {"scatter": ["ต้นทุน"], "hist": ["ยอดขาย"], "pie": ["ยอดขาย"]}.get(kind, ["ยอดขาย", "ต้นทุน"])  # fmt: skip
    x = "ยอดขาย" if kind == "scatter" else "เดือน"
    with warnings.catch_warnings():
        warnings.filterwarnings("error", message=".*[Gg]lyph.*missing")  # matplotlib: tofu
        png = ds.chart_png(r, kind, x, ycols, "ยอดขายรายเดือน (บาท)", "th")
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) <= ds.CHART_MAX_BYTES
    assert png_info(png) == ("PNG", (1200, 700))


def test_chart_defaults_pick_first_column_and_numeric_columns():
    df = ds.pd.DataFrame({"city": ["a", "b", "c"], "n": [1, 2, 3], "m": [3.5, 2.5, 1.5]})
    assert ds.chart_png(df, "bar", title="x", lang="en")[:4] == b"\x89PNG"
    assert ds.chart_png(df, "line", "city", ["n", "m"], lang="en")[:4] == b"\x89PNG"
    assert ds.chart_png(df, "hist", y="n", lang="en")[:4] == b"\x89PNG"
    dated = ds.pd.DataFrame({"d": ds.pd.to_datetime(["2025-01-01", "2025-02-01"]), "v": [1, 2]})
    assert ds.chart_png(dated, "line", "d", ["v"], lang="en")[:4] == b"\x89PNG"


def test_cjk_and_devanagari_labels_use_script_fonts():
    from dory.pdf_report import find_fonts

    df = ds.pd.DataFrame({"k": ["销售", "利润", "विक्रय"], "v": [3, 2, 1]})
    needs = ds.scripts_in(df["k"])
    assert needs == ["zh", "hi"]
    if not all(find_fonts("", n).script == n for n in needs):
        pytest.skip("CJK / Devanagari fonts are not installed here")
    with warnings.catch_warnings():
        warnings.filterwarnings("error", message=".*[Gg]lyph.*missing")
        assert ds.chart_png(df, "bar", "k", ["v"], "销售 विक्रय", "th")[:4] == b"\x89PNG"
    paths = ds.chart_font_paths(df["k"], "th")
    assert len(paths) == 3  # Thai/Latin, then zh, then hi


def test_category_cap_keeps_the_top_29_and_sums_the_rest_as_other():
    df = ds.pd.DataFrame({"k": [f"c{i}" for i in range(45)], "v": list(range(45))})
    out = ds.cap_categories(df, "k", ["v"], "Other")
    assert len(out) == ds.MAX_CATEGORIES == 30
    assert out["k"].iloc[-1] == "Other" and out["v"].iloc[-1] == sum(range(16))  # 0..15 left out
    assert out["k"].iloc[0] == "c44" and "c16" in set(out["k"]) and "c15" not in set(out["k"])
    small = df.head(10)
    assert ds.cap_categories(small, "k", ["v"], "Other") is small
    assert ds.chart_png(df, "bar", "k", ["v"], lang="en")[:4] == b"\x89PNG"
    assert ds.chart_png(df, "pie", "k", ["v"], lang="th")[:4] == b"\x89PNG"
    assert t("data_other", "th") == "อื่นๆ" and t("data_other", "en") == "Other"


def test_chart_errors_are_reasons():
    df = ds.pd.DataFrame({"k": ["a", "b"], "v": [1, -2], "s": ["x", "y"]})
    for kw, msg in [
        (dict(kind="donut"), "Chart kinds"),
        (dict(kind="bar", x="zzz"), "zzz"),
        (dict(kind="bar", y=["zzz"]), "k, v, s"),
        (dict(kind="bar", y=["s"]), "numeric"),
        (dict(kind="pie", y=["v"]), "non-negative"),
    ]:
        with pytest.raises(ds.DataError, match=msg):
            ds.chart_png(df, lang="en", **kw)
    with pytest.raises(ds.DataError, match="no data"):
        ds.chart_png(df.iloc[0:0], "bar", lang="en")


def test_chart_png_stays_under_one_megabyte_for_dense_data():
    import random

    rnd = random.Random(1)
    df = ds.pd.DataFrame({"x": [rnd.random() for _ in range(20000)], "y": [rnd.random() for _ in range(20000)]})  # fmt: skip
    png = ds.chart_png(df, "scatter", "x", ["y"], lang="en")
    assert len(png) <= ds.CHART_MAX_BYTES


# ------------------------------------------------------------------ stored datasets
@pytest.fixture(params=["sqlite", "firestore"])
def docs(request):
    if request.param == "sqlite":
        return SQLiteDocStore(":memory:")
    return FirestoreDocStore(client=FakeFirestore())


def repo_of(docs):
    return ds.DatasetRepo(docs, clock=lambda: datetime(2026, 10, 17, 9, 0, tzinfo=UTC))


def test_save_and_load_round_trip_keeps_types_nulls_and_text_ids(docs):
    repo = repo_of(docs)
    csv = 'รหัส,วันที่,ยอด,ส่วนลด,ใช่\n007,15/10/2568,"1,200",12%,\n010,16/10/2568,30,-,\n'.encode()
    d = ds.load_tables(csv, "orders.csv")
    rec = repo.save(7, d, name="Orders", owner_id=1, owner_name="A", channel_id=10)
    assert rec["id"].startswith("ds-") and len(rec["id"]) == 9 and rec["pinned"] is False
    back = repo.load(7, repo.get(7, rec["id"]))
    a, b = d.tables[0], back.tables[0]
    assert [c.dump() for c in a.columns] == [c.dump() for c in b.columns]
    assert list(b.frame["col_1"]) == ["007", "010"] and b.frame["col_2"][0].year == 2025
    assert list(b.frame["col_3"]) == [1200, 30] and b.frame["col_4"].isna()[1]
    r = ds.run_sql(back, "SELECT SUM(col_3) FROM t_orders")
    assert r.rows == [[1230]]


def test_big_tables_are_split_into_small_documents(docs, monkeypatch):
    monkeypatch.setattr(ds, "CHUNK_CHARS", 200)
    repo = repo_of(docs)
    import random

    rnd = random.Random(3)
    body = "a,b\n" + "\n".join(f"{rnd.random()},{rnd.random()}" for _ in range(400))
    d = ds.load_tables(body.encode(), "r.csv")
    rec = repo.save(7, d, name="r", owner_id=1)
    n = rec["tables"][0]["chunks"]
    assert n > 3
    assert len(docs.find(ds.DATASET_DATA, 7)) == n
    assert repo.load(7, rec).tables[0].rows == 400
    repo.delete(7, rec)
    assert docs.find(ds.DATASET_DATA, 7) == [] and repo.all(7) == []


def test_too_big_dataset_is_refused_not_half_saved(docs, monkeypatch):
    monkeypatch.setattr(ds, "MAX_STORED_BYTES", 100)
    repo = repo_of(docs)
    import random

    rnd = random.Random(2)
    d = ds.load_tables(("a\n" + "\n".join(str(rnd.random()) for _ in range(300))).encode(), "r.csv")
    with pytest.raises(ds.DataError, match="3 MB"):
        repo.save(7, d, name="r", owner_id=1)
    assert docs.find(ds.DATASETS, 7) == [] and docs.find(ds.DATASET_DATA, 7) == []


def test_prune_keeps_30_per_guild_oldest_first_and_never_pinned(docs):
    now = [datetime(2026, 10, 1, tzinfo=UTC)]

    def tick():
        now[0] = now[0].replace(day=now[0].day, minute=now[0].minute) + (datetime(2026, 1, 1, 0, 1) - datetime(2026, 1, 1))  # fmt: skip
        return now[0]

    repo = ds.DatasetRepo(docs, clock=tick)
    d = ds.load_tables(b"a\n1\n", "x.csv")
    first = repo.save(7, d, name="n0", owner_id=1)
    docs.put(ds.DATASETS, 7, first["id"], {**first, "pinned": True})
    ids = [first["id"]]
    for i in range(1, 33):
        ids.append(repo.save(7, d, name=f"n{i}", owner_id=1)["id"])
    left = {r["id"] for r in repo.all(7)}
    assert len(left) == ds.MAX_DATASETS == 30
    assert ids[0] in left  # pinned survives although it is the oldest
    assert ids[1] not in left and ids[2] not in left and ids[3] not in left  # oldest unpinned went
    assert ids[-1] in left
    other = repo.save(8, d, name="other guild", owner_id=1)
    assert len(repo.all(7)) == 30 and [r["id"] for r in repo.all(8)] == [other["id"]]


def test_find_by_id_name_or_file_fragment(docs):
    repo = repo_of(docs)
    d = ds.load_tables(b"a\n1\n", "Q3-sales.csv")
    r1 = repo.save(7, d, name="Sales Q3", owner_id=1)
    repo.save(7, d, name="Costs", owner_id=1)
    assert [r["id"] for r in repo.find(7, "sales q3")] == [r1["id"]]
    assert len(repo.find(7, "q3-sales")) == 2  # both came from that file
    assert [r["id"] for r in repo.find(7, r1["id"])] == [r1["id"]]
    assert repo.find(7, "") == [] and repo.find(8, "sales") == []


def test_visibility_channel_readers_audience_and_private_datasets():
    chan = {"id": "x", "owner_id": 1, "channel_id": 10}
    can = lambda cid: cid == 10  # noqa: E731
    # a channel dataset: whoever (owner included) can read the channel and so can the audience
    assert ds.can_read(chan, 2, can) and ds.can_read(chan, 1, can)
    assert not ds.can_read(chan, 2, lambda _c: False) and not ds.can_read(chan, 1, lambda _c: False)  # fmt: skip
    assert not ds.can_read({**chan, "channel_id": 11}, 2, can)
    # a DM dataset: the owner, in a private conversation only
    dm = {"id": "y", "owner_id": 1, "channel_id": 0}
    assert ds.can_read(dm, 1, can, private=True)
    assert not ds.can_read(dm, 1, can, private=False)
    assert not ds.can_read(dm, 2, lambda _c: True, private=True)


# ------------------------------------------------------------------ strings and playbook
def test_playbook_is_a_short_english_file_with_a_thai_header():
    text = ds.playbook()
    assert text.count("\n") <= 62 and "ทักษะวิเคราะห์ข้อมูล" in text
    for needle in ("data_load", "data_sql", "data_chart", "median", "data_save"):
        assert needle in text


def test_every_data_string_has_thai_and_english_and_matching_placeholders():
    import re

    ph = re.compile(r"\{(\w+)\}")
    assert ds.STRINGS
    for key, v in ds.STRINGS.items():
        assert CATALOG[key] is v
        assert v["th"] and v["en"], key
        assert set(ph.findall(v["th"])) == set(ph.findall(v["en"])), key
        t(key, "th", **dict.fromkeys(ph.findall(v["en"]), "x"))
        assert t(key, "ja") is not None and EXTRA  # second tier falls back, never raises
