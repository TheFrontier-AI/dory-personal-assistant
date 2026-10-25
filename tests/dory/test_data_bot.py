# ruff: noqa: F811, F401
"""Data studio tools and glue (`dory.data_bot`): attachments, per-request state, saved datasets and
who may see them, voice parity, the prompt playbook and the feature flag."""

from __future__ import annotations

import io
import json
from pathlib import Path
from types import SimpleNamespace

from conftest import GUILD, fc_resp, scripted, text_resp
from openpyxl import Workbook, load_workbook
from test_coordinator import BOSS, INTERN, OWNER_ID, PEER, STAFF, world  # noqa: F401
from test_media import matt
from test_modes import TMsg, bot  # noqa: F401

from dory import data_studio as ds
from dory import qa
from dory.data_bot import DataCtx, data_tools
from dory.guild_settings import FEATURES, Features, GuildSettingsUpdate
from dory.qa import FileInput, tool_files
from dory.voice_actions import ACTION_TOOL_NAMES

ALL = {"data_load", "data_sql", "data_chart", "data_export", "data_save", "data_list", "data_forget"}  # fmt: skip


def xlsx_bytes(rows=None, title="ยอดขาย") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows or [
        ["รายงานยอดขาย"],
        ["เดือน", "สินค้า", "ยอด"],
        ["ม.ค.", "กาแฟ", "1,200"],
        ["ก.พ.", "ชา", "800"],
        ["มี.ค.", "กาแฟ", "1,500"],
    ]:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def names(tools):
    return {f.__name__ for f in tools}


def tool(tools, name):
    return next(f for f in tools if f.__name__ == name)


def set_flag(bot, on: bool):
    bot.guilds_store.save_settings(GUILD, GuildSettingsUpdate(features={"data": on}), "x")
    bot.team._cache.clear()


async def inputs(bot, world, member, attachments, channel=None):
    msg = TMsg(world.g, channel or world.pub, member, "วิเคราะห์ให้หน่อย", attachments=attachments)
    return msg, await bot._file_inputs(msg)


async def tools_for(bot, world, member, attachments=(), channel=None, is_exec=False):
    msg, files = await inputs(bot, world, member, list(attachments), channel)
    return msg, bot._data_tools(msg, world.g, member, files, "th", is_exec)


# ------------------------------------------------------------------ the files the agent gets
async def test_xlsx_keeps_its_bytes_and_the_text_part_points_to_data_load(bot, world):
    data = xlsx_bytes()
    _msg, files = await inputs(bot, world, world.staff, [matt("sales.xlsx", "", data)])
    (f,) = files
    assert f.skip == "" and f.raw == data and f.mime == "text/plain"
    text = f.data.decode()
    assert "ยอดขาย" in text and text.endswith("(use data_load for analysis: 1 sheets, 5 rows)")
    parts, bad = qa.file_parts(files)
    assert bad == [] and "use data_load" in parts[0].text


async def test_csv_tsv_and_json_are_data_files_even_when_big(bot, world):
    big = ("a,b\n" + "\n".join(f"{i},{i * 2}" for i in range(60000))).encode()  # > 200 KB
    _m, files = await inputs(
        bot,
        world,
        world.staff,
        [
            matt("big.csv", "text/csv", big),
            matt("t.tsv", "", b"a\tb\n1\t2\n"),
            matt("j.json", "", b'[{"a":1}]'),
        ],  # fmt: skip
    )
    assert [f.skip for f in files] == ["", "", ""] and all(f.raw for f in files)
    assert files[0].raw == big and "use data_load" in files[0].data.decode()
    assert len(files[0].data) < 31_000  # the text part is still clipped


async def test_old_xls_that_cannot_be_opened_is_skipped_as_before(bot, world):
    _m, (f,) = await inputs(bot, world, world.staff, [matt("old.xls", "", b"x")])
    assert ".xls" in f.skip and not f.raw


async def test_flag_off_means_no_raw_bytes_and_no_hint_line(bot, world):
    set_flag(bot, False)
    _m, (f,) = await inputs(bot, world, world.staff, [matt("s.xlsx", "", xlsx_bytes())])
    assert f.raw == b"" and "data_load" not in f.data.decode()
    _m, (c,) = await inputs(bot, world, world.staff, [matt("s.csv", "text/csv", b"a,b\n1,2\n")])
    assert c.raw == b"" and c.mime == "text/plain"  # the plain text path as before


async def test_other_documents_are_untouched(bot, world):
    _m, (f,) = await inputs(bot, world, world.staff, [matt("n.txt", "text/plain", b"hello")])
    assert f.raw == b"" and f.data == b"hello"


# ------------------------------------------------------------------ the tools of one request
async def test_tools_chat_voice_flag_and_interjection(bot, world):
    m = world.staff
    msg, tools = await tools_for(bot, world, m)
    assert names(tools) == ALL
    assert all(f.outbox is tools[0].outbox for f in tools)
    set_flag(bot, False)
    assert bot._data_tools(msg, world.g, m, [], "th") == []
    set_flag(bot, True)
    kw = dict(setting="th", lang="th", question="x", memory=None, profile=None)
    assert names(bot._agent_kwargs(msg, world.g, m, False, files=[], **kw)["extra_tools"]) >= ALL
    inter = bot._agent_kwargs(msg, world.g, m, False, interjection=True, **kw)
    assert not ALL & names(inter.get("extra_tools", []))
    room = bot.room_message(world.g, m, world.pub, room={m.id})
    voice = names(bot._agent_kwargs(room, world.g, m, False, files=[], **kw)["extra_tools"])
    assert voice >= ALL
    # voice: the actions are re-bound to the speaker; reading tools are not
    assert {"data_chart", "data_export", "data_save", "data_forget"} <= ACTION_TOOL_NAMES
    assert not {"data_load", "data_sql", "data_list"} & ACTION_TOOL_NAMES
    assert "data" in FEATURES and Features().data is True


async def test_load_profile_sql_default_dataset_and_a_second_file(bot, world):
    m = world.staff
    first = xlsx_bytes()
    second = xlsx_bytes([["k", "v"], ["a", 1], ["b", 2]], title="Other")
    _msg, tools = await tools_for(
        bot, world, m, [matt("sales.xlsx", "", first), matt("other.xlsx", "", second)]
    )
    load, sql = tool(tools, "data_load"), tool(tools, "data_sql")
    assert sql("select 1")["status"] == "error"  # nothing loaded yet
    out = load()  # -1 = the last data file
    assert out["status"] == "ok" and out["tables"][0]["name"] == "t_other" and not out["saved"]
    assert out["dataset_id"].startswith("ds-") and "`k`" in out["profile_md"]
    assert out["sample"]["t_other"][0] == {"k": "a", "v": 1}
    # the default dataset is the one loaded last
    assert sql("SELECT SUM(v) AS s FROM t_other")["rows"] == [[3]]
    out1 = load(file_index=0)
    assert out1["tables"][0]["name"] == "t_1" and out1["notes"] == []
    r = sql("SELECT nope FROM t_1")
    assert r["status"] == "error" and "col_1, col_2, col_3" in r["error"]  # lists the real columns
    ok = sql("SELECT col_2 AS item, SUM(col_3) AS total FROM t_1 GROUP BY 1 ORDER BY 1")
    assert ok["rows"] == [["กาแฟ", 2700], ["ชา", 800]] and "```" in ok["result_md"]
    assert ok["columns"] == ["item", "total"] and ok["truncated"] is False
    # the first one is still there by name
    assert sql("SELECT COUNT(*) FROM t_other", dataset=out["dataset_id"])["rows"] == [[2]]
    assert sql("SELECT 1", dataset="nonexistent")["status"] == "error"
    assert load(file_index=9)["status"] == "error"


async def test_load_without_a_file_or_dataset_is_a_reason_not_an_exception(bot, world):
    _msg, tools = await tools_for(bot, world, world.staff)
    out = tool(tools, "data_load")()
    assert out["status"] == "error" and "attached" in out["error"]
    assert tool(tools, "data_load")(dataset="nothing")["status"] == "error"
    assert tool(tools, "data_sql")("DROP TABLE x")["status"] == "error"


async def test_corrupt_file_is_reported(bot, world):
    _msg, tools = await tools_for(bot, world, world.staff, [matt("x.csv", "text/csv", b"")])
    out = tool(tools, "data_load")()
    assert out["status"] == "error"


async def test_chart_and_export_are_attached_through_the_outbox(bot, world):
    _msg, tools = await tools_for(bot, world, world.staff, [matt("s.xlsx", "", xlsx_bytes())])
    tool(tools, "data_load")()
    out = tool(tools, "data_chart")(
        query="SELECT col_1 AS เดือน, col_3 AS ยอด FROM t_1 ORDER BY 1",
        kind="bar",
        x="เดือน",
        y=["ยอด"],
        title="ยอดขายรายเดือน",
    )
    assert out["status"] == "ok" and out["attached"] is True
    exp = tool(tools, "data_export")(query="SELECT * FROM t_1", file_format="xlsx", name="สรุป")
    assert exp["attached"] and "rows attached" in exp["note"]
    csv = tool(tools, "data_export")(
        query="SELECT * FROM t_1", file_format="csv", name="clean data"
    )
    files = tool_files(tools)
    assert [n for n, _ in files] == ["chart-data.png", "data.xlsx", "clean-data.csv"]
    assert files[0][1][:8] == b"\x89PNG\r\n\x1a\n" and files[0][0].endswith(".png")
    assert files[0][0].isascii() and files[1][0].isascii() and files[1][0].endswith(".xlsx")
    rows = list(load_workbook(io.BytesIO(files[1][1])).active.iter_rows(values_only=True))
    assert len(rows) == 4 and csv["attached"]  # header + 3 rows
    assert files[2][1].startswith(b"\xef\xbb\xbf")
    bad = tool(tools, "data_chart")(query="SELECT 1 AS a", kind="donut")
    assert bad["status"] == "error" and len(tool_files(tools)) == 3


async def test_save_list_load_later_and_forget_in_a_public_channel(bot, world):
    staff = world.staff
    _m, tools = await tools_for(bot, world, staff, [matt("s.xlsx", "", xlsx_bytes())])
    tool(tools, "data_load")()
    saved = tool(tools, "data_save")(name="Sales Q1")
    did = saved["dataset_id"]
    assert saved["status"] == "ok" and "channel" in saved["visible_to"]
    assert tool(tools, "data_save")()["note"] == "already saved"
    # a later request (no file) by someone else who can read the channel
    _m2, tools2 = await tools_for(bot, world, world.peer)
    listed = tool(tools2, "data_list")()["datasets"]
    assert [d["dataset_id"] for d in listed] == [did] and listed[0]["owner"] == "Somchai"
    out = tool(tools2, "data_load")(dataset="sales q1")
    assert out["saved"] is True and out["dataset_id"] == did
    sql = tool(tools2, "data_sql")
    assert sql("SELECT SUM(col_3) FROM t_1")["rows"] == [[3500]]
    # forget: not the peer, but the owner
    assert "only the owner" in tool(tools2, "data_forget")(did)["error"]
    assert tool(tools, "data_forget")(did)["deleted"] == "Sales Q1"
    assert tool(tools, "data_list")()["datasets"] == []
    _m3, tools3 = await tools_for(bot, world, world.peer)  # a new request: it is gone
    assert tool(tools3, "data_load")(dataset=did)["status"] == "error"


async def test_an_executive_can_forget_someone_elses_dataset(bot, world):
    _m, tools = await tools_for(bot, world, world.staff, [matt("s.csv", "text/csv", b"a,b\n1,2\n")])
    tool(tools, "data_load")()
    did = tool(tools, "data_save")()["dataset_id"]
    _m2, boss_tools = await tools_for(bot, world, world.boss, is_exec=True)
    assert tool(boss_tools, "data_forget")(did)["status"] == "ok"


async def test_a_dataset_from_a_hidden_channel_is_invisible_elsewhere(bot, world):
    # saved by the boss in the secret channel (readers: owner, admin, boss)
    _m, tools = await tools_for(
        bot, world, world.boss, [matt("pay.csv", "text/csv", b"who,pay\nA,100\n")], world.secret
    )
    tool(tools, "data_load")()
    did = tool(tools, "data_save")(name="payroll")["dataset_id"]
    # the boss in the secret channel sees it; so does the admin
    _m, in_secret = await tools_for(bot, world, world.boss, channel=world.secret)
    assert [d["dataset_id"] for d in tool(in_secret, "data_list")()["datasets"]] == [did]
    _m, admin = await tools_for(bot, world, world.admin, channel=world.secret)
    assert tool(admin, "data_list")()["datasets"][0]["dataset_id"] == did
    # staff cannot read the channel, so neither list nor load nor name guessing works
    _m, staff = await tools_for(bot, world, world.staff)
    assert tool(staff, "data_list")()["datasets"] == []
    assert tool(staff, "data_load")(dataset=did)["status"] == "error"
    assert tool(staff, "data_load")(dataset="payroll")["status"] == "error"
    # the boss asking in the PUBLIC channel must not leak it to the people reading there
    _m, boss_pub = await tools_for(bot, world, world.boss, channel=world.pub)
    assert tool(boss_pub, "data_list")()["datasets"] == []
    assert tool(boss_pub, "data_load")(dataset=did)["status"] == "error"
    # in a private chat / DM only the audience of one counts
    dm = TMsg(None, SimpleNamespace(id=777, log=[]), world.boss, "hi")
    dm_tools = bot._data_tools(dm, world.g, world.boss, [], "th")
    assert tool(dm_tools, "data_list")()["datasets"][0]["dataset_id"] == did


async def test_a_dm_dataset_is_the_owners_alone(bot, world):
    dm = TMsg(None, SimpleNamespace(id=778, log=[]), world.staff, "hi")
    files = [FileInput(0, "mine.csv", "text/plain", b"a\n1\n", raw=b"a,b\n1,2\n")]
    tools = bot._data_tools(dm, world.g, world.staff, files, "th")
    tool(tools, "data_load")()
    saved = tool(tools, "data_save")()
    assert saved["visible_to"] == "owner only"
    rec = bot.docs.get(ds.DATASETS, GUILD, saved["dataset_id"])
    assert rec["channel_id"] == 0 and rec["owner_id"] == world.staff.id
    _m, peer = await tools_for(bot, world, world.peer)
    assert tool(peer, "data_list")()["datasets"] == []
    dm2 = TMsg(None, SimpleNamespace(id=779, log=[]), world.staff, "again")
    again = bot._data_tools(dm2, world.g, world.staff, [], "th")
    assert tool(again, "data_load")(dataset="mine")["status"] == "ok"


async def test_voice_room_audience_must_all_read_the_source_channel(bot, world):
    _m, tools = await tools_for(
        bot, world, world.boss, [matt("pay.csv", "text/csv", b"a,b\n1,2\n")], world.secret
    )
    tool(tools, "data_load")()
    did = tool(tools, "data_save")()["dataset_id"]
    only_bosses = bot.room_message(world.g, world.boss, world.pub, room={world.boss.id, world.admin.id})  # fmt: skip
    ok = bot._data_tools(only_bosses, world.g, world.boss, [], "th")
    assert tool(ok, "data_load")(dataset=did)["status"] == "ok"
    mixed = bot.room_message(world.g, world.boss, world.pub, room={world.boss.id, world.staff.id})
    no = bot._data_tools(mixed, world.g, world.boss, [], "th")
    assert tool(no, "data_load")(dataset=did)["status"] == "error"


async def test_dataset_too_big_to_keep_says_so(bot, world, monkeypatch):
    monkeypatch.setattr(ds, "MAX_STORED_BYTES", 10)
    _m, tools = await tools_for(bot, world, world.staff, [matt("s.csv", "text/csv", b"a,b\n" + b"1.5,2\n" * 50)])  # fmt: skip
    tool(tools, "data_load")()
    out = tool(tools, "data_save")()
    assert out["status"] == "error" and "3 MB" in out["error"]
    assert tool(tools, "data_sql")("SELECT COUNT(*) FROM t_s")["rows"] == [
        [50]
    ]  # still usable here


def test_tools_never_raise_and_log_no_user_text(caplog):
    ctx = DataCtx(
        guild_id=1, asker_id=1, asker_name="a", is_exec=False, channel_id=1, lang="th", private=False,
        files=[], repo=None,  # type: ignore[arg-type]
    )  # fmt: skip
    tools = data_tools(ctx)
    ctx.repo = SimpleNamespace(
        all=lambda _g: (_ for _ in ()).throw(RuntimeError("boom secret-text")),
    )
    out = tool(tools, "data_list")()
    assert out["status"] == "error" and "secret-text" not in json.dumps(out)
    assert all("secret-text" not in r.getMessage() for r in caplog.records)


# ------------------------------------------------------------------ the agent
async def test_ask_picks_the_playbook_from_the_tools_and_attaches_the_chart(bot, world):
    m = world.staff
    msg, tools = await tools_for(bot, world, m, [matt("s.xlsx", "", xlsx_bytes())])
    seen = {}

    def first(req):
        seen["prompt"] = req.config.system_instruction
        return fc_resp(("data_load", {}))

    llm = scripted(
        first,
        fc_resp(
            (
                "data_chart",
                {
                    "query": "SELECT col_1 AS m, col_3 AS v FROM t_1 ORDER BY 1",
                    "kind": "line",
                    "x": "m",
                    "y": ["v"],
                    "title": "ยอดขาย",
                },
            )
        ),
        text_resp("แนบกราฟให้แล้วค่ะ"),
    )
    ans = await qa.ask_async(
        bot.s, bot.store, "ทำกราฟยอดขายให้หน่อย", guild_id=GUILD, is_exec=False, user_id=m.id,
        model=llm, extra_tools=tools,
    )  # fmt: skip
    assert "ทักษะวิเคราะห์ข้อมูล" in seen["prompt"] and "SANITY CHECKS" in seen["prompt"]
    assert "วิเคราะห์ข้อมูลจากไฟล์ Excel/CSV/JSON" in seen["prompt"]  # the capabilities line
    assert [n for n, _ in ans.attachments] == ["chart-data.png"]  # Thai title: ascii file name
    plain = scripted(lambda req: seen.update(p2=req.config.system_instruction) or text_resp("ok"))
    await qa.ask_async(
        bot.s, bot.store, "สวัสดี", guild_id=GUILD, is_exec=False, user_id=m.id, model=plain
    )
    assert "SANITY CHECKS" not in seen["p2"]


def test_prompt_section_follows_the_flag():
    from dory.config import Settings

    s = Settings(_env_file=None)
    on = qa.system_prompt(s, scoped=False, data=True)
    off = qa.system_prompt(s, scoped=False)
    assert "SANITY CHECKS" in on and "SANITY CHECKS" not in off
    text = Path(ds.__file__).with_name("skills").joinpath("data_analysis.md").read_text("utf-8")
    assert len(text.splitlines()) <= 60 and text.splitlines()[0].startswith("# ทักษะ")


def test_every_data_string_is_in_all_second_tier_catalogs():
    from dory.i18n import EXTRA

    locales = Path(ds.__file__).parent / "locales"
    for lang in EXTRA:
        strings = json.loads((locales / f"{lang}.json").read_text("utf-8"))["strings"]
        missing = [k for k in ds.STRINGS if k not in strings]
        assert not missing, (lang, missing)


def test_the_hub_lists_the_flag_in_thai_and_english():
    html = (Path(__file__).parents[2] / "src/meeting_scribe/static/hub.html").read_text("utf-8")
    for k in ("f_data_n", "f_data_d", "f_data_w"):
        assert k in html
    assert '{ id: "data" }' in html and html.count("data: true") >= 2
