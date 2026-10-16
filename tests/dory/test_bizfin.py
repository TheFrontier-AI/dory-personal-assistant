# ruff: noqa: F811, F401
"""Business finance math (known textbook values), market benchmark validation, security / scam
heuristics, the payment-slip checklist, tool wiring and the passive link guard. Offline."""

import io
import json
from datetime import UTC, date, datetime
from decimal import Decimal
from types import SimpleNamespace

import pytest
from conftest import GUILD, FakeClient, text_resp
from test_coordinator import BOSS, STAFF, world
from test_modes import TMsg, bot  # noqa: F401  (fixture)

from dory import bizfin as bf
from dory import security_check as sec
from dory.benchmark import DORY_FACTS, run_benchmark, validate
from dory.bizfin_tools import BizfinCtx, bizfin_tools
from dory.guild_settings import FEATURES, Features, GuildSettingsUpdate
from dory.i18n import CATALOG, t
from dory.models import Receipt
from dory.qa import system_prompt, tool_files

NOW = datetime(2026, 10, 16, 12, 0, tzinfo=UTC)


# ------------------------------------------------------------------ investment metrics
def test_textbook_npv_irr_roi_payback():
    r = bf.investment_metrics(1000, [300, 400, 500, 600], 10, "year")
    assert r["npv"] == 388.77  # -1000 + 272.73 + 330.58 + 375.66 + 409.81
    assert r["irr_period_pct"] == 24.89
    assert r["roi_pct"] == 80.0  # (1800 - 1000) / 1000
    assert r["payback_periods"] == 2.6  # 300 + 400 + 300 of 500 in year 3
    assert r["discounted_payback_periods"] == 3.05
    assert r["verdict"] == "npv_positive" and "NPV" in r["text"] and "IRR" in r["text"]


def test_irr_single_flow_and_negative_irr_and_none():
    assert bf.investment_metrics(100, [110], 5, "year")["irr_period_pct"] == 10.0
    assert bf.investment_metrics(1000, [100, 100], 5, "year")["irr_period_pct"] < 0
    r = bf.investment_metrics(1000, [-10, -10], 5, "year")  # never pays back
    assert r["irr_period_pct"] is None and r["payback_periods"] is None
    assert r["verdict"] == "npv_negative"


def test_monthly_flows_use_a_monthly_rate_and_annualise_the_irr():
    r = bf.investment_metrics(12000, [1500] * 12, 12, "month")
    assert abs(r["discount_rate_period_pct"] - 0.9489) < 1e-3  # (1.12)^(1/12) - 1
    assert abs(r["irr_annual_pct"] - ((1 + r["irr_period_pct"] / 100) ** 12 - 1) * 100) < 0.5
    assert r["payback_periods"] == 8.0


def test_investment_input_errors_are_clear():
    for args in [(0, [1]), (-5, [1]), (100, [])]:
        with pytest.raises(bf.FinError):
            bf.investment_metrics(*args)
    with pytest.raises(bf.FinError):
        bf.investment_metrics(100, [float("nan")])
    with pytest.raises(bf.FinError):
        bf.investment_metrics(True, [1])


# ------------------------------------------------------------------ unit economics / pricing
def test_unit_economics_break_even():
    r = bf.unit_economics(50000, 40, 100, 1000, 200000)
    assert r["contribution_margin_unit"] == 60.0 and r["gross_margin_pct"] == 60.0
    assert r["break_even_units_month"] == 834  # ceil(50000 / 60)
    assert r["break_even_revenue_month"] == 83400.0
    assert r["expected_profit_month"] == 10000.0 and r["months_to_recoup_investment"] == 20.0
    assert "834" in r["text"]


def test_unit_economics_price_below_cost_never_breaks_even():
    r = bf.unit_economics(1000, 120, 100)
    assert r["break_even"] == "never" and "warning" in r
    with pytest.raises(bf.FinError):
        bf.unit_economics(1000, 10, 0)


def test_unit_economics_low_volume_does_not_recoup():
    r = bf.unit_economics(50000, 40, 100, 100, 1000)
    assert r["months_to_recoup_investment"] is None and "warning" in r


def test_cost_plus_margin_vs_markup_and_vat():
    r = bf.price_from_cost(["วัตถุดิบ: 40", "box=10"], 30)
    assert r["cost_total"] == 50.0 and r["cost_plus_price"] == 71.43  # 50 / 0.7
    std = next(t for t in r["tiers"] if t["tier"] == "standard")
    assert std["price_ex_vat"] == 71.43 and std["vat"] == 5.0 and std["price_inc_vat"] == 76.43
    m = bf.price_from_cost([50], markup_pct=30)
    assert m["cost_plus_price"] == 65.0 and m["basis"] == "markup"
    with pytest.raises(bf.FinError):
        bf.price_from_cost([50], 30, 20)
    with pytest.raises(bf.FinError):
        bf.price_from_cost([50], 100)
    with pytest.raises(bf.FinError):
        bf.price_from_cost([])


def test_competitor_band_sets_tiers_and_never_below_cost():
    r = bf.price_from_cost([50], 30, competitor_prices=[120, 150, 200])
    assert r["competitor_band"]["median"] == 150.0
    tiers = {t["tier"]: t["price_ex_vat"] for t in r["tiers"]}
    assert tiers["entry"] <= tiers["standard"] <= tiers["premium"] and tiers["entry"] > 50
    inc = bf.price_from_cost([50], 30, competitor_prices=[107], competitors_include_vat=True)
    assert inc["competitor_band"]["median"] == 100.0  # 107 / 1.07
    assert any("below the competitor" in w for w in r["warnings"])


def test_parse_items_accepts_numbers_strings_and_dicts():
    items = bf.parse_items([5, "ค่าส่ง: 1,200.50", "x 3", {"name": "a", "amount": 2}])
    assert [a for _n, a in items] == [Decimal(5), Decimal("1200.50"), Decimal(3), Decimal(2)]
    with pytest.raises(bf.FinError):
        bf.parse_items(["no number here"])


# ------------------------------------------------------------------ projection
def proj(**kw):
    base = dict(arpu_month=100, new_customers_month=10, churn_pct_month=0, cac=0,
                fixed_costs_month=0, months=12)  # fmt: skip
    return bf.financial_projection(**{**base, **kw})


def test_projection_is_arithmetic_not_guesswork():
    r = proj()
    assert [row["customers"] for row in r["table"]][:3] == [10, 20, 30]
    assert r["table"][11]["customers"] == 120 and r["table"][11]["revenue"] == 12000.0
    assert r["metrics"]["end_cash"] == sum(row["revenue"] for row in r["table"])
    assert r["metrics"]["cash_out_month"] is None and r["metrics"]["first_profitable_month"] == 1


def test_ltv_cac_payback_and_runway():
    r = proj(churn_pct_month=5, cac=800, variable_cost_customer=20, fixed_costs_month=60000,
             opening_cash=100000, months=24)  # fmt: skip
    m = r["metrics"]
    assert m["ltv"] == 1600.0 and m["ltv_cac"] == 2.0 and m["cac_payback_months"] == 10.0
    assert m["cash_out_month"] is not None and m["lowest_cash"] < 0


def test_scenarios_are_ordered_and_bounds_checked():
    r = proj(churn_pct_month=3, cac=100, fixed_costs_month=500, acquisition_growth_pct_month=2)
    s = r["scenarios"]
    assert s["low"]["end_cash"] < s["base"]["end_cash"] < s["high"]["end_cash"]
    assert bf.financial_projection(100, 10, 3, 100, 500, scenario="low")["scenario"] == "low"
    for bad in ({"months": 6}, {"months": 40}, {"churn_pct_month": 120}, {"scenario": "wild"}):
        with pytest.raises(bf.FinError):
            proj(**bad)


def test_data_checklist_lists_missing_inputs_by_product_type():
    r = bf.financial_data_checklist("แอปจัดประชุม", ["price", "cac"])
    assert r["product_type"] == "saas" and "price" not in r["missing"]
    assert {"churn_pct_month", "initial_investment"} <= set(r["missing"])
    assert bf.financial_data_checklist("ร้านอาหาร")["product_type"] == "retail"
    assert bf.financial_data_checklist("???")["product_type"] == "generic"
    assert "[x]" in r["text"] and bf.financial_data_checklist("saas", lang="en")["text"].startswith(
        "Data"
    )


def test_xlsx_and_csv_files_roundtrip():
    from openpyxl import load_workbook

    name, data = bf.table_file("p", {"T": [{"a": 1, "b": 2.5}], "U": [{"x": "ไทย"}]})
    wb = load_workbook(io.BytesIO(data))
    assert name == "p.xlsx" and wb.sheetnames == ["T", "U"] and wb["T"]["B2"].value == 2.5
    assert bf.to_csv([{"a": 1}]).startswith(b"\xef\xbb\xbfa")


# ------------------------------------------------------------------ URL heuristics
@pytest.mark.parametrize(
    "url,verdict,sign",
    [
        ("https://www.google.com/search?q=x", "safe", None),
        ("https://accounts.google.com/signin", "safe", None),
        ("https://github.com/a/b", "safe", None),
        ("https://www.ktb.go.th/", "safe", None),
        ("https://www.kasikornbank.com/th", "safe", None),
        ("https://example.com/page", "safe", None),
        ("https://pineapple.com", "safe", None),
        ("http://google.com@evil.xyz/a", "dangerous", "userinfo"),
        ("https://xn--ggle-0nda.com/", "dangerous", "punycode"),
        ("https://kbank-verify.top/login", "dangerous", "brand_lookalike"),
        ("https://scb.co.th.evil.com/x", "dangerous", "brand_lookalike"),
        ("https://kasikornbamk.com", "dangerous", "brand_typo"),
        ("https://paypa1.com", "dangerous", "brand_lookalike"),
        ("https://shop.example.co.th/download/app.apk", "dangerous", "risky_file"),
        ("http://192.168.1.5/pay", "suspicious", "ip_host"),
        ("https://bit.ly/3abcd", "suspicious", "shortener"),
        ("https://free-gift.click/win", "suspicious", "bad_tld"),
    ],
)
def test_url_heuristics(url, verdict, sign):
    rep = sec.analyze_url(url)
    assert rep.verdict == verdict, [f.id for f in rep.findings]
    if sign:
        assert sign in [f.id for f in rep.findings]


def test_lookalike_of_an_allowlisted_domain_is_not_allowlisted():
    assert sec.analyze_url("https://google.com.evil.xyz/login").findings
    assert not sec.is_allowlisted("notgithub.com") and sec.is_allowlisted("gist.github.com")
    assert sec.registrable("a.b.example.co.th") == "example.co.th"


def test_masked_link_text_must_match_the_destination():
    links = sec.find_links("ดูที่ [https://kasikornbank.com/x](https://evil.example.net/p) ด่วน")
    assert links == [("https://evil.example.net/p", "https://kasikornbank.com/x")]
    rep = sec.analyze_url(*links[0])
    assert rep.verdict == "dangerous" and rep.findings[0].id == "masked_mismatch"
    ok = sec.analyze_url("https://github.com/a", "https://github.com/a")
    assert ok.verdict == "safe"


def test_find_links_bare_domains_and_cap():
    got = [u for u, _ in sec.find_links("go to kbank-th.top/login now, or https://a.com/x.")]
    assert "kbank-th.top/login" in got and "https://a.com/x" in got
    assert len(sec.find_links(" ".join(f"https://s{i}.com" for i in range(30)))) == 10


# ------------------------------------------------------------------ message heuristics
SCAM_TH = "ด่วน! บัญชีของคุณถูกระงับ ยืนยันตัวตนที่ https://kbank-secure.top/login ภายใน 24 ชั่วโมง"


def test_thai_bank_scam_message_is_dangerous_with_reasons_in_both_languages():
    res = sec.check(SCAM_TH)
    assert res.verdict == "dangerous"
    assert {"brand_lookalike", "urgency", "account_threat", "verify_credentials"} <= set(
        res.reasons
    )
    th, en = sec.render_check(res, "th"), sec.render_check(res, "en")
    assert "อันตราย" in th and "1441" in th and "Dangerous" in en and "thaipoliceonline" in en


@pytest.mark.parametrize(
    "text,pid",
    [
        ("กรุณาส่งรหัส OTP ที่ได้รับให้เจ้าหน้าที่", "otp"),
        ("Please share the one-time code you received", "otp"),
        ("ตำรวจ แจ้งว่าบัญชีคุณเกี่ยวข้องกับการฟอกเงิน", "gov_threat"),
        ("ยินดีด้วย คุณได้รับรางวัลมูลค่า 50,000 บาท", "prize"),
        ("โอนค่าธรรมเนียม 500 บาทก่อน เพื่อรับเงินรางวัล", "pay_first"),
        ("ลงทุนวันละ 5% การันตี ไม่มีความเสี่ยง", "invest_promise"),
        ("งานพาร์ทไทม์ รายได้วันละ 800 กดไลก์", "job_scam"),
        ("พัสดุตกค้าง ที่อยู่ไม่ถูกต้อง กรุณาอัปเดต", "parcel"),
        ("Your account will be suspended, verify your account", "account_threat"),
        ("ติดตั้งแอปเพื่อยืนยัน ธนาคาร ไฟล์ apk", "install_app"),
    ],
)
def test_scam_language_patterns(text, pid):
    assert pid in [f.id for f in sec.analyze_text(text)]


def test_normal_work_chat_is_safe_and_empty_is_nothing():
    for msg in ["พรุ่งนี้ประชุมสิบโมง ส่งไฟล์ให้ด้วยนะ", "Please review https://github.com/org/repo/pull/3"]:
        assert sec.check(msg).verdict == "safe"
    assert sec.check("   ") is None
    assert sec.render_check(None, "en") == t("sec_nothing_to_check", "en")
    safe = sec.render_check(sec.check("hello"), "th")
    assert "ไม่ได้เปิดดูหน้าเว็บจริง" in safe  # never a guarantee


def test_urgency_alone_is_not_enough_but_with_a_foreign_link_it_is_suspicious():
    assert sec.check("ด่วน ส่งงานภายใน 2 ชั่วโมง").verdict == "safe"
    assert sec.check("ด่วน เปิดดูภายใน 2 ชั่วโมง https://example.org/x").verdict == "suspicious"


# ------------------------------------------------------------------ Safe Browsing
def test_safe_browsing_off_without_a_key_and_fails_open():
    calls = []
    assert sec.safe_browsing_lookup(["https://a.com"], "", lambda *a: calls.append(a)) == set()
    assert calls == []

    def boom(*a):
        raise RuntimeError("secret-key-in-message")

    assert sec.safe_browsing_lookup(["https://a.com"], "K", boom) == set()


def test_safe_browsing_match_makes_it_dangerous():
    seen = {}

    def post(url, key, body):
        seen["body"] = body
        return {"matches": [{"threat": {"url": "https://example.com/p"}}]}

    res = sec.check("https://example.com/p", safe_browsing_key="K", post=post)
    assert res.verdict == "dangerous" and "safe_browsing" in res.reasons and res.safe_browsing
    assert seen["body"]["threatInfo"]["threatEntries"] == [{"url": "https://example.com/p"}]
    assert sec.check("https://example.com/p").verdict == "safe"  # no key: heuristics only


# ------------------------------------------------------------------ passive guard
def test_link_guard_hits_only_high_confidence_never_allowlisted():
    assert sec.link_guard_hits("see https://bit.ly/abc and http://1.2.3.4/x") == []
    assert sec.link_guard_hits("https://github.com/x https://docs.google.com/d") == []
    hits = sec.link_guard_hits("โอนที่ https://kbank-th.top/pay ด่วน")
    assert [h.host for h in hits] == ["kbank-th.top"]
    masked = sec.link_guard_hits("[https://scb.co.th](https://evil.example.net/p)")
    assert [h.host for h in masked] == ["evil.example.net"]
    assert "kbank" not in sec.guard_message(hits, "en") or "Careful" in sec.guard_message(
        hits, "en"
    )
    assert "ระวัง" in sec.guard_message(hits, "th")


def test_guard_limiter_keeps_quiet():
    now = [0.0]
    lim = sec.GuardLimiter(lambda: now[0])
    assert lim.allow(1, 10, ["a.top"])
    assert not lim.allow(1, 10, ["b.top"])  # same channel within 10 minutes
    assert not lim.allow(1, 11, ["a.top"])  # same host within 6 hours
    now[0] = 601
    assert lim.allow(1, 10, ["b.top"])
    now[0] = 7 * 3600
    assert lim.allow(1, 11, ["a.top"])  # host window over
    assert lim.allow(2, 10, ["a.top"])  # other guild is independent
    lim2 = sec.GuardLimiter(lambda: now[0])
    sent = 0
    for i in range(40):  # different channels and hosts, 10 s apart: only the hourly cap bites
        now[0] += 10
        sent += lim2.allow(1, 100 + i, [f"h{i}.top"])
    assert sent == 20
    now[0] += 3600
    assert lim2.allow(1, 500, ["new.top"])


# ------------------------------------------------------------------ slips
def slip(**kw):
    base = dict(bank="กสิกรไทย", amount=1500, date_iso="2026-10-16", time_hhmm="09:30",
                ref_no="015289123456ABC", sender="สมชาย ใจดี", receiver="บริษัท เอ จำกัด",
                now=NOW, lang="th")  # fmt: skip
    return sec.verify_slip_data(**{**base, **kw})


def test_a_consistent_slip_has_no_issues_but_is_never_confirmed():
    r = slip()
    assert r["status"] == "no_issues_found" and r["issues"] == [] and not r["can_confirm_with_bank"]
    assert "ยืนยันกับธนาคารไม่ได้" in r["text"] and "QR" in r["text"]
    assert "cannot confirm with the bank" in slip(lang="en")["text"]


def test_slip_flags_inconsistencies():
    assert any("อนาคต" in i for i in slip(date_iso="2026-10-20")["issues"])
    assert any("เก่ามาก" in i for i in slip(date_iso="2026-06-01")["issues"])
    assert any("1,000.00" in i for i in slip(expected_amount=1000)["issues"])
    assert not slip(expected_amount=1500.00)["issues"]
    assert any("เลขอ้างอิง" in i for i in slip(ref_no="1111")["issues"])
    assert any("ชื่อเดียวกัน" in i for i in slip(receiver="สมชาย  ใจดี")["issues"])
    assert any("ผู้รับ" in i for i in slip(expected_receiver="ร้านอื่น")["issues"])
    assert any("ไม่รู้จักชื่อธนาคาร" in i for i in slip(bank="ธนาคารแห่งหนึ่ง")["issues"])
    assert slip(time_hhmm="27:99")["issues"]
    assert slip(date_iso="2569-10-16")["status"] == "no_issues_found"  # Buddhist year converted
    assert slip(amount=-5)["issues"]


def test_slip_future_time_today():
    assert any("อนาคต" in i for i in slip(time_hhmm="18:00")["issues"])
    assert slip(time_hhmm="09:00")["status"] == "no_issues_found"


def test_slip_reports_unreadable_fields_and_duplicate_reference():
    r = slip(bank=None, ref_no=None, sender=None)
    assert {"bank", "ref", "sender"} <= set(r["missing_fields"]) and "อ่านไม่ได้" in r["text"]
    d = slip(ledger_refs=lambda ref: 42 if ref == "015289123456ABC" else None)
    assert any("#42" in i for i in d["issues"]) and d["status"] == "issues"


# ------------------------------------------------------------------ benchmark
def fake_search(client, model, q, *, today=None):
    return {
        "answer": f"Tool A costs $10 per month [1]. Tool B has no Thai support [2]. ({q[:12]})",
        "sources": [
            {"n": 1, "title": "A pricing", "domain": "a.example", "link": "https://a.example/p"},
            {"n": 2, "title": "B review", "domain": "b.example", "link": "https://b.example/r"},
        ],
    }


GOOD = {
    "summary": "Few tools handle Thai well.",
    "gaps": [{"gap": "No Thai speaker labels", "sources": [2]}, {"gap": "invented", "sources": [9]}],
    "competitors": [
        {"name": "Tool A", "pricing": "$10/mo", "strengths": "cheap", "weaknesses": "English only",
         "thai_support": "no", "sources": [1]},
        {"name": "Ghost", "pricing": "$1", "strengths": "x", "weaknesses": "y",
         "thai_support": "yes", "sources": [7]},
        {"name": "NoSource", "pricing": "unknown", "sources": []},
    ],
    "ours_wins": [{"point": "Thai voice meetings in Discord", "vs": "Tool A"}],
    "ours_loses": [{"point": "No Zoom bot", "vs": "Tool A", "sources": [1, 9]}],
    "next_moves": ["Add Meet"],
}  # fmt: skip


def test_validation_drops_rows_without_a_real_source():
    clean, dropped = validate(GOOD, {1, 2})
    assert [c["name"] for c in clean["competitors"]] == ["Tool A"] and dropped == 3
    assert [g["gap"] for g in clean["gaps"]] == ["No Thai speaker labels"]
    assert clean["ours_loses"][0]["sources"] == [1]
    assert validate({"competitors": "junk", "gaps": [5]}, {1})[0]["competitors"] == []


def test_run_benchmark_renders_code_built_report_with_real_sources_only():
    client = FakeClient([text_resp(json.dumps(GOOD))])
    res = run_benchmark(client, topic="AI meeting notes", lang="en", today=date(2026, 10, 16),
                        search_model="m", synth_model="m", search=fake_search)  # fmt: skip
    assert res.status == "ok" and res.dropped == 3
    md = res.markdown
    assert (
        "| Tool A | $10/mo |" in md and "[1]" in md and "Ghost" not in md and "NoSource" not in md
    )
    assert "https://a.example/p" in md and "https://b.example/r" in md and "[9]" not in md
    assert "Where we win" in md and "Where we lose" in md and "estimate" in md
    prompt = client.models.calls[0][1][0].parts[0].text
    assert "Dory (ดอรี่) is an AI teammate" in prompt  # default: facts from code
    th = run_benchmark(FakeClient([text_resp(json.dumps(GOOD))]), topic="x", lang="th",
                       today=date(2026, 10, 16), search_model="m", synth_model="m",
                       search=fake_search)  # fmt: skip
    assert "ตารางเทียบคู่แข่ง" in th.markdown and "ไม่รองรับ" in th.markdown


def test_benchmark_uses_the_users_product_description_and_handles_failures():
    client = FakeClient([text_resp(json.dumps(GOOD))])
    run_benchmark(client, topic="t", our_product="MyApp does X", lang="en",
                  today=date(2026, 10, 16), search_model="m", synth_model="m",
                  search=fake_search)  # fmt: skip
    prompt = client.models.calls[0][1][0].parts[0].text
    assert "MyApp does X" in prompt and "Dory (ดอรี่)" not in prompt

    def nothing(*a, **k):
        return {"error": "down"}

    r = run_benchmark(FakeClient([]), topic="t", today=date(2026, 10, 16), search_model="m",
                      synth_model="m", search=nothing)  # fmt: skip
    assert r.status == "no_results"
    junk = run_benchmark(FakeClient([text_resp("not json")]), topic="t", today=date(2026, 10, 16),
                         search_model="m", synth_model="m", search=fake_search)  # fmt: skip
    assert junk.status == "no_results"


def test_dory_facts_are_short_and_state_limits():
    assert len(DORY_FACTS) < 1600 and "Limits" in DORY_FACTS


# ------------------------------------------------------------------ tools
def make_ctx(**kw):
    saved = []
    base = dict(lang="th", now=lambda: NOW, client=None, model="m",
                save_idea=lambda title, section, md: saved.append((title, section, md))
                or {"status": "saved"})  # fmt: skip
    ctx = BizfinCtx(**{**base, **kw})
    return ctx, saved


def by_name(tools):
    return {f.__name__: f for f in tools}


def test_tool_set_and_web_switch():
    ctx, _ = make_ctx()
    names = set(by_name(bizfin_tools(ctx)))
    assert {"unit_economics", "price_from_cost", "investment_metrics", "financial_projection",
            "financial_data_checklist", "security_check", "verify_slip",
            "market_benchmark"} == names  # fmt: skip
    assert "market_benchmark" not in by_name(bizfin_tools(ctx, web=False))
    assert all(f.outbox is ctx.outbox for f in bizfin_tools(ctx))


def test_calculator_tools_return_text_numbers_and_files():
    ctx, saved = make_ctx()
    t_ = by_name(bizfin_tools(ctx))
    r = t_["unit_economics"](50000, 40, 100, 1000, 0, save_to_idea="CCTV Watch")
    assert (
        r["status"] == "ok" and r["break_even_units_month"] == 834 and r["saved_to_idea"] == "saved"
    )
    assert saved[0][0] == "CCTV Watch" and "834" in saved[0][2]
    p = t_["financial_projection"](500, 20, 5, 800, 60000, 100, 0, 5, 300000, 24)
    assert p["status"] == "ok" and p["attached_file"] == "projection.xlsx" and "table" not in p
    assert p["metrics"]["ltv"] == 8000.0 and p["table_rows"] == 24
    i = t_["investment_metrics"](1000, [300, 400, 500, 600], 10, "year", file_format="csv")
    assert i["npv"] == 388.77 and i["attached_file"] == "investment.csv"
    pr = t_["price_from_cost"](["40", "10"], 30, file_format="pdf")
    assert pr["attached_file"] == "pricing.pdf"
    names = [n for n, _d in tool_files(list(t_.values()))]
    assert names == ["projection.xlsx", "investment.csv", "pricing.pdf"]
    assert dict(ctx.outbox)["pricing.pdf"][:5] == b"%PDF-"


def test_tool_errors_are_data_and_saving_needs_ideas():
    ctx, _ = make_ctx(save_idea=None)
    t_ = by_name(bizfin_tools(ctx))
    assert t_["unit_economics"](10, 5, 0)["status"] == "error"
    assert t_["investment_metrics"](100, [])["status"] == "error"
    assert t_["unit_economics"](10, 5, 10, save_to_idea="X")["saved_to_idea"] == "unavailable"
    assert t_["financial_data_checklist"]("saas", ["price"])["missing"]


def test_security_and_slip_tools():
    ctx, _ = make_ctx(ledger_ref=lambda ref: 7 if ref == "ABCDEF123456" else None)
    t_ = by_name(bizfin_tools(ctx))
    r = t_["security_check"](SCAM_TH)
    assert r["verdict"] == "dangerous" and r["hosts"] == ["kbank-secure.top"]
    assert t_["security_check"]("")["status"] == "nothing_to_check"
    s = t_["verify_slip"](bank="SCB", amount=100, date_iso="2026-10-16", time_hhmm="08:00",
                          ref_no="ABCDEF123456", sender="A B", receiver="C D")  # fmt: skip
    assert s["status"] == "issues" and any("#7" in i for i in s["issues"])
    assert t_["verify_slip"]()["missing_fields"]


def test_market_benchmark_tool_attaches_a_pdf_and_needs_a_client():
    ctx, _ = make_ctx()
    assert by_name(bizfin_tools(ctx))["market_benchmark"]("AI notes")["status"] == "error"
    ctx, _ = make_ctx(client=FakeClient([text_resp(json.dumps(GOOD))]), search=fake_search)
    r = by_name(bizfin_tools(ctx))["market_benchmark"]("AI meeting notes")
    assert r["status"] == "ok" and r["attached_file"].endswith(".pdf")
    assert r["sources"][0]["link"] == "https://a.example/p" and "Tool A" in r["report_markdown"]
    assert ctx.outbox[0][1][:5] == b"%PDF-"


# ------------------------------------------------------------------ bot wiring
def kw_for(bot, world, member, **extra):
    msg = TMsg(world.g, world.pub, member, "x")
    return bot._agent_kwargs(msg, world.g, member, False, setting="th", lang="th", question="x",
                             memory=None, profile=None, **extra)  # fmt: skip


async def test_tools_registered_for_chat_and_voice_and_prompt_follows(bot, world):
    member = world.g.members[STAFF]
    names = {f.__name__ for f in kw_for(bot, world, member)["extra_tools"]}
    assert {"unit_economics", "security_check", "verify_slip", "financial_projection"} <= names
    from dory.voice_actions import ACTION_TOOL_NAMES

    assert {"unit_economics", "price_from_cost", "investment_metrics", "financial_projection",
            "verify_slip"} <= ACTION_TOOL_NAMES  # fmt: skip
    assert "การเงินธุรกิจ" in system_prompt(bot.s, scoped=False, bizfin=True)
    assert "การเงินธุรกิจ" not in system_prompt(bot.s, scoped=False)
    inter = kw_for(bot, world, member, interjection=True)
    assert "unit_economics" not in {f.__name__ for f in inter.get("extra_tools", [])}


async def test_flag_off_means_no_tools_and_benchmark_needs_web(bot, world):
    member = world.g.members[STAFF]
    bot.guilds_store.save_settings(GUILD, GuildSettingsUpdate(features={"bizfin": False}), "x")
    bot.team._cache.clear()
    assert bot._bizfin_tools(world.g, member, False, "th") == []
    bot.guilds_store.save_settings(GUILD, GuildSettingsUpdate(features={"bizfin": True}), "x")
    bot.team._cache.clear()
    assert "unit_economics" in {f.__name__ for f in bot._bizfin_tools(world.g, member, False, "th")}
    # no model client in this bot: no web research tool, calculators remain
    assert "market_benchmark" not in {
        f.__name__ for f in bot._bizfin_tools(world.g, member, False, "th")
    }


async def test_slip_duplicate_check_is_limited_to_what_the_asker_may_see(bot, world):
    def add(uid, ref, name):
        bot.store.add_receipt(
            Receipt(guild_id=GUILD, submitter_id=uid, submitter_name=name, sha256=ref + str(uid),
                    effective_date="2026-10-10", invoice_no=ref, total=100.0)
        )  # fmt: skip

    add(BOSS, "REF-AAAA-111111", "boss")
    add(STAFF, "REF-BBBB-222222", "staff")
    staff, boss = world.g.members[STAFF], world.g.members[BOSS]
    assert bot._bizfin_ledger_ref(world.g, staff, False, "refbbbb222222") is not None
    assert bot._bizfin_ledger_ref(world.g, staff, False, "REFAAAA111111") is None  # not theirs
    assert bot._bizfin_ledger_ref(world.g, boss, True, "REFAAAA111111") is not None  # executive
    assert bot._bizfin_ledger_ref(world.g, boss, True, "x1") is None  # too short to match


async def test_save_to_idea_uses_the_askers_own_idea_document(bot, world):
    member = world.g.members[STAFF]
    tools = by_name(bot._bizfin_tools(world.g, member, False, "th"))
    r = tools["unit_economics"](1000, 10, 50, save_to_idea="Cafe Plan")
    assert r["saved_to_idea"] == "saved"
    doc = bot.ideas.get(GUILD, STAFF, "cafe-plan")
    assert doc is not None and "Unit economics" in [s.title for _k, s in doc.ordered()]


# ------------------------------------------------------------------ passive link guard in the bot
DANGER = "โอนด่วนที่นี่ https://kbank-th.top/pay"


async def guard_send(bot, world, text, who=STAFF, **kw):
    m = TMsg(world.g, world.pub, world.g.members[who], text, **kw)
    await bot.on_message(m)
    return m


async def test_guard_warns_once_under_a_dangerous_link(bot, world):
    m = await guard_send(bot, world, DANGER)
    assert len(m.replies) == 1 and "ระวัง" in m.replies[0]
    again = await guard_send(bot, world, "ลองอีกที " + DANGER)
    assert again.replies == []  # rate-limited: same host / channel
    assert sorted(bot.calls) == []  # the agent was never asked


async def test_guard_ignores_safe_links_allowlist_and_plain_chat(bot, world):
    for text in ["https://github.com/a/b/pull/1", "https://www.google.com/search?q=x",
                 "https://bit.ly/abc", "ไม่มีลิงก์ค่ะ"]:  # fmt: skip
        assert (await guard_send(bot, world, text)).replies == []


async def test_guard_flag_off_and_quiet_modes_and_direct_questions(bot, world):
    bot.guilds_store.save_settings(GUILD, GuildSettingsUpdate(features={"link_guard": False}), "x")
    bot.team._cache.clear()
    assert (await guard_send(bot, world, DANGER)).replies == []
    bot.guilds_store.save_settings(GUILD, GuildSettingsUpdate(features={"link_guard": True}), "x")
    bot.team._cache.clear()
    bot.guilds_store.set_channel_mode(GUILD, world.pub.id, "listen")
    bot.team._cache.clear()
    assert (await guard_send(bot, world, DANGER)).replies == []
    bot.guilds_store.set_channel_mode(GUILD, world.pub.id, "call")
    bot.team._cache.clear()
    asked = await guard_send(bot, world, "ลิงก์นี้ปลอดภัยไหม " + DANGER, mention=True)
    assert all("ระวัง" not in r for r in asked.replies)  # the agent answers, not the guard


# ------------------------------------------------------------------ flags and strings
def test_feature_flags_default_on():
    for f in ("bizfin", "link_guard"):
        assert f in FEATURES and getattr(Features(), f) is True


def test_every_user_facing_string_exists_in_thai_and_english():
    keys = [k for k in CATALOG if k.startswith(("sec_", "slip_"))]
    assert len(keys) > 25
    for k in keys:
        assert CATALOG[k]["th"].strip() and CATALOG[k]["en"].strip()
        assert set(__import__("re").findall(r"\{(\w+)\}", CATALOG[k]["th"])) == set(
            __import__("re").findall(r"\{(\w+)\}", CATALOG[k]["en"])
        ), k
    for th, en in sec.REASONS.values():
        assert th.strip() and en.strip()
    for rid in {f for tp in [sec.TEXT_PATTERNS] for f in tp} | {
        "userinfo", "ip_host", "punycode", "mixed_script", "shortener", "bad_tld",
        "brand_lookalike", "brand_typo", "keywords", "many_labels", "no_https", "risky_file",
        "masked_mismatch", "safe_browsing",
    }:  # fmt: skip
        assert rid in sec.REASONS, rid
