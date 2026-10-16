"""Slide decks: outline models, validation, filenames, rendering and charts (dory.deck).
Offline: rendered files are read back with python-pptx."""

import io
from datetime import date

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pydantic import ValidationError

from dory.deck import (
    LATIN_FONT,
    THAI_FONT,
    ChartData,
    ChartQuery,
    DeckSpec,
    chart_data,
    deck_filename,
    render_deck,
)


def bullets(n=2, text="ข้อ"):
    return {"kind": "bullets", "title": "หัวข้อ", "bullets": [f"{text} {i}" for i in range(n)]}


# ------------------------------------------------------------------ outline models
def test_valid_outline_parses():
    spec = DeckSpec.model_validate({"title": "สรุปเดือนตุลาคม", "slides": [bullets()]})
    assert spec.slides[0].kind == "bullets" and spec.subtitle == ""


@pytest.mark.parametrize(
    "bad, field",
    [
        ({"title": "x" * 81, "slides": [bullets()]}, "title"),
        ({"title": "t", "slides": []}, "slides"),
        ({"title": "t", "slides": [bullets()] * 13}, "slides"),
        ({"title": "t", "slides": [bullets(7)]}, "bullets"),
        (
            {"title": "t", "slides": [{"kind": "bullets", "title": "t", "bullets": ["x" * 141]}]},
            "bullets",
        ),
        ({"title": "t", "slides": [{"kind": "bullets", "title": "t", "bullets": []}]}, "bullets"),
        ({"title": "t", "slides": [{"kind": "chart", "title": "t"}]}, "chart"),
        ({"title": "t", "slides": [{"kind": "video", "title": "t"}]}, "kind"),
    ],
)
def test_invalid_outlines_name_the_field(bad, field):
    with pytest.raises(ValidationError) as e:
        DeckSpec.model_validate(bad)
    assert field in str(e.value)


@pytest.mark.parametrize(
    "q, ok",
    [
        ({"group_by": "category", "start_date": "2026-10-01", "end_date": "2026-10-31"}, True),
        (
            {"group_by": "day", "start_date": "2026-09-01", "end_date": "2026-10-31"},
            True,
        ),  # 61 days
        ({"group_by": "day", "start_date": "2026-08-01", "end_date": "2026-10-31"}, False),  # 92
        ({"group_by": "category", "start_date": "2026-10-31", "end_date": "2026-10-01"}, False),
        ({"group_by": "category", "start_date": "31/10/2026", "end_date": "2026-10-31"}, False),
        (
            {
                "group_by": "category",
                "start_date": "2026-10-01",
                "end_date": "2026-10-31",
                "top": 11,
            },
            False,
        ),
    ],
)
def test_chart_query_dates_and_limits(q, ok):
    if ok:
        ChartQuery.model_validate(q)
    else:
        with pytest.raises(ValidationError):
            ChartQuery.model_validate(q)


def test_deck_filename_slug_and_fallback():
    d = date(2026, 10, 16)
    assert (
        deck_filename("Spending report: Oct 2026!", d) == "spending-report-oct-2026-20261016.pptx"
    )
    assert deck_filename("สรุปค่าใช้จ่าย", d) == "dory-deck-20261016.pptx"
    assert deck_filename("🎉 Q4 / ../plan", d) == "q4-plan-20261016.pptx"
    assert len(deck_filename("word " * 40, d)) <= 64


# ------------------------------------------------------------------ rendering: text slides
def render(spec_dict, lang="th", charts=None):
    data = render_deck(
        DeckSpec.model_validate(spec_dict), charts or {}, lang=lang, today=date(2026, 10, 16)
    )
    return Presentation(io.BytesIO(data))


def texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def test_title_slide_then_one_slide_per_outline_slide():
    prs = render(
        {
            "title": "สรุปประชุม",
            "subtitle": "ทีมขาย",
            "slides": [
                {"kind": "section", "title": "ภาพรวม"},
                {
                    "kind": "bullets",
                    "title": "สิ่งที่ตกลง",
                    "bullets": ["เลื่อนเปิดร้านเป็นวันศุกร์"],
                    "notes": "รายละเอียด",
                },
                {"kind": "two_column", "title": "ก่อน / หลัง", "bullets": ["เดิม"], "right": ["ใหม่"]},
            ],
        }
    )
    assert len(prs.slides) == 4 and prs.slide_width == Inches(13.333)
    cover = " ".join(texts(prs.slides[0]))
    assert "สรุปประชุม" in cover and "ทีมขาย" in cover and "16 ตุลาคม 2569" in cover
    assert "ดอรี่" in cover
    assert "เลื่อนเปิดร้านเป็นวันศุกร์" in " ".join(texts(prs.slides[2]))
    assert prs.slides[2].notes_slide.notes_text_frame.text == "รายละเอียด"
    assert {"เดิม", "ใหม่"} <= set(" ".join(texts(prs.slides[3])).split())


def test_every_run_sets_the_thai_font_after_latin():
    prs = render({"title": "ทดสอบ", "slides": [bullets()]})
    runs = [
        r
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
        for r in p.runs
    ]
    assert runs
    for r in runs:
        rpr = r._r.find(qn("a:rPr"))
        tags = [c.tag for c in rpr]
        assert rpr.find(qn("a:latin")).get("typeface") == LATIN_FONT
        assert rpr.find(qn("a:cs")).get("typeface") == THAI_FONT
        assert tags.index(qn("a:cs")) == tags.index(qn("a:latin")) + 1


def _bullet_size(prs):
    """Largest run size on the first content slide, ignoring its title ("t")."""
    return max(
        r.font.size
        for s in prs.slides[1].shapes
        if s.has_text_frame
        for p in s.text_frame.paragraphs
        if p.text != "t"
        for r in p.runs
    )


def test_long_text_gets_a_smaller_font():
    short = render({"title": "t", "slides": [{"kind": "bullets", "title": "t", "bullets": ["สั้น"]}]})
    long_ = render(
        {
            "title": "t",
            "slides": [{"kind": "bullets", "title": "t", "bullets": ["ยาวมาก" * 23] * 6}],
        }
    )
    assert _bullet_size(long_) < _bullet_size(short) and _bullet_size(long_) >= Pt(16)


def test_control_characters_and_newlines_are_cleaned():
    prs = render(
        {
            "title": "a\x00b",
            "slides": [{"kind": "bullets", "title": "t", "bullets": ["line1\nline2\x0b end"]}],
        }
    )
    body = " ".join(texts(prs.slides[1]))
    assert "line1 line2 end" in body and "\x0b" not in body
    assert "ab" in " ".join(texts(prs.slides[0]))


def test_english_fixed_strings():
    cover = " ".join(texts(render({"title": "Q4", "slides": [bullets()]}, lang="en").slides[0]))
    assert "Made with Dory" in cover and "16 October 2026" in cover


def test_page_numbers_on_content_slides():
    prs = render({"title": "t", "slides": [bullets(), bullets()]})
    assert "2" in texts(prs.slides[1]) and "3" in texts(prs.slides[2])


def test_notes_master_is_listed_in_the_presentation():
    # python-pptx adds the notes master part but not p:notesMasterIdLst; Apple's importer
    # (Quick Look, Keynote) then fails to open slides that have speaker notes.
    prs = render({"title": "t", "slides": [{**bullets(), "notes": "n"}]})
    pres = prs.part._element
    listed = pres.findall(f"{qn('p:notesMasterIdLst')}/{qn('p:notesMasterId')}")
    rel_ids = {r.rId for r in prs.part.rels.values() if r.reltype == RT.NOTES_MASTER}
    assert len(listed) == 1 and listed[0].get(qn("r:id")) in rel_ids
    tags = [c.tag for c in pres]
    assert tags.index(qn("p:notesMasterIdLst")) == tags.index(qn("p:sldMasterIdLst")) + 1


# ------------------------------------------------------------------ charts
def summary(groups):
    return {"groups": [{"key": k, "totals_by_currency": by} for k, by in groups]}


Q = {"group_by": "category", "start_date": "2026-10-01", "end_date": "2026-10-31"}


def test_top_groups_plus_other_per_currency():
    s = summary(
        [
            ("วัตถุดิบ", {"THB": 9000}),
            ("ค่าเดินทาง", {"THB": 250.5}),
            ("สาธารณูปโภค", {"THB": 120}),
            ("อุปกรณ์", {"USD": 10}),
        ]
    )
    cd = chart_data(s, ChartQuery.model_validate({**Q, "top": 2}), "th")
    assert cd.categories == ["วัตถุดิบ", "ค่าเดินทาง", "อื่นๆ"]
    assert cd.series == {"THB": [9000, 250.5, 120], "USD": [0, 0, 10]}


def test_day_groups_keep_date_order_and_no_other():
    s = summary([("2026-10-01", {"THB": 5}), ("2026-10-02", {"THB": 7})])
    q = {**Q, "group_by": "day", "top": 1, "end_date": "2026-10-02"}
    cd = chart_data(s, ChartQuery.model_validate(q), "en")
    assert cd.categories == ["10-01", "10-02"] and cd.series == {"THB": [5, 7]}


def test_missing_keys_and_empty_data():
    q = ChartQuery.model_validate(Q)
    assert chart_data(summary([(None, {"THB": 1})]), q, "th").categories == ["—"]
    assert chart_data(summary([]), q, "th").empty


def test_chart_slide_draws_a_native_chart_with_the_given_numbers():
    spec = {"title": "t", "slides": [{"kind": "chart", "title": "ค่าใช้จ่ายตามหมวด", "chart": Q}]}
    cd = ChartData(categories=["วัตถุดิบ", "อื่นๆ"], series={"THB": [9000.0, 370.5]})
    prs = render(spec, charts={0: cd})
    chart = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
    assert list(chart.plots[0].categories) == ["วัตถุดิบ", "อื่นๆ"]
    assert list(chart.plots[0].series[0].values) == [9000.0, 370.5]


@pytest.mark.parametrize(
    "style, expected",
    [
        ("bar", XL_CHART_TYPE.COLUMN_CLUSTERED),
        ("line", XL_CHART_TYPE.LINE_MARKERS),
        ("pie", XL_CHART_TYPE.PIE),
    ],
)
def test_chart_styles(style, expected):
    spec = {
        "title": "t",
        "slides": [{"kind": "chart", "title": "c", "chart": {**Q, "style": style}}],
    }
    prs = render(spec, charts={0: ChartData(["a", "b"], {"THB": [1.0, 2.0]})})
    assert next(s.chart for s in prs.slides[1].shapes if s.has_chart).chart_type == expected


def test_empty_chart_shows_the_placeholder():
    spec = {"title": "t", "slides": [{"kind": "chart", "title": "c", "chart": Q}]}
    prs = render(spec, charts={0: ChartData([], {})})
    assert not any(s.has_chart for s in prs.slides[1].shapes)
    assert "ไม่มีข้อมูลในช่วงนี้" in " ".join(texts(prs.slides[1]))


# ------------------------------------------------------------------ review fixes
def _chart_of(prs, slide=1):
    return next(s.chart for s in prs.slides[slide].shapes if s.has_chart)


def test_xml_invalid_characters_anywhere_do_not_break_the_deck():
    spec = {
        "title": "t￿",
        "slides": [
            {"kind": "bullets", "title": "t", "bullets": ["a￿b"]},
            {"kind": "chart", "title": "c", "chart": Q},
        ],
    }
    cd = ChartData(["Mak\x0bro", "￿", "\ud800ok"], {"TH\x00B": [1.0, 2.0, 3.0]})
    prs = render(spec, charts={1: cd})
    chart = _chart_of(prs, 2)
    assert list(chart.plots[0].categories) == ["Makro", "—", "ok"]
    assert chart.plots[0].series[0].name == "THB"
    assert "ab" in " ".join(texts(prs.slides[1]))


def test_a_real_other_category_folds_into_the_other_bucket():
    s = summary(
        [
            ("วัตถุดิบ", {"THB": 10000}),
            ("อื่นๆ", {"THB": 5000}),
            ("ค่าเช่า", {"THB": 4000}),
            ("ค่าแรง", {"THB": 3000}),
            ("การตลาด", {"THB": 1000}),
        ]
    )
    cd = chart_data(s, ChartQuery.model_validate({**Q, "top": 3}), "th")
    assert cd.categories == ["วัตถุดิบ", "ค่าเช่า", "ค่าแรง", "อื่นๆ"]
    assert cd.series == {"THB": [10000, 4000, 3000, 6000]}


def test_pie_slices_get_distinct_colours():
    cats = [f"c{i}" for i in range(11)]  # top 10 + Other
    spec = {
        "title": "t",
        "slides": [{"kind": "chart", "title": "c", "chart": {**Q, "style": "pie"}}],
    }
    prs = render(spec, charts={0: ChartData(cats, {"THB": [float(i + 1) for i in range(11)]})})
    points = _chart_of(prs).plots[0].series[0].points
    assert len({str(points[i].format.fill.fore_color.rgb) for i in range(11)}) == 11


def test_day_charts_have_one_category_per_day_in_the_range():
    s = summary([("2026-10-01", {"THB": 5}), ("2026-10-03", {"THB": 7})])
    q = {"group_by": "day", "start_date": "2026-10-01", "end_date": "2026-10-04"}
    cd = chart_data(s, ChartQuery.model_validate(q), "en")
    assert cd.categories == ["10-01", "10-02", "10-03", "10-04"]
    assert cd.series == {"THB": [5, 0, 7, 0]}


def test_a_multi_currency_pie_falls_back_to_bars():
    spec = {
        "title": "t",
        "slides": [{"kind": "chart", "title": "c", "chart": {**Q, "style": "pie"}}],
    }
    cd = ChartData(["SaaS", "Makro"], {"THB": [0.0, 500.0], "USD": [900.0, 0.0]})
    chart = _chart_of(render(spec, charts={0: cd}))
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED and len(chart.plots[0].series) == 2


def test_charts_have_no_automatic_title():
    spec = {"title": "t", "slides": [{"kind": "chart", "title": "c", "chart": Q}]}
    chart = _chart_of(render(spec, charts={0: ChartData(["a"], {"THB": [1.0]})}))
    assert chart._chartSpace.chart.autoTitleDeleted.val is True
