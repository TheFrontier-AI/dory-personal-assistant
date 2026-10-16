import io
import time

import pytest
from PIL import Image, ImageDraw, ImageFont

from dory.config import Settings
from dory.extract import extract_receipt
from dory.llm import make_client

pytestmark = pytest.mark.live
GUILD = 1000


def make_receipt_png() -> bytes:
    img = Image.new("RGB", (520, 560), "white")
    d = ImageDraw.Draw(img)
    try:
        f = ImageFont.load_default(size=22)
    except TypeError:
        f = ImageFont.load_default()
    lines = [
        "MAKRO SUKHUMVIT BRANCH",
        "TAX INVOICE / RECEIPT",
        "Date: 12/10/2026",
        "-" * 32,
        "Ice 10kg x 2        100.00",
        "Pork belly 3kg      540.00",
        "Cooking oil 5L      325.00",
        "-" * 32,
        "Subtotal           965.00",
        "VAT 7%             67.55",
        "TOTAL THB        1032.55",
        "Paid: CASH",
    ]
    for i, t in enumerate(lines):
        d.text((20, 20 + i * 38), t, fill="black", font=f)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def test_live_extract():
    s = Settings.from_env()
    t = time.monotonic()
    ex = extract_receipt(make_client(s), s, make_receipt_png(), "image/png", "ค่าวัตถุดิบ สาขา 2")
    print(f"\n{time.monotonic() - t:.1f}s\n{ex.model_dump_json(indent=2, ensure_ascii=False)}")
    assert ex.is_receipt and "makro" in (ex.vendor or "").lower()
    assert ex.total == pytest.approx(1032.55, abs=0.01)
    assert ex.receipt_date == "2026-10-12"
    assert ex.category == "วัตถุดิบ"


def test_live_firestore_store_and_gcs_archive():
    """Real Firestore `(default)` + the existing audio bucket, in throwaway names only.

    Writes only to `dory_receipts_test_<random>` (+ its `_meta` counter) and to
    `dory-receipts/test-<random>/` in the bucket; everything created is deleted in `finally`.
    """
    import uuid

    from google.cloud import firestore

    from dory.archive import ReceiptArchive
    from dory.firestore_store import FirestoreStore
    from dory.models import Receipt

    s = Settings.from_env()
    token = uuid.uuid4().hex[:10]
    coll = f"dory_receipts_test_{token}"
    client = firestore.Client(project=s.gcp_project)
    store = FirestoreStore(client=client, collection=coll)
    bucket = "dory-personal-assistant-audio"
    arch = ReceiptArchive(bucket, f"dory-receipts/test-{token}/", project=s.gcp_project)
    uri = None
    try:

        def mk(n, uid, name, vendor, day, total, conf=0.9):
            return store.add_receipt(
                Receipt(
                    guild_id=GUILD, submitter_id=uid, submitter_name=name, sha256=f"{token}-{n}", vendor=vendor,
                    receipt_date=day, effective_date=day, total=total, category="วัตถุดิบ",
                    confidence=conf, created_at=day,
                )
            )  # fmt: skip

        a = mk(1, 100, "Alice", "Makro", "2026-10-01", 1000.0)
        b = mk(2, 200, "Bob", "Makro", "2026-10-05", 9000.0)
        c = mk(3, 200, "Bob", "Ice", "2026-10-14", 120.0, conf=0.4)
        assert (a.id, b.id, c.id) == (1, 2, 3)  # transactional counter
        assert store.get_by_sha(GUILD, f"{token}-2").id == 2 and store.get(3).vendor == "Ice"
        with pytest.raises(ValueError):
            mk(2, 200, "Bob", "Makro", "2026-10-05", 9000.0)  # same sha -> rejected
        assert store.summarize(GUILD, "2026-10-01", "2026-10-31")["total"] == 10120.0
        assert (
            store.summarize(GUILD, "2026-10-01", "2026-10-31", submitter_id=100)["total"] == 1000.0
        )
        assert [r.vendor for r in store.flagged(GUILD)] == ["Ice"]
        assert store.find_possible_duplicate(a.model_copy(update={"sha256": "x"})).id == 1
        store.set_reply_message(1, 42)
        assert store.get_by_reply_message(42).id == 1
        assert store.update_receipt(1, {"total": 5.0}).total == 5.0

        uri = arch.save(f"{token}-obj", b"hello", "application/pdf")
        assert uri == f"gs://{bucket}/dory-receipts/test-{token}/{token}-obj.pdf"
        blob = arch.client.bucket(bucket).blob(uri.split(f"{bucket}/", 1)[1])
        assert blob.download_as_bytes() == b"hello"
    finally:
        for name in (coll, f"{coll}_meta"):
            for snap in client.collection(name).stream():
                snap.reference.delete()
        if uri:
            arch.delete(uri)


def test_embedding_probe():
    """One-off probe for memory search: model availability in the configured location, the batch
    limit, the dimension, and Thai/English similarity (prints the numbers used to calibrate
    dory.embed.EMBED_BATCH and qa.MIN_SIMILARITY)."""
    from dory.embed import VertexEmbedder, cosine

    s = Settings()
    texts = ["งบโฆษณาเดือนนี้ 50,000 บาท", "marketing budget for this month", "ซื้อน้ำแข็ง 3 ถุง"]
    client = make_client(s)
    try:
        vecs = VertexEmbedder(client, s.embed_model, 256, batch=3).embed(texts, "document")
        batch_ok = True
    except Exception as e:  # noqa: BLE001  (the probe reports which batch size works)
        print(f"batch of 3 failed: {type(e).__name__}: {e}")
        vecs = VertexEmbedder(client, s.embed_model, 256, batch=1).embed(texts, "document")
        batch_ok = False
    th, en, ice = vecs
    q = VertexEmbedder(client, s.embed_model, 256).embed(["how big is the ads budget?"], "query")[0]
    print(
        f"model={s.embed_model} location={s.embed_location or s.gemini_location} "
        f"batch3_ok={batch_ok} dims={[len(v) for v in vecs]}\n"
        f"cos(th,en)={cosine(th, en):.3f} cos(th,ice)={cosine(th, ice):.3f} "
        f"cos(en,ice)={cosine(en, ice):.3f} cos(q,th)={cosine(q, th):.3f} "
        f"cos(q,ice)={cosine(q, ice):.3f}"
    )
    assert all(len(v) == 256 for v in vecs)
    assert cosine(th, en) > cosine(th, ice) and cosine(q, th) > cosine(q, ice)


def test_live_memory_store_queries():
    """Real Firestore for the memory store's queries (three equality filters in messages_after,
    the guild-wide all_summaries, bytes vectors), in throwaway `dory_*_test_<random>` collections
    that are deleted in `finally`."""
    import uuid
    from datetime import UTC, datetime

    from google.cloud import firestore

    from dory.firestore_memory import FirestoreMemoryStore
    from dory.memory import build_message
    from dory.memory_store import DaySummary

    s = Settings.from_env()
    token = uuid.uuid4().hex[:10]
    names = {k: f"dory_{k}_test_{token}" for k in ("messages", "summaries", "channels", "guilds")}
    client = firestore.Client(project=s.gcp_project)
    store = FirestoreMemoryStore(client=client, project=s.gcp_project, **names)
    try:
        at = datetime(2026, 10, 14, 3, 0, tzinfo=UTC)
        for mid, cid in ((1, 10), (2, 11), (3, 10)):
            store.add_message(build_message(
                s, guild_id=GUILD, channel_id=cid, message_id=mid, author_id=1, author_name="a",
                created_at=at.replace(minute=mid), text=f"m{mid}",
            ))  # fmt: skip
        assert [m.message_id for m in store.messages_after(GUILD, 10, "2026-10-14")] == [1, 3]
        got = store.messages(GUILD, {10, 11}, "2026-10-13", "2026-10-14", newest_first=False)
        assert [m.message_id for m in got] == [1, 2, 3]
        store.put_summary(DaySummary(guild_id=GUILD, channel_id=10, local_date="2026-10-14",
                                     summary={"overview": "x"}, updated_at=""))  # fmt: skip
        assert store.put_summary_vectors(GUILD, 10, "2026-10-14", "m@2", ["h"], b"\x00" * 8)
        [summary] = store.all_summaries(GUILD)
        assert summary.vec_model == "m@2" and summary.vecs == b"\x00" * 8
    finally:
        for name in names.values():
            for snap in client.collection(name).stream():
                snap.reference.delete()


def test_live_deck_request(tmp_path):
    """Real Gemini builds a spending deck through create_deck (also proves Gemini accepts the
    tool's nested schema). Local throwaway SQLite ledger only; nothing is posted or stored."""
    from pptx import Presentation

    from dory.models import Receipt
    from dory.qa import ask
    from dory.store import SQLiteStore

    store = SQLiteStore(str(tmp_path / "ledger.db"))
    for n, (vendor, day, total, cat) in enumerate(
        [("Makro", "2026-10-03", 4200.0, "วัตถุดิบ"), ("Grab", "2026-10-05", 350.0, "ค่าเดินทาง")]
    ):
        store.add_receipt(
            Receipt(
                guild_id=GUILD, submitter_id=1, submitter_name="Exec", sha256=f"live-deck-{n}",
                vendor=vendor, receipt_date=day, effective_date=day, total=total, category=cat,
                confidence=0.95, created_at=day,
            )
        )  # fmt: skip
    ans = ask(
        Settings.from_env(),
        store,
        "ทำสไลด์สรุปค่าใช้จ่ายเดือนตุลาคม 2026 แยกตามหมวด",
        guild_id=GUILD,
        is_exec=True,
        user_id=1,
        deck=True,
    )
    print(ans.text, ans.tool_calls)
    [(name, data)] = ans.files
    prs = Presentation(io.BytesIO(data))
    assert name.endswith(".pptx") and len(prs.slides) >= 2
    assert any(s.has_chart for slide in prs.slides for s in slide.shapes)
