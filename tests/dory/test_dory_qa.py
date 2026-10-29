import io
import json
from datetime import datetime
from zoneinfo import ZoneInfo

import pytest
from conftest import ALICE, BOB, EXEC_ID, GUILD, fc_resp, scripted, text_resp, usage
from pptx import Presentation

from dory.digest import build_daily_digest
from dory.qa import (
    ACTION_RULE,
    CAPABILITIES_PROMPT,
    CLAIM_CHECK,
    MAX_ROUNDS,
    MEMORY_PROMPT,
    NO_REPLY,
    LedgerTools,
    ask,
    claims_done,
    stable_prompt,
    system_prompt,
)


def test_aggregation_math(seeded):
    t = LedgerTools(seeded, GUILD)
    s = t.spending_summary("2026-10-01", "2026-10-31", "category")
    assert s["total"] == 10370.5 and s["count"] == 4
    assert {g["key"]: g["total"] for g in s["groups"]} == {"วัตถุดิบ": 10120.0, "ค่าเดินทาง": 250.5}
    assert s["total_fmt"] == "฿10,370.50"
    v = t.spending_summary("2026-10-01", "2026-10-31", "vendor")["groups"]
    assert v[0] == {
        "key": "Makro",
        "total": 10000.0,
        "totals_by_currency": {"THB": 10000.0},
        "count": 2,
        "total_fmt": "฿10,000",
    }
    assert [g["key"] for g in t.spending_summary("2026-10-01", "2026-10-31", "day")["groups"]] == [
        "2026-10-01",
        "2026-10-05",
        "2026-10-14",
    ]
    assert t.spending_summary("2026-10-01", "2026-10-31", "submitter")["groups"][0]["key"] == "Bob"
    assert t.spending_summary("2030-01-01", "2030-01-31")["total"] == 0


def test_compare_periods(seeded):
    c = LedgerTools(seeded, GUILD).compare_periods(
        "2026-10-01", "2026-10-31", "2026-09-01", "2026-09-30", "category"
    )
    assert c["overall"]["diff"] == 7370.5 and c["overall"]["pct_change_vs_b"] == 245.7
    util = next(g for g in c["groups"] if g["key"] == "สาธารณูปโภค")
    assert util["a"] == 0 and util["b"] == 3000 and util["diff"] == -3000


def test_list_and_flagged(seeded):
    t = LedgerTools(seeded, GUILD)
    assert t.list_receipts("2026-10-01", "2026-10-31", vendor_contains="mak")["count"] == 2
    assert (
        t.list_receipts("2026-10-01", "2026-10-31", category="ค่าเดินทาง")["receipts"][0]["vendor"]
        == "Grab"
    )
    assert t.list_receipts("2026-01-01", "2026-12-31", limit=2)["count"] == 2
    assert [r["vendor"] for r in t.flagged_receipts()["receipts"]] == ["Ice Co"]


def test_non_exec_scoped_in_every_tool(seeded):
    t = LedgerTools(seeded, GUILD, forced_submitter_id=ALICE)
    assert t.spending_summary("2026-01-01", "2026-12-31")["total"] == 1250.5
    assert t.spending_summary("2026-01-01", "2026-12-31", "submitter")["groups"] == [
        {
            "key": "Alice",
            "total": 1250.5,
            "totals_by_currency": {"THB": 1250.5},
            "count": 2,
            "total_fmt": "฿1,250.50",
        }
    ]
    assert {
        r["submitter"]
        for r in t.list_receipts("2026-01-01", "2026-12-31", submitter="Bob")["receipts"]
    } == set()
    assert t.list_receipts("2026-01-01", "2026-12-31", vendor_contains="Makro")["count"] == 1
    c = t.compare_periods("2026-10-01", "2026-10-31", "2026-09-01", "2026-09-30")
    assert c["period_a"]["total"] == 1250.5 and c["period_b"]["total"] == 0
    assert t.flagged_receipts()["count"] == 0  # Bob's low-confidence receipt is invisible to Alice


def test_model_cannot_override_scope(seeded):
    t = LedgerTools(seeded, GUILD, forced_submitter_id=ALICE)
    r = t.call(
        "spending_summary",
        {
            "start_date": "2026-01-01",
            "end_date": "2026-12-31",
            "submitter_id": BOB,
            "forced_submitter_id": BOB,
        },
    )
    assert r["total"] == 1250.5
    assert "error" in t.call("drop_table", {})


def _tool_results(llm):
    """Every function_response the model was shown, across all requests, as one string."""
    out = []
    for req in llm.requests:
        for c in req.contents:
            out += [str(p.function_response) for p in c.parts or [] if p.function_response]
    return json.dumps(out, ensure_ascii=False)


EVERYTHING = {"start_date": "2026-01-01", "end_date": "2026-12-31"}


def test_non_exec_cannot_see_others_even_if_model_asks_for_everything(seeded, settings):
    """A hostile/confused model asks for all data, other users, and a forged submitter scope."""
    llm = scripted(
        fc_resp(
            ("spending_summary", {**EVERYTHING, "group_by": "submitter", "submitter_id": BOB}),
            ("list_receipts", {**EVERYTHING, "submitter": "Bob", "limit": 100}),
            (
                "list_receipts",
                {**EVERYTHING, "vendor_contains": "Makro", "forced_submitter_id": BOB},
            ),
            (
                "compare_periods",
                {
                    "a_start": "2026-10-01",
                    "a_end": "2026-10-31",
                    "b_start": "2026-09-01",
                    "b_end": "2026-09-30",
                    "submitter_id": 0,
                },
            ),
            ("flagged_receipts", {"submitter_id": BOB}),
        ),
        text_resp("ok"),
    )
    a = ask(
        settings, seeded, "ใครเบิกเยอะสุด", guild_id=GUILD, is_exec=False, user_id=ALICE, model=llm
    )
    blob = _tool_results(llm)
    assert a.text == "ok" and len(a.tool_calls) == 5
    assert "Alice" in blob and "1250.5" in blob
    for secret in ("Bob", "9000", "3000", "13370", "Ice Co", "PEA"):
        assert secret not in blob, secret
    prompt = llm.requests[0].config.system_instruction
    assert "ไม่ใช่ผู้บริหาร" in prompt and "2026-10-14" in prompt and "Asia/Bangkok" in prompt
    # no tool schema offers a submitter scope parameter
    for t in llm.requests[0].config.tools:
        for d in t.function_declarations:
            schema = d.parameters_json_schema or (d.parameters and d.parameters.model_dump()) or {}
            props = set(schema.get("properties") or {})
            assert d.name != "spending_summary" or "group_by" in props  # schema is really read
            assert not {"submitter_id", "forced_submitter_id", "user_id"} & props


def test_exec_sees_everything(seeded, settings):
    llm = scripted(fc_resp(("spending_summary", EVERYTHING)), text_resp("รวม ฿13,370.50"))
    a = ask(
        settings, seeded, "ทั้งหมดเท่าไร", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm
    )
    assert a.text == "รวม ฿13,370.50" and a.rounds == 1 and a.tool_calls[0][0] == "spending_summary"
    assert "13370.5" in _tool_results(llm)
    assert "ไม่ใช่ผู้บริหาร" not in llm.requests[0].config.system_instruction


def test_tools_are_the_ledger_tools(seeded, settings):
    llm = scripted(text_resp("ok"))
    ask(settings, seeded, "x", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm)
    names = {d.name for t in llm.requests[0].config.tools for d in t.function_declarations}
    assert names == {
        "spending_summary",
        "list_receipts",
        "compare_periods",
        "flagged_receipts",
        "find_duplicates",
        "void_receipts",
        "restore_receipts",
    }


def test_bad_tool_args_return_error_to_model(seeded, settings):
    llm = scripted(
        fc_resp(("spending_summary", {**EVERYTHING, "group_by": "nonsense"})), text_resp("ขออภัย")
    )
    a = ask(settings, seeded, "x", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm)
    assert a.text == "ขออภัย" and "bad group_by" in _tool_results(llm)


def test_a_crashing_tool_is_reported_to_the_model_and_the_turn_goes_on(
    seeded, settings, monkeypatch, caplog
):
    """A Firestore / network error inside a tool used to abort the whole answer."""

    def down(self, **kw):
        raise ConnectionError("firestore unavailable")

    monkeypatch.setattr(LedgerTools, "spending_summary", down)
    llm = scripted(fc_resp(("spending_summary", EVERYTHING)), text_resp("ตอนนี้ดึงยอดไม่ได้ค่ะ"))
    with caplog.at_level("WARNING", logger="dory.qa"):
        a = ask(settings, seeded, "ยอดปีนี้", guild_id=GUILD, is_exec=True, user_id=EXEC_ID,
                model=llm)  # fmt: skip
    assert a.text == "ตอนนี้ดึงยอดไม่ได้ค่ะ" and a.rounds == 1
    shown = _tool_results(llm)
    assert "internal error" in shown and "firestore unavailable" not in shown  # no internals
    assert "spending_summary" in a.tool_ms  # still timed
    (rec,) = [r for r in caplog.records if r.name == "dory.qa" and r.exc_info]
    assert "spending_summary" in rec.getMessage()
    assert "2026-01-01" not in rec.getMessage() and "ยอดปีนี้" not in rec.getMessage()


def test_round_cap(seeded, settings):
    llm = scripted(*[fc_resp(("flagged_receipts", {}))] * MAX_ROUNDS, text_resp("สรุป"))
    a = ask(settings, seeded, "x", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm)
    assert a.rounds == MAX_ROUNDS and a.text == "สรุป"
    assert len(llm.requests) == MAX_ROUNDS + 1
    assert not llm.requests[-1].config.tools  # the final call has no tools


def test_digest(seeded):
    seeded_day = "2026-10-05"
    d = build_daily_digest(seeded, GUILD, seeded_day, "th")
    assert "฿9,250.5" in d and "Makro" in d and "วัตถุดิบ" in d and "ค่าเดินทาง" in d
    d2 = build_daily_digest(seeded, GUILD, "2026-10-14", "th")
    assert "Ice Co" in d2 and "ต้องตรวจสอบ" in d2
    assert "ไม่มีใบเสร็จ" in build_daily_digest(seeded, GUILD, "2026-01-01", "th")


def test_digest_top3(store):
    from conftest import mk

    for i, (v, t) in enumerate([("A", 1), ("B", 5), ("C", 3), ("D", 4)]):
        mk(store, i, 1, "x", v, "2026-10-01", t)
    d = build_daily_digest(store, GUILD, "2026-10-01", "th")
    assert "1. B" in d and "2. D" in d and "3. C" in d and "A:" not in d.split("3 อันดับ")[1]


def test_is_exec(settings):
    assert settings.is_exec(EXEC_ID) and not settings.is_exec(5, ["staff"])
    assert settings.is_exec(5, ["ผู้บริหาร"]) and settings.is_exec(5, ["Executive"])


def test_system_prompt_language_instruction(settings):
    from dory.qa import system_prompt

    th = system_prompt(settings, scoped=False, language="th")
    assert "พูดเหมือนคนจริงๆ" in th and "\nLanguage:" not in th
    assert "ภาษาหลักของเซิร์ฟเวอร์นี้คือภาษาไทย" in th
    en = system_prompt(settings, scoped=False)  # English is the default
    assert en == system_prompt(settings, scoped=False, language="en")
    assert (
        "reply in English" in en
        and "\nLanguage: this server's main language is English" in en
        and "พูดเหมือนคนจริงๆ ไม่ใช่หุ่นยนต์: สั้น" not in en
    )
    auto = system_prompt(settings, scoped=True, language="auto")
    assert (
        "same language as the user's latest message" in auto
        and "English when you cannot tell" in auto
        and "ไม่ใช่ผู้บริหาร" in auto
    )


def test_system_prompt_lists_capabilities(settings):
    from dory.qa import system_prompt

    for lang in ("th", "en", "auto"):
        p = system_prompt(settings, scoped=True, language=lang)
        for cmd in (
            "/โหมด",
            "/mode",
            "/ภาษา",
            "/language",
            "/ความจำ",
            "/memory",
            "/join-meeting",
            "/leave-meeting",
        ):
            assert cmd in p
        assert "งานใหญ่:" in p and "ค้นอินเทอร์เน็ตไม่ได้" not in p and "Google Search" in p


def test_ask_passes_the_language_to_the_model(seeded, settings):
    llm = scripted(text_resp("Total is ฿1,000"))
    a = ask(
        settings, seeded, "total?", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm,
        language="en",
    )  # fmt: skip
    assert a.text == "Total is ฿1,000"
    assert "reply in English" in llm.requests[0].config.system_instruction


def test_empty_answer_falls_back_in_the_users_language(seeded, settings):
    def empty(question, language):
        llm = scripted(text_resp(""))
        return ask(
            settings, seeded, question, guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm,
            language=language,
        ).text  # fmt: skip

    assert "ขออภัย" in empty("ยอดเท่าไร", "th")
    assert "Sorry" in empty("total?", "en")
    assert "Sorry" in empty("total?", "auto") and "ขออภัย" in empty("ยอดเท่าไร", "auto")


def test_digest_in_english(seeded):
    d = build_daily_digest(seeded, GUILD, "2026-10-14", "en")
    assert "Daily expense summary" in d and "Needs review" in d and "unclear data" in d
    assert "สรุปค่าใช้จ่าย" not in d
    assert "No receipts" in build_daily_digest(seeded, GUILD, "2026-01-01", "en")


def test_spoken_hint_asks_for_a_short_answer_without_markdown(seeded, settings):
    store = seeded
    for spoken in (True, False):
        llm = scripted(text_resp("ok"))
        ask(settings, store, "สรุปให้หน่อย", guild_id=GUILD, is_exec=True, user_id=EXEC_ID,
            model=llm, spoken=spoken)  # fmt: skip
        text = "".join(p.text or "" for c in llm.requests[0].contents for p in c.parts or [])
        assert ("1-3 short sentences" in text) is spoken and (
            "no lists or markdown" in text
        ) is spoken


def test_a_model_that_keeps_calling_tools_after_the_cap_gets_an_answer_not_a_crash(
    seeded, settings
):
    """Real incident: after MAX_ROUNDS the tools are withdrawn but the model still emitted
    chat_search; ADK raised "Tool 'chat_search' not found. Available tools: " (empty)."""
    llm = scripted(*[fc_resp(("flagged_receipts", {}))] * (MAX_ROUNDS + 1), text_resp("สรุป"))
    a = ask(settings, seeded, "x", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm)
    assert a.text == "สรุป"
    assert not llm.requests[MAX_ROUNDS].config.tools  # tools withdrawn on the call after the cap
    assert "Stop using tools now" in _tool_results(llm)  # the stray call was answered, not run
    assert len([c for c in a.tool_calls if c[0] == "flagged_receipts"]) == MAX_ROUNDS


# ------------------------------------------------------------------ metrics
def test_answer_reports_tokens_and_tool_time(settings, seeded):
    llm = scripted(
        usage(
            fc_resp(("spending_summary", {"start_date": "2026-10-01", "end_date": "2026-10-14"})),
            prompt=100, cached=40, out=10, thought=5,
        ),
        usage(text_resp("ยอด ฿10,370.50"), prompt=120, cached=60, out=20),
    )  # fmt: skip
    a = ask(settings, seeded, "เดือนนี้เท่าไร", guild_id=GUILD, is_exec=True, user_id=EXEC_ID,
            model=llm)  # fmt: skip
    assert (a.model_calls, a.tokens_in, a.tokens_cached, a.tokens_out, a.tokens_thought) == (
        2, 220, 100, 30, 5,
    )  # fmt: skip
    assert set(a.tool_ms) == {"spending_summary"} and a.tool_ms["spending_summary"] >= 0


def test_answer_without_usage_metadata_reports_zero_tokens(settings, seeded):
    a = ask(settings, seeded, "hi", guild_id=GUILD, is_exec=True, user_id=EXEC_ID,
            model=scripted(text_resp("hi")))  # fmt: skip
    assert (a.model_calls, a.tokens_in, a.tool_ms) == (1, 0, {})


# ------------------------------------------------------------------ prompt caching
@pytest.mark.parametrize("lang", ["th", "en", "auto"])
def test_prompt_prefix_is_identical_across_times_and_askers(settings, lang):
    t1 = datetime(2026, 10, 14, 9, 1, tzinfo=ZoneInfo("Asia/Bangkok"))
    t2 = datetime(2026, 10, 16, 18, 47, tzinfo=ZoneInfo("Asia/Bangkok"))
    a = system_prompt(settings, scoped=False, now=t1, language=lang, memory=True)
    b = system_prompt(settings, scoped=True, now=t2, language=lang, memory=True, interjection=True)
    stable = stable_prompt(lang, memory=True)
    assert a.startswith(stable) and b.startswith(stable)
    assert stable.endswith(MEMORY_PROMPT) and CAPABILITIES_PROMPT in stable
    assert "2026-10-14" not in stable and a.index("2026-10-14") > len(stable)
    assert "ผู้ถามไม่ใช่ผู้บริหาร" in b[len(stable) :] and NO_REPLY in b[len(stable) :]


def test_one_action_rule_for_every_language_and_the_statuses_it_names_exist():
    """Text and voice share one rule for tools, failures, sources and suggestions; each "done
    before" status it names is really returned by its tool (so the rule never drifts)."""
    import importlib
    import inspect

    for lang in ("th", "en", "ja"):
        assert ACTION_RULE in stable_prompt(lang, memory=False)
    for status, module in (
        ("already_set", "reminders"),
        ("already_booked", "trips"),
        ("already_scheduled", "secretary"),
        ("already_on_the_list", "coo_tools"),
    ):
        assert status in ACTION_RULE
        assert f'"{status}"' in inspect.getsource(importlib.import_module(f"dory.{module}"))


# ------------------------------------------------------------------ slide decks
DECK = {
    "title": "สรุปค่าใช้จ่าย",
    "slides": [
        {
            "kind": "chart",
            "title": "ตามหมวด",
            "chart": {"group_by": "category", "start_date": "2026-10-01", "end_date": "2026-10-31"},
        }
    ],
}


def deck_ask(settings, store, *script, is_exec=True, user_id=EXEC_ID,
             question="ทำสไลด์สรุปค่าใช้จ่ายเดือนนี้", language="th"):  # fmt: skip
    llm = scripted(*script)
    a = ask(settings, store, question, guild_id=GUILD, is_exec=is_exec, user_id=user_id,
            model=llm, deck=True, language=language)  # fmt: skip
    return llm, a


def chart_values(data):
    prs = Presentation(io.BytesIO(data))
    chart = next(s.chart for sl in prs.slides for s in sl.shapes if s.has_chart)
    return dict(zip(chart.plots[0].categories, chart.plots[0].series[0].values, strict=True))


def test_create_deck_attaches_a_pptx_with_ledger_numbers(settings, seeded):
    _, a = deck_ask(settings, seeded, fc_resp(("create_deck", DECK)), text_resp("ทำสไลด์ให้แล้วครับ"))
    [(name, data)] = a.attachments
    assert name.endswith(".pptx") and data[:2] == b"PK"
    assert chart_values(data) == {"วัตถุดิบ": 10120.0, "ค่าเดินทาง": 250.5}


def test_a_non_executive_chart_covers_only_their_own_receipts(settings, seeded):
    _, a = deck_ask(settings, seeded, fc_resp(("create_deck", DECK)), text_resp("ok"),
                    is_exec=False, user_id=ALICE)  # fmt: skip
    assert chart_values(a.attachments[0][1]) == {"วัตถุดิบ": 1000.0, "ค่าเดินทาง": 250.5}


def test_an_invalid_outline_returns_an_error_and_a_fixed_call_succeeds(settings, seeded):
    bad = {"title": "t", "slides": [{"kind": "chart", "title": "c"}]}
    llm, a = deck_ask(settings, seeded, fc_resp(("create_deck", bad)),
                      fc_resp(("create_deck", DECK)), text_resp("ok"))  # fmt: skip
    first = [
        p.function_response.response
        for c in llm.requests[1].contents
        for p in c.parts or []
        if p.function_response
    ][0]
    blob = json.dumps(first, ensure_ascii=False)
    assert "chart" in blob and "error" in blob
    assert len(a.attachments) == 1


def test_invalid_outline_text_never_reaches_the_logs(settings, seeded, caplog):
    bad = {"title": "t", "slides": [{"kind": "chart", "title": "c", "notes": "ลับมาก"}]}
    with caplog.at_level("DEBUG"):
        deck_ask(settings, seeded, fc_resp(("create_deck", bad)), text_resp("ok"))
    assert not [r.name for r in caplog.records if "ลับมาก" in r.getMessage()]


def test_a_second_create_deck_replaces_the_first(settings, seeded):
    other = {"title": "Second", "slides": [{"kind": "bullets", "title": "b", "bullets": ["x"]}]}
    _, a = deck_ask(settings, seeded, fc_resp(("create_deck", DECK)),
                    fc_resp(("create_deck", other)), text_resp("ok"))  # fmt: skip
    assert [n.split("-")[0] for n, _ in a.attachments] == ["second"]


def test_deck_fixed_strings_follow_the_reply_language(settings, seeded):
    _, a = deck_ask(settings, seeded, fc_resp(("create_deck", DECK)), text_resp("done"),
                    question="make slides of this month's spending", language="auto")  # fmt: skip
    cover = Presentation(io.BytesIO(a.attachments[0][1])).slides[0]
    words = " ".join(s.text_frame.text for s in cover.shapes if s.has_text_frame)
    assert "Made with Dory" in words


def test_no_deck_tool_without_deck_and_prompt_only_with_it(settings, seeded):
    llm = scripted(text_resp("hi"))
    ask(settings, seeded, "hi", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm)
    names = {d.name for t_ in llm.requests[0].config.tools for d in t_.function_declarations}
    assert "create_deck" not in names
    assert "create_deck" not in llm.requests[0].config.system_instruction
    llm2, _ = deck_ask(settings, seeded, text_resp("hi"))
    decls = {d.name: d for t_ in llm2.requests[0].config.tools for d in t_.function_declarations}
    assert '"kind"' in json.dumps(decls["create_deck"].parameters_json_schema)  # nested schema
    assert "create_deck" in llm2.requests[0].config.system_instruction


# ----------------------------------------------------------------------------- "done" with no tool (DEF-07)
@pytest.mark.parametrize(
    "text",
    ["I've saved that: passport name SOMCHAI/TEST, aisle seat preference.",
     "Booked! Your receipts are filed under Travel.", "I have gone ahead and booked both for you.",
     "Your profile has been saved.", "I've set a reminder for 9am.", "บันทึกไว้แล้วค่ะ",
     "จองให้เรียบร้อยแล้วค่ะ", "ตั้งเตือนให้แล้วค่ะ"],
)  # fmt: skip
def test_an_answer_that_says_it_did_something_is_a_claim(text):
    assert claims_done(text)


@pytest.mark.parametrize(
    "text",
    ["Shall I book it?", "I haven't booked anything yet.", "Was it booked?",
     "Once the trip is booked, the receipts will be filed.", "I can save that if you like.",
     "What I can do is set a 3am reminder for you.", "I've set out two options below.",
     "I've updated the numbers below.", "I've added up the totals.", "ยังไม่ได้จองค่ะ",
     "I checked the trip records: nothing has been booked.", "There are no trips booked.",
     "None of them is booked.",
     "ถ้าจองแล้ว ดอรี่จะส่งใบเสร็จให้ค่ะ", "ตั้งแต่เมื่อวานแล้วค่ะ", "ไฟล์ที่คุณส่งมาแล้วค่ะ"],
)  # fmt: skip
def test_explanations_questions_and_negations_are_not(text):
    assert not claims_done(text)


def _profile_tool(saved):
    def travel_profile(update: str = "") -> dict:
        """Save the travel profile."""
        saved.append(update)
        return {"status": "saved"}

    return travel_profile


def test_a_claim_without_a_tool_call_is_asked_again_and_the_tool_runs(seeded, settings):
    saved = []
    llm = scripted(
        text_resp("I've saved that: passport name SOMCHAI/TEST."),
        fc_resp(("travel_profile", {"update": "passport name SOMCHAI/TEST"})),
        text_resp("Saved: passport name SOMCHAI/TEST."),
    )
    a = ask(settings, seeded, "Remember my passport name is SOMCHAI/TEST", guild_id=GUILD,
            is_exec=True, user_id=EXEC_ID, model=llm, extra_tools=[_profile_tool(saved)])  # fmt: skip
    assert saved == ["passport name SOMCHAI/TEST"] and a.rounds == 1
    assert a.text == "Saved: passport name SOMCHAI/TEST."
    assert CLAIM_CHECK in str(llm.requests[1].contents[-1])


def test_a_claim_made_twice_without_a_tool_is_replaced_by_the_truth(seeded, settings):
    saved = []
    llm = scripted(text_resp("I've saved that."), text_resp("Done, I've saved it for you."))
    a = ask(settings, seeded, "Remember my seat is 12A", guild_id=GUILD, is_exec=True,
            user_id=EXEC_ID, model=llm, extra_tools=[_profile_tool(saved)])  # fmt: skip
    assert saved == [] and a.rounds == 0 and len(llm.requests) == 2
    assert a.text.startswith("Sorry, I can't confirm that") and "saved" not in a.text


def test_a_claim_backed_by_a_tool_call_or_no_claim_is_left_alone(seeded, settings):
    saved = []
    llm = scripted(
        fc_resp(("travel_profile", {"update": "seat 12A"})), text_resp("I've saved that.")
    )
    a = ask(settings, seeded, "Remember my seat is 12A", guild_id=GUILD, is_exec=True,
            user_id=EXEC_ID, model=llm, extra_tools=[_profile_tool(saved)])  # fmt: skip
    assert a.text == "I've saved that." and len(llm.requests) == 2 and saved == ["seat 12A"]
    llm = scripted(text_resp("Hello! How can I help?"))
    a = ask(settings, seeded, "hi", guild_id=GUILD, is_exec=True, user_id=EXEC_ID, model=llm)
    assert a.text == "Hello! How can I help?" and len(llm.requests) == 1
