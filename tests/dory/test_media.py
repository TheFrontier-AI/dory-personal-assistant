# ruff: noqa: F811, F401
"""Audio/video understanding (dory.media), documents (dory.docs), memory notes, backlog and the
open_file tool. Offline: ffmpeg is real (skipped when missing), every model call is faked."""

import io
import shutil
import subprocess
import tempfile
from types import SimpleNamespace

import pytest
from conftest import GUILD, FakeClient, json_resp
from test_bot_multitenant import member
from test_memory_import import (
    FakeAtt,
    HChan,
    ago,
    att,
    hmsg,
    iw,
    logged,
    run_import,
    set_llm,
    settle,
)
from test_modes import STAFF, TMsg, bot, set_mode, world
from test_dory_files import wire

from dory import docs, media, media_bot
from dory.media import Seg, Transcript
from dory.qa import FileInput, MemoryTools, file_parts

needs_ffmpeg = pytest.mark.skipif(
    shutil.which("ffmpeg") is None or shutil.which("ffprobe") is None, reason="no ffmpeg"
)


# ------------------------------------------------------------------ helpers
class MAtt(SimpleNamespace):
    """Attachment with bytes."""

    data = b""

    async def read(self):
        return self.data


def matt(name, ctype, data=b"x", **kw):
    return MAtt(filename=name, content_type=ctype, size=len(data), data=data, **kw)


def tr(*lines, duration=60.0):
    return Transcript([Seg(s, sp, tx) for s, sp, tx in lines], duration, "th")


def make_media(tmp_path, name, args):
    out = tmp_path / name
    subprocess.run(
        ["ffmpeg", "-loglevel", "error", "-y", *args, str(out)], check=True, capture_output=True
    )
    return out.read_bytes()


SINE = ["-f", "lavfi", "-i", "sine=frequency=440:duration=3"]


def seg_resp(*segs, language="th"):
    return json_resp(
        {
            "language": language,
            "segments": [{"start": s, "speaker": sp, "text": tx} for s, sp, tx in segs],
        }
    )


def tiny_pdf(text="Invoice INV-1234 total 500 THB"):
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET"
    objs[3] = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def blank_pdf(pages=1):
    from pypdf import PdfWriter

    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(100, 100)
    b = io.BytesIO()
    w.write(b)
    return b.getvalue()


def make_docx():
    from docx import Document

    d = Document()
    d.sections[0].header.paragraphs[0].text = "HEAD"
    d.add_paragraph("ใบเสนอราคา 12,500 บาท")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "A", "B"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_xlsx(rows=3):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ยอดขาย"
    for i in range(rows):
        ws.append([f"item{i}", i * 10])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_pptx():
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "แผนไตรมาส"
    s.placeholders[1].text = "เป้าหมาย 1M"
    s.notes_slide.notes_text_frame.text = "พูดเรื่องงบ"
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()


# ------------------------------------------------------------------ detection
def test_media_detection_by_type_extension_and_voice_flag():
    assert media.media_kind(att("a.mp3", "audio/mpeg")) == "audio"
    assert media.media_kind(att("clip.MOV", "")) == "video"
    assert media.media_kind(att("x.opus", "application/octet-stream")) == "audio"
    assert media.media_kind(att("a.webm", "video/webm")) == "video"
    assert media.media_kind(att("a.png", "image/png")) is None
    assert media.media_kind(att("a.pdf", "application/pdf")) is None
    voice = FakeAtt(
        filename="voice-message.ogg", content_type="audio/ogg", waveform="AAAA", duration=4.2
    )
    assert media.media_kind(voice) == "voice"
    flagged = SimpleNamespace(flags=SimpleNamespace(voice=True), attachments=[att("v.ogg", "")])
    assert media.is_voice_message(flagged) and media.media_candidates(flagged) == [
        (0, flagged.attachments[0])
    ]
    plain = SimpleNamespace(flags=SimpleNamespace(voice=False), attachments=[att("a.mp3")])
    assert not media.is_voice_message(plain)


def test_time_helpers_and_parse_message_ref():
    assert media.fmt_time(222) == "3:42" and media.fmt_time(3725) == "1:02:05"
    assert media.parse_time("03:42") == 222 and media.parse_time("1:02:05") == 3725
    assert media.parse_time("junk") == 0
    assert media.parse_message_ref("https://discord.com/channels/1/22/333") == (22, 333)
    assert media.parse_message_ref("333") == (None, 333)
    assert media.parse_message_ref("hello") is None


# ------------------------------------------------------------------ chunk planning and merging
def test_plan_chunks_offsets_and_overlap():
    assert media.plan_chunks(300) == [(0.0, 300.0)]
    assert media.plan_chunks(600) == [(0.0, 600.0)]
    assert media.plan_chunks(1500) == [(0.0, 602.0), (600.0, 602.0), (1200.0, 300.0)]
    assert media.plan_chunks(601) == [(0.0, 601.0)]  # a tail of 1 s only repeats the overlap
    plan = media.plan_chunks(90 * 60)
    assert len(plan) == 9 and plan[-1] == (4800.0, 600.0)
    # every second is covered and consecutive windows overlap by 2 s
    assert all(b[0] == a[0] + 600 for a, b in zip(plan, plan[1:], strict=False))
    assert all(a[0] + a[1] == b[0] + 2 for a, b in zip(plan, plan[1:], strict=False))


def test_merge_drops_overlap_duplicates_and_keeps_order():
    a = [Seg(1, "A", "hello"), Seg(599, "A", "end of first")]
    b = [Seg(600.5, "A", "end of first"), Seg(601.5, "B", "dup in overlap"), Seg(605, "B", "next")]
    merged = media.merge_chunks([(0.0, a), (600.0, b)])
    assert [s.text for s in merged] == ["hello", "end of first", "next"]


def test_segments_from_raw_adds_offset_and_clamps():
    raw = media.RawTranscript(
        segments=[
            media.RawSeg(start="00:05", speaker="Speaker A", text=" สวัสดี  ครับ "),
            media.RawSeg(start="99:00", speaker="", text="late"),
            media.RawSeg(start="00:07", speaker="B", text="[inaudible]"),
            media.RawSeg(start="00:08", speaker="B", text="..."),
        ]
    )
    segs = media.segments_from_raw(raw, 600.0, 120.0)
    assert [(s.start, s.speaker, s.text) for s in segs] == [
        (605.0, "Speaker A", "สวัสดี ครับ"),
        (720.0, "Speaker A", "late"),
    ]


# ------------------------------------------------------------------ formatting
def test_prompt_text_cap_truncation_and_label():
    segs = [Seg(i, "Speaker A", "x" * 100) for i in range(600)]
    t = Transcript(segs, 3600.0, "th", truncated=True)
    body = media.prompt_text(t, 5000)
    assert len(body) < 5600 and "ตัดทอน" in body and "ถอดเสียงเฉพาะ 90 นาทีแรก" in body
    short = tr((5, "Speaker A", "สวัสดี"), duration=222)
    assert media.prompt_text(short) == "[0:05] Speaker A: สวัสดี"
    assert media.label_for(2, "a.mp3", short) == "[ถอดเสียงจากไฟล์ #2: a.mp3, ยาว 3:42]"
    assert "no speech" in media.prompt_text(Transcript([], 3.0))


def test_memory_note_format_and_cap():
    t = tr(
        (0, "Speaker A", "สวัสดี"),
        (5, "Speaker A", "เรื่องงบ"),
        (9, "Speaker B", "โอเค"),
        duration=222,
    )
    assert media.memory_note("audio", t) == "[เสียง: 3:42] Speaker A: สวัสดี เรื่องงบ | Speaker B: โอเค"
    assert media.memory_note("video", t, "ประชุมงบ").startswith("[วิดีโอ: 3:42] ประชุมงบ | Speaker A:")
    big = tr(*[(i, "Speaker A", "ก" * 200) for i in range(50)], duration=100)
    note = media.memory_note("audio", big)
    assert len(note) <= media.NOTE_CHARS and note.endswith("…")
    assert "ไม่พบเสียงพูด" in media.memory_note("audio", Transcript([], 5.0))


# ------------------------------------------------------------------ ffmpeg pipeline
@needs_ffmpeg
def test_extract_audio_from_video_and_ogg(tmp_path):
    mp4 = make_media(
        tmp_path, "a.mp4",
        [*SINE, "-f", "lavfi", "-i", "color=c=blue:s=64x64:d=3", "-shortest",
         "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
    )  # fmt: skip
    ogg = make_media(tmp_path, "a.ogg", [*SINE, "-c:a", "libvorbis"])
    for name, data in (("a.mp4", mp4), ("a.ogg", ogg)):
        src, dst = tmp_path / f"in_{name}", tmp_path / f"out_{name}.flac"
        src.write_bytes(data)
        dur = media.extract_audio(src, dst)
        assert 2.5 < dur < 3.5 and dst.read_bytes()[:4] == b"fLaC"
    with pytest.raises(media.MediaError):
        (tmp_path / "bad.bin").write_bytes(b"not media")
        media.extract_audio(tmp_path / "bad.bin", tmp_path / "bad.flac")


@needs_ffmpeg
def test_transcribe_media_with_mocked_gemini_and_cleanup(tmp_path, monkeypatch):
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    mp4 = make_media(
        tmp_path, "v.mp4",
        [*SINE, "-f", "lavfi", "-i", "color=c=red:s=64x64:d=3", "-shortest",
         "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
    )  # fmt: skip
    client = FakeClient([seg_resp(("00:01", "Speaker A", "สวัสดีครับ"), ("00:02", "B", "hello"))])
    t = media.transcribe_media(client, "gemini-x", mp4, "v.mp4")
    assert [(round(s.start), s.speaker, s.text) for s in t.segments] == [
        (1, "Speaker A", "สวัสดีครับ"),
        (2, "B", "hello"),
    ]
    assert 2.5 < t.duration < 3.5 and t.language == "th" and not t.truncated
    (model, contents, cfg) = client.models.calls[0]
    assert model == "gemini-x" and contents[0].inline_data.mime_type == "audio/flac"
    assert contents[0].inline_data.data[:4] == b"fLaC"  # audio only, no video sent
    assert not list(tmp_path.glob("dory-media-*"))  # temp files gone


@needs_ffmpeg
def test_chunked_transcription_adds_offsets_and_cleans_up_on_failure(tmp_path, monkeypatch):
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    monkeypatch.setattr(media, "CHUNK_WORKERS", 1)
    monkeypatch.setattr(media, "plan_chunks", lambda d: [(0.0, 1.8), (1.0, 2.0)])
    ogg = make_media(tmp_path, "c.ogg", [*SINE, "-c:a", "libvorbis"])
    client = FakeClient(
        [
            seg_resp(("00:00", "A", "first"), ("00:01", "A", "tail")),
            seg_resp(("00:00", "A", "dup of the overlap"), ("00:02", "B", "second")),
        ]
    )
    t = media.transcribe_media(client, "m", ogg, "c.ogg")
    assert [(s.start, s.text) for s in t.segments] == [
        (0.0, "first"),
        (1.0, "tail"),
        (3.0, "second"),
    ]
    boom = FakeClient([lambda c, cfg: (_ for _ in ()).throw(RuntimeError("secret"))] * 2)
    with pytest.raises(media.MediaError):
        media.transcribe_media(boom, "m", ogg, "c.ogg")
    assert not list(tmp_path.glob("dory-media-*"))
    with pytest.raises(media.MediaError):
        media.transcribe_media(client, "m", b"", "x.mp3")
    assert not list(tmp_path.glob("dory-media-*"))


def test_daily_budget():
    b = media.MediaBudget(2)  # 120 s
    assert b.available(1, "d1")
    b.add(1, "d1", 100)
    assert b.available(1, "d1")
    b.add(1, "d1", 30)
    assert not b.available(1, "d1") and b.available(2, "d1") and b.available(1, "d2")
    assert not media.MediaBudget(0).available(1, "d")


# ------------------------------------------------------------------ addressed messages
async def test_audio_sent_to_dory_reaches_the_agent_as_a_transcript_part(bot, world, monkeypatch):
    w = wire(bot, monkeypatch)
    got = []

    def fake(client, model, data, filename="", **kw):
        got.append((model, filename, data))
        return tr((5, "Speaker A", "งบโฆษณา 5 หมื่น"), (9, "Speaker B", "ตกลง"), duration=222)

    monkeypatch.setattr(media, "transcribe_media", fake)
    m = TMsg(
        world.g, world.chat, member(world.g, STAFF), "สรุปให้หน่อย",
        attachments=[matt("meeting.m4a", "audio/mp4", b"AUDIO")],
    )  # fmt: skip
    await bot.on_message(m)
    await settle(bot)
    ((q, kw),) = w.agent.calls
    (f,) = kw["files"]
    assert q == "สรุปให้หน่อย" and got[0][1:] == ("meeting.m4a", b"AUDIO")
    assert f.mime == "text/plain" and f.skip == ""
    parts, bad = file_parts([f])
    assert bad == [] and parts[0].text.startswith("[ถอดเสียงจากไฟล์ #0: meeting.m4a, ยาว 3:42]\n")
    assert "[0:05] Speaker A: งบโฆษณา 5 หมื่น" in parts[0].text
    assert m.replies == ["agent says"]


async def test_failed_transcription_is_a_skip_note_not_a_crash(bot, world, monkeypatch):
    w = wire(bot, monkeypatch)

    def boom(*a, **k):
        raise media.MediaError("nope")

    monkeypatch.setattr(media, "transcribe_media", boom)
    m = TMsg(world.g, world.chat, member(world.g, STAFF), "ฟังหน่อย",
             attachments=[matt("a.mp3", "audio/mpeg")])  # fmt: skip
    await bot.on_message(m)
    ((_q, kw),) = w.agent.calls
    assert kw["files"][0].skip == "audio could not be transcribed"


async def test_short_voice_message_is_the_question_itself(bot, world, monkeypatch):
    w = wire(bot, monkeypatch)
    monkeypatch.setattr(
        media, "transcribe_media",
        lambda *a, **k: tr((0, "Speaker A", "ยอดใช้จ่ายเดือนนี้เท่าไหร่"), duration=6),
    )  # fmt: skip
    voice = matt("voice-message.ogg", "audio/ogg", b"v", waveform="AA", duration=6.0)
    m = TMsg(world.g, world.chat, member(world.g, STAFF), "", attachments=[voice])
    m.flags = SimpleNamespace(voice=True)
    await bot.on_message(m)
    ((q, kw),) = w.agent.calls
    assert q == media_bot.VOICE_PREFIX + "ยอดใช้จ่ายเดือนนี้เท่าไหร่" and "files" not in kw


async def test_long_voice_message_or_one_with_text_stays_a_file(bot, world, monkeypatch):
    w = wire(bot, monkeypatch)
    monkeypatch.setattr(
        media, "transcribe_media",
        lambda *a, **k: tr((0, "Speaker A", "ยาวมาก"), duration=300),
    )  # fmt: skip
    voice = matt("voice-message.ogg", "audio/ogg", b"v", waveform="AA", duration=300.0)
    m = TMsg(world.g, world.chat, member(world.g, STAFF), "", attachments=[voice])
    m.flags = SimpleNamespace(voice=True)
    await bot.on_message(m)
    ((q, kw),) = w.agent.calls
    assert q == "" and len(kw["files"]) == 1


# ------------------------------------------------------------------ memory: audio / video
async def test_unaddressed_media_is_transcribed_in_the_background_into_memory(
    iw, bot, world, monkeypatch
):
    calls = []
    monkeypatch.setattr(
        media, "transcribe_media",
        lambda *a, **k: calls.append(1) or tr((0, "Speaker A", "สวัสดี"), (3, "Speaker B", "ครับ"), duration=222),
    )  # fmt: skip
    monkeypatch.setattr(media, "summarize_note", lambda *a: "ทักทายกัน")
    m = TMsg(world.g, iw.a, member(world.g, STAFF), "ฟังนี่",
             attachments=[matt("clip.mp4", "video/mp4")])  # fmt: skip
    await bot.on_message(m)
    await settle(bot)
    (rec,) = logged(bot)
    assert (
        rec.text == ("ฟังนี่ [วิดีโอ: 3:42] ทักทายกัน | Speaker A: สวัสดี | Speaker B: ครับ")
        and len(calls) == 1
    )


async def test_background_media_respects_the_daily_cap_and_the_flag(iw, bot, world, monkeypatch):
    calls = []
    monkeypatch.setattr(
        media, "transcribe_media",
        lambda *a, **k: calls.append(1) or tr((0, "Speaker A", "x"), duration=50),
    )  # fmt: skip
    monkeypatch.setattr(media, "summarize_note", lambda *a: "")
    bot.media_budget = media.MediaBudget(1)  # 60 s a day
    for _ in range(3):
        m = TMsg(
            world.g, iw.a, member(world.g, STAFF), "a", attachments=[matt("a.mp3", "audio/mpeg")]
        )
        await bot.on_message(m)
        await settle(bot)
    assert len(calls) == 2  # 50 s, then 100 s >= cap, the third is skipped
    texts = [r.text for r in logged(bot)]
    assert texts.count("a [audio: a.mp3]") == 1 and bot.media_stats["skipped"] == 1
    from dory.guild_settings import GuildSettingsUpdate

    bot.guilds_store.save_settings(
        GUILD, GuildSettingsUpdate(features=dict(media_memory=False)), "x"
    )
    bot.team._cache.clear()
    bot.media_budget = media.MediaBudget(60)
    m = TMsg(world.g, iw.a, member(world.g, STAFF), "b", attachments=[matt("b.mp3", "audio/mpeg")])
    await bot.on_message(m)
    await settle(bot)
    assert len(calls) == 2  # flag off: nothing new


async def test_off_channels_are_not_transcribed(iw, bot, world, monkeypatch):
    calls = []
    monkeypatch.setattr(media, "transcribe_media", lambda *a, **k: calls.append(1) or tr())
    set_mode(bot, 12, "off")
    m = TMsg(world.g, iw.b, member(world.g, STAFF), "x", attachments=[matt("a.mp3", "audio/mpeg")])
    await bot.on_message(m)
    await settle(bot)
    assert calls == []


async def test_addressed_media_is_transcribed_once_and_reused_for_memory(bot, world, monkeypatch):
    wire(bot, monkeypatch)
    calls = []
    monkeypatch.setattr(
        media, "transcribe_media",
        lambda *a, **k: calls.append(1) or tr((0, "Speaker A", "งานด่วน"), duration=30),
    )  # fmt: skip
    monkeypatch.setattr(media, "summarize_note", lambda *a: "")
    m = TMsg(
        world.g,
        world.chat,
        member(world.g, STAFF),
        "สรุป",
        attachments=[matt("a.mp3", "audio/mpeg")],
    )
    await bot.on_message(m)
    await settle(bot)
    assert len(calls) == 1
    (rec,) = [
        r
        for r in bot.memory_store.messages(
            GUILD, {world.chat.id}, "2026-01-01", "2026-12-31", limit=50
        )
    ]
    assert "[เสียง: 0:30] Speaker A: งานด่วน" in rec.text


async def test_media_backlog_counts_against_the_cap_newest_first(iw, bot, world, monkeypatch):
    order = []
    monkeypatch.setattr(
        media, "transcribe_media",
        lambda c, m, data, name="", **k: order.append(name) or tr((0, "Speaker A", name), duration=50),
    )  # fmt: skip
    monkeypatch.setattr(media, "summarize_note", lambda *a: "")
    bot.media_budget = media.MediaBudget(1)
    for i in range(1, 5):  # 4 is the newest
        hmsg(
            world, iw.a, i, ago(10 - i), f"m{i}", attachments=[matt(f"{i}.mp3", "audio/mpeg", b"z")]
        )
    await run_import(bot)
    assert order == ["4.mp3", "3.mp3"]  # 50 s + 50 s >= 60 s cap
    texts = {m.message_id: m.text for m in logged(bot)}
    assert "[เสียง: 0:50] Speaker A: 4.mp3" in texts[4] and "[audio: 2.mp3]" in texts[2]


# ------------------------------------------------------------------ documents
def test_pdf_text_layer_and_scanned_detection():
    dt = docs.extract_text(tiny_pdf(), "pdf")
    assert "INV-1234" in dt.text and dt.pages == 1 and not docs.is_scanned("pdf", dt)
    blank = docs.extract_text(blank_pdf(2), "pdf")
    assert blank.pages == 2 and docs.is_scanned("pdf", blank)
    assert docs.first_pages(blank_pdf(40), 30) != blank_pdf(40)
    assert len(docs.extract_text(docs.first_pages(blank_pdf(40)), "pdf").text) == 0


def test_office_extractors_and_the_shared_helper():
    kind, text, pages = docs.extract_document(make_docx(), "q.docx", None)
    assert kind == "docx" and "HEAD" in text and "ใบเสนอราคา 12,500 บาท" in text and "A | B" in text
    kind, text, pages = docs.extract_document(make_xlsx(), "s.xlsx", None)
    assert kind == "xlsx" and text.startswith("## ยอดขาย") and "item2,20" in text and pages == 1
    big = docs.extract_text(make_xlsx(700), "xlsx").text
    assert "item499" in big and "item500" not in big and "first 500 rows" in big
    kind, text, pages = docs.extract_document(make_pptx(), "p.pptx", None)
    assert (
        kind == "pptx" and "แผนไตรมาส" in text and "เป้าหมาย 1M" in text and "[โน้ต] พูดเรื่องงบ" in text
    )
    assert docs.extract_document(b"a,b\n1,2", "x.csv", "text/csv")[:2] == (
        "text",
        "a,b\n1,2",
    )
    with pytest.raises(docs.DocUnreadable):
        docs.extract_document(b"x", "old.xls", None)
    with pytest.raises(docs.DocUnreadable):
        docs.extract_document(b"x", "a.zip", None)
    if shutil.which("antiword") is None:
        with pytest.raises(docs.DocUnreadable, match="docx or PDF"):
            docs.extract_document(b"x", "old.doc", None)
    assert docs.clip_for_agent("x" * 40000).endswith("truncated]")


def test_doc_note_format_and_caps():
    dt = docs.DocText("ข้อความ " * 500, 3)
    note = docs.DocNote(doc_type="invoice", summary="ใบแจ้งหนี้ ABC", amounts="500 THB",
                        references="INV-1")  # fmt: skip
    out = docs.doc_note("inv.pdf", "pdf", dt, note)
    assert out.startswith("[PDF: inv.pdf, 3 หน้า, invoice] ใบแจ้งหนี้ ABC | ยอด: 500 THB | เลขที่: INV-1")
    assert "| ข้อความสำคัญ: ข้อความ" in out and len(out) <= docs.NOTE_CHARS
    key = out.split("ข้อความสำคัญ: ")[1]
    assert len(key) <= docs.KEY_TEXT_CHARS
    assert docs.doc_note("s.xlsx", "xlsx", docs.DocText("a", 2), None).startswith(
        "[XLSX: s.xlsx, 2 ชีต]"
    )
    assert (
        docs.doc_note("n.txt", "text", docs.DocText("a", 1), None)
        == "[ข้อความ: n.txt] | ข้อความสำคัญ: a"
    )


def test_scanned_pdf_goes_to_the_cheap_model_as_a_pdf_part():
    client = FakeClient(
        [json_resp({"doc_type": "receipt", "summary": "ใบเสร็จร้านกาแฟ", "amounts": "85 THB",
                    "text": "กาแฟ 85"})]
    )  # fmt: skip
    pdf = blank_pdf(40)
    dt = docs.extract_text(pdf, "pdf")
    note = docs.analyze(client, "flash-lite", pdf, "pdf", dt)
    (model, contents, _cfg) = client.models.calls[0]
    assert model == "flash-lite" and contents[0].inline_data.mime_type == "application/pdf"
    sent = docs.extract_text(contents[0].inline_data.data, "pdf")
    assert sent.pages == docs.OCR_MAX_PAGES  # first 30 pages only
    out = docs.doc_note("scan.pdf", "pdf", dt, note)
    assert "ใบเสร็จร้านกาแฟ" in out and "ข้อความสำคัญ: กาแฟ 85" in out
    with pytest.raises(ValueError):
        docs.analyze(client, "m", b"x" * (docs.MAX_DOC_BYTES + 1), "pdf", docs.DocText("", 1))
    # a text layer sends no PDF bytes, only its text
    client2 = FakeClient([json_resp({"summary": "ok"})])
    docs.analyze(client2, "m", tiny_pdf(), "pdf", docs.extract_text(tiny_pdf(), "pdf"))
    assert isinstance(client2.models.calls[0][1][0], str)


async def test_pdf_in_a_remembered_channel_is_read_into_the_memory_note(iw, bot, world):
    client = FakeClient([json_resp({"doc_type": "invoice", "summary": "ใบแจ้งหนี้ ABC",
                                    "amounts": "500 THB", "references": "INV-1234"})])  # fmt: skip
    set_llm(bot, client)
    m = TMsg(world.g, iw.a, member(world.g, STAFF), "",
             attachments=[matt("Invoice-1234.pdf", "application/pdf", tiny_pdf())])  # fmt: skip
    await bot.on_message(m)
    await settle(bot)
    (rec,) = logged(bot)
    assert rec.text.startswith("[PDF: Invoice-1234.pdf, 1 หน้า, invoice] ใบแจ้งหนี้ ABC")
    assert "INV-1234" in rec.text and "ข้อความสำคัญ: Invoice INV-1234 total 500 THB" in rec.text
    assert bot.doc_budget.used(GUILD, bot._day()) == 1


async def test_docx_xlsx_pptx_are_read_in_the_background_without_a_model_for_tiny_text(
    iw, bot, world
):
    set_llm(bot, FakeClient([]))
    m = TMsg(world.g, iw.a, member(world.g, STAFF), "ไฟล์",
             attachments=[matt("n.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", make_docx()),
                          matt("s.xlsx", "", make_xlsx())])  # fmt: skip
    await bot.on_message(m)
    await settle(bot)
    (rec,) = logged(bot)
    assert "[DOCX: n.docx, 1 หน้า] | ข้อความสำคัญ: HEAD ใบเสนอราคา 12,500 บาท" in rec.text
    assert "[XLSX: s.xlsx, 1 ชีต]" in rec.text and "item2,20" in rec.text


async def test_doc_cap_and_flag(iw, bot, world):
    from dory.guild_settings import GuildSettingsUpdate

    set_llm(bot, FakeClient([json_resp({"summary": "s"}) for _ in range(5)]))
    bot.doc_budget = media.DailyBudget(1)
    for i in range(3):
        m = TMsg(world.g, iw.a, member(world.g, STAFF), f"p{i}",
                 attachments=[matt(f"{i}.pdf", "application/pdf", tiny_pdf())])  # fmt: skip
        await bot.on_message(m)
        await settle(bot)
    texts = [r.text for r in logged(bot)]
    assert sum("[PDF:" in t for t in texts) == 1 and sum("[file: " in t for t in texts) == 2
    bot.doc_budget = media.DailyBudget(500)
    bot.guilds_store.save_settings(GUILD, GuildSettingsUpdate(features=dict(doc_memory=False)), "x")
    bot.team._cache.clear()
    m = TMsg(
        world.g,
        iw.a,
        member(world.g, STAFF),
        "z",
        attachments=[matt("z.pdf", "application/pdf", tiny_pdf())],
    )
    await bot.on_message(m)
    await settle(bot)
    assert "[file: z.pdf]" in logged(bot)[-1].text


async def test_receipt_pdf_reuses_the_receipt_note_and_is_never_ingested_twice(
    iw, bot, world, monkeypatch
):
    from dory import bot as botmod

    set_llm(bot, FakeClient([]))  # any model call would fail the test
    ingested = []
    monkeypatch.setattr(botmod, "ingest_attachment", lambda *a, **k: ingested.append(1))
    m = TMsg(
        world.g,
        iw.a,
        member(world.g, STAFF),
        "",
        attachments=[matt("r.pdf", "application/pdf", tiny_pdf())],
    )
    bot._agent_receipts[m.id] = ["[ใบเสร็จ] Makro | ยอด 500 THB | 2026-10-14 | วัตถุดิบ"]
    await bot.on_message(m)
    await settle(bot)
    (rec,) = logged(bot)
    assert rec.text.startswith("[PDF: r.pdf, 1 หน้า] [ใบเสร็จ] Makro | ยอด 500 THB")
    assert "ข้อความสำคัญ: Invoice INV-1234" in rec.text and ingested == []
    assert bot.doc_budget.used(GUILD, bot._day()) == 0


async def test_backlog_reads_old_pdfs_newest_first_within_the_page_cap(iw, bot, world):
    bot.doc_budget = media.DailyBudget(2)
    set_llm(bot, FakeClient([json_resp({"summary": f"sum{i}"}) for i in range(6)]))
    for i in range(1, 5):
        hmsg(world, iw.a, i, ago(10 - i), f"d{i}",
             attachments=[matt(f"{i}.pdf", "application/pdf", tiny_pdf(f"doc number {i}"))])  # fmt: skip
    hmsg(world, iw.a, 9, ago(1), "zip", attachments=[matt("a.zip", "application/zip", b"z")])
    await run_import(bot)
    texts = {m.message_id: m.text for m in logged(bot)}
    assert "doc number 4" in texts[4] and "doc number 3" in texts[3]
    assert "[file: 2.pdf]" in texts[2] and "[file: 1.pdf]" in texts[1]
    assert texts[9] == "zip [file: a.zip]"
    await run_import(bot)  # cap holds the same day
    assert "[file: 2.pdf]" in {m.message_id: m.text for m in logged(bot)}[2]


async def test_backlog_marks_unreadable_documents(iw, bot, world):
    set_llm(bot, FakeClient([]))
    hmsg(
        world, iw.a, 1, ago(3), "bad", attachments=[matt("x.pdf", "application/pdf", b"not a pdf")]
    )
    bot.memory_store.add_message(
        __import__("test_memory").mk(bot.s, 2, ago(2), "gone [file: g.pdf]", cid=10)
    )
    await run_import(bot)
    texts = {m.message_id: m.text for m in logged(bot)}
    assert "[file-unread: g.pdf]" in texts[2] and "[file: x.pdf]" in texts[1].replace("-unread", "")


# ------------------------------------------------------------------ addressed office files
async def test_docx_sent_to_dory_is_a_text_part_not_a_skip(bot, world, monkeypatch):
    w = wire(bot, monkeypatch)
    m = TMsg(world.g, world.chat, member(world.g, STAFF), "สรุปหน่อย",
             attachments=[matt("q.docx", "", make_docx()), matt("p.pptx", "", make_pptx()),
                          matt("s.xlsx", "", make_xlsx()), matt("old.xls", "", b"x")])  # fmt: skip
    await bot.on_message(m)
    ((_q, kw),) = w.agent.calls
    docx, pptx, xlsx, xls = kw["files"]
    assert docx.skip == pptx.skip == xlsx.skip == ""
    parts, bad = file_parts([docx, pptx, xlsx])
    assert (
        bad == []
        and "ใบเสนอราคา 12,500 บาท" in parts[0].text
        and parts[0].text.startswith("[ไฟล์ #0: q.docx")
    )
    assert "แผนไตรมาส" in parts[1].text and "item2,20" in parts[2].text
    assert "old .xls" in xls.skip or ".xls" in xls.skip


# ------------------------------------------------------------------ open_file
def fake_channel(world, iw):
    return iw.a, iw.b


async def test_open_file_access_control_happy_path_and_cap(iw, bot, world):
    set_llm(bot, FakeClient([]))
    long_text = "word " * 12000
    allowed_msg = hmsg(world, iw.a, 5, ago(1), "doc",
                       attachments=[matt("a.pdf", "application/pdf", tiny_pdf(long_text[:200]))])  # fmt: skip
    secret = hmsg(
        world, iw.b, 6, ago(1), "secret", attachments=[matt("s.pdf", "application/pdf", tiny_pdf())]
    )
    allowed_msg.jump_url = secret.jump_url = "https://discord.com/x"
    for m in (allowed_msg, secret):
        bot.memory_store.add_message(
            __import__("test_memory").mk(bot.s, m.id, ago(1), "doc [file: x.pdf]", cid=m.channel.id)
        )
    allowed = {iw.a.id: iw.a.name}  # the asker may read channel 10 only
    ok = await bot._open_file(world.g, allowed, iw.a.id, 5, 0)
    assert "Invoice" not in ok.get("error", "") and ok["file"] == "a.pdf" and "word" in ok["text"]
    refused = await bot._open_file(world.g, allowed, iw.b.id, 6, 0)
    assert "not allowed" in refused["error"] and "text" not in refused
    # a bare id outside the readable channels is not found either
    assert "not found" in (await bot._open_file(world.g, allowed, None, 6, 0))["error"]
    assert (await bot._open_file(world.g, allowed, None, 5, 0))["file"] == "a.pdf"
    assert "out of range" in (await bot._open_file(world.g, allowed, iw.a.id, 5, 3))["error"]
    assert "message not found" in (await bot._open_file(world.g, allowed, iw.a.id, 404, 0))["error"]
    # the 30k char cap
    hmsg(
        world,
        iw.a,
        7,
        ago(1),
        "big",
        attachments=[matt("big.txt", "text/plain", ("y" * 50000).encode())],
    )
    out = await bot._open_file(world.g, allowed, iw.a.id, 7, 0)
    assert out["truncated"] and len(out["text"]) == docs.OPEN_FILE_CHARS


async def test_open_file_transcribes_audio_and_ocrs_scanned_pdfs(iw, bot, world, monkeypatch):
    monkeypatch.setattr(
        media, "transcribe_media", lambda *a, **k: tr((7, "Speaker A", "เรื่องสัญญา"), duration=20)
    )
    hmsg(world, iw.a, 5, ago(1), "a", attachments=[matt("c.mp3", "audio/mpeg", b"a")])
    out = await bot._open_file(world.g, {iw.a.id: "x"}, iw.a.id, 5, 0)
    assert out["text"] == "[0:07] Speaker A: เรื่องสัญญา"
    set_llm(bot, FakeClient([lambda c, cfg: SimpleNamespace(text="ข้อความจากการสแกน")]))
    hmsg(
        world, iw.a, 6, ago(1), "b", attachments=[matt("scan.pdf", "application/pdf", blank_pdf())]
    )
    out = await bot._open_file(world.g, {iw.a.id: "x"}, iw.a.id, 6, 0)
    assert out["text"] == "ข้อความจากการสแกน"


def test_open_file_tool_is_exposed_only_with_an_opener_and_parses_links(seeded, settings):
    from dory.qa import LedgerTools, build_tools

    mem = MemoryTools(None, GUILD, {10: "a"}, settings)
    assert mem.open_file("nonsense")["error"].startswith("give the message link")
    assert "not available" in mem.open_file("1")["error"]
    seen = []
    mem = MemoryTools(
        None, GUILD, {10: "a"}, settings, opener=lambda c, m, i: seen.append((c, m, i)) or {"ok": 1}
    )
    assert mem.open_file("https://discord.com/channels/1/10/99", 2) == {"ok": 1} and seen == [
        (10, 99, 2)
    ]
    ledger = LedgerTools(seeded, GUILD)
    names = lambda m: {t.__name__ for t in build_tools(ledger, None, m)}  # noqa: E731
    assert "open_file" in names(mem)
    assert "open_file" not in names(MemoryTools(None, GUILD, {10: "a"}, settings))
